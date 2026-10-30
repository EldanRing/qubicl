#!/usr/bin/env python3
"""Convert a CSV file to a styled .xlsx workbook with type inference.

Conservative homogeneous-column inference (disable with --no-infer):
  int, float, bool ("true"/"false", case-insensitive), ISO date
  (YYYY-MM-DD) and ISO datetime. Leading zeros, long numbers and mixed-type
  columns remain text. Formulas require --allow-formulas or an explicit column
  type. --column-types accepts a JSON map of column letters to types.

Styling applied by default (disable with --plain):
  bold header row with a light fill, frozen top row, autofilter over the
  data range, and column widths sized to the longest cell (capped at 60).

Usage:
  csv_to_xlsx.py data.csv out.xlsx
  csv_to_xlsx.py data.csv out.xlsx --sheet-name Import --encoding cp1252
  csv_to_xlsx.py data.csv out.xlsx --delimiter ';' --no-infer
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
import re
import math
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_COL_WIDTH = 60
COL_PADDING = 2
DEFAULT_COL_WIDTH = 8
# Excel worksheet limits, independent of what openpyxl permits or truncates.
# https://support.microsoft.com/en-us/excel/excel-specifications-and-limits
MAX_ROWS = 1_048_576
MAX_COLUMNS = 16_384
MAX_CELL_CHARACTERS = 32_767
MAX_FORMULA_CHARACTERS = 8_192


def read_rows(stream, delimiter):
    rows = []
    reader = csv.reader(stream, delimiter=delimiter)
    try:
        for index, row in enumerate(reader, start=1):
            if index > MAX_ROWS:
                raise ValueError(f"CSV row {index} exceeds Excel's {MAX_ROWS}-row limit")
            if len(row) > MAX_COLUMNS:
                raise ValueError(f"CSV row {index}, column {MAX_COLUMNS + 1} (XFE{index}) exceeds Excel's {MAX_COLUMNS}-column limit")
            for column, value in enumerate(row, start=1):
                if len(value) > MAX_CELL_CHARACTERS:
                    coordinate = f"{get_column_letter(column)}{index}"
                    raise ValueError(f"CSV row {index}, column {column} ({coordinate}) contains {len(value)} characters; Excel allows {MAX_CELL_CHARACTERS} per cell")
            rows.append(row)
    except csv.Error as exc:
        raise ValueError(f"CSV record {len(rows) + 1}, source line {reader.line_num}: {exc}") from exc
    return rows


def save_verified(workbook, output, input_rows):
    """Publish only a complete archive whose text and formulas survived saving."""
    destination = Path(output)
    with tempfile.NamedTemporaryFile(prefix='.qubicl-csv-', suffix='.xlsx',
                                     dir=destination.parent, delete=False) as staged:
        staged_path = Path(staged.name)
    try:
        workbook.save(staged_path)
        saved = load_workbook(staged_path, data_only=False)
        try:
            expected, actual = workbook.active, saved.active
            if (actual.max_row, actual.max_column) != (expected.max_row, expected.max_column):
                raise ValueError("Saved workbook dimensions differ from the converted CSV")
            # Visit CSV cells only: a single wide row must not expand every
            # narrower row into a large blank rectangle during verification.
            for row_index, source_row in enumerate(input_rows, start=1):
                for column_index in range(1, len(source_row) + 1):
                    original = expected.cell(row_index, column_index)
                    cell = actual.cell(row_index, column_index)
                    value = original.value
                    # Empty CSV strings are represented as blank cells by Excel.
                    if isinstance(value, str) and value and (cell.value != value or cell.data_type != original.data_type):
                        raise ValueError(f"Saved workbook changed text or formula at {original.coordinate}; output was not replaced")
                    if value not in (None, '') and cell.value is None:
                        raise ValueError(f"Saved workbook lost a value at {original.coordinate}; output was not replaced")
        finally:
            saved.close()
        os.replace(staged_path, destination)
    finally:
        staged_path.unlink(missing_ok=True)


def infer(text):
    if text == "":
        return None
    if text != text.strip() or re.fullmatch(r"[+-]?0[0-9]+", text) or len(re.sub(r"[^0-9]", "", text)) > 15:
        return text
    low = text.lower()
    if low in ("true", "false"):
        return low == "true"
    for caster in (int, float):
        try:
            value = caster(text)
            return value if not isinstance(value, float) or math.isfinite(value) else text
        except ValueError:
            pass
    for parser in (date.fromisoformat, datetime.fromisoformat):
        try:
            return parser(text)
        except ValueError:
            pass
    return text


def main(argv=None):
    ap = argparse.ArgumentParser(description="CSV -> styled .xlsx converter.")
    ap.add_argument("csv_file", help="input CSV path")
    ap.add_argument("output", help="output .xlsx path")
    ap.add_argument("--sheet-name", default="Sheet1")
    ap.add_argument("--encoding", default="utf-8-sig",
                    help="CSV file encoding (default UTF-8, accepting an initial BOM)")
    ap.add_argument("--delimiter", default=",")
    ap.add_argument("--no-infer", action="store_true",
                    help="keep every cell as a string")
    ap.add_argument("--plain", action="store_true",
                    help="skip header styling / freeze / autofilter")
    ap.add_argument("--column-types", default="{}", help='JSON mapping of column letters to text, number, boolean, date or formula')
    ap.add_argument("--allow-formulas", action="store_true", help="Interpret = expressions deliberately; default is literal text")
    ap.add_argument("--empty-cells", choices=("blank", "reject"), default="blank",
                    help="Empty explicitly typed cells become blanks by default; reject enforces nonempty typed values")
    args = ap.parse_args(argv)
    column_types = json.loads(args.column_types)
    if not isinstance(column_types, dict) or any(v not in {"text", "number", "boolean", "date", "formula"} for v in column_types.values()):
        raise ValueError("Invalid column types")

    with open(args.csv_file, newline="", encoding=args.encoding) as fh:
        rows = read_rows(fh, args.delimiter)

    wb = Workbook()
    ws = wb.active
    ws.title = args.sheet_name
    conversions = []
    # Infer a type only when all nonempty values in a column agree.
    inferred = {}
    for col in range(max((len(row) for row in rows), default=0)):
        types = {type(infer(row[col])) for row in rows[1:] if col < len(row) and row[col] != ""}
        inferred[col] = len(types) == 1 and str not in types
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            kind = column_types.get(get_column_letter(j + 1)) if i else "text"
            value = text
            try:
                if i and kind not in (None, "text") and text == "":
                    if args.empty_cells == "reject": raise ValueError("Empty typed value is not allowed")
                    value = None
                elif i and kind == "number":
                    value = float(text) if any(c in text.lower() for c in ".e") else int(text)
                    if not math.isfinite(value): raise ValueError("Numeric overflow or nonfinite value cannot be stored in Excel")
                elif i and kind == "boolean":
                    if text.lower() not in ("true", "false"): raise ValueError("Invalid boolean")
                    value = text.lower() == "true"
                elif i and kind == "date": value = date.fromisoformat(text)
                elif i and kind is None and not args.no_infer and inferred.get(j): value = infer(text)
            except (ValueError, OverflowError) as exc:
                coordinate = f"{get_column_letter(j + 1)}{i + 1}"
                raise ValueError(f"CSV row {i + 1}, column {j + 1} ({coordinate}), type {kind}: {exc}") from exc
            if isinstance(value, str) and value.startswith('=') and (args.allow_formulas or kind == 'formula') and len(value) > MAX_FORMULA_CHARACTERS:
                coordinate = f"{get_column_letter(j + 1)}{i + 1}"
                raise ValueError(f"CSV row {i + 1}, column {j + 1} ({coordinate}) formula contains {len(value)} characters; Excel allows {MAX_FORMULA_CHARACTERS}")
            cell = ws.cell(i + 1, j + 1, value)
            if isinstance(value, str) and not (args.allow_formulas or kind == "formula"):
                cell.data_type = "s"
            if not isinstance(value, str): conversions.append({"cell": cell.coordinate, "type": type(value).__name__})

    if rows and not args.plain:
        header_font = Font(bold=True)
        header_fill = PatternFill("solid", fgColor="DDEBF7")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col_idx in range(1, ws.max_column + 1):
            longest = max((len(str(r[col_idx - 1])) for r in rows
                           if len(r) >= col_idx), default=DEFAULT_COL_WIDTH)
            ws.column_dimensions[get_column_letter(col_idx)].width = \
                min(longest + COL_PADDING, MAX_COL_WIDTH)

    save_verified(wb, args.output, rows)
    print(json.dumps({"ok": True, "output": args.output,
                      "rows": len(rows), "conversions": conversions}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": str(exc)}), file=sys.stderr)
        sys.exit(1)
