#!/usr/bin/env python3
"""Read an .xlsx workbook: inventory, JSON/CSV dumps, formula listing.

Modes (pick one):
  --sheets     JSON inventory: sheet names, dimensions, row/col counts
  --json       dump one sheet's rows as a JSON array of arrays
  --cells      typed cell records with stored values, formats and cache status
  --csv        dump one sheet as CSV to stdout or --out
  --formulas   JSON list of formula cells {"cell", "formula", "cached"};
               array/data-table entries also include formula_type, ref and
               data-table attributes (inputs and boolean flags)
  --notes      JSON list of cell notes/comments across sheets
  --names      JSON map of workbook defined names

Options:
  --sheet NAME     sheet to dump (default: active sheet)
  --data-only      load cached formula RESULTS instead of formula strings.
                   Caveat: openpyxl never computes formulas; cached values
                   exist only if the file was last saved by Excel/LibreOffice.
  --encoding ENC   encoding for --csv --out files (default utf-8)
  --out PATH       write --csv output to a file instead of stdout

Elapsed-time cells export as signed ISO 8601 durations. JSON/CSV sheet dumps
represent array/data-table formulas as JSON objects with explicit metadata.

Usage:
  xlsx_read.py book.xlsx --sheets
  xlsx_read.py book.xlsx --json --sheet Data
  xlsx_read.py book.xlsx --csv --sheet Data --out data.csv
  xlsx_read.py book.xlsx --formulas
  xlsx_read.py book.xlsx --notes
  xlsx_read.py book.xlsx --names
"""
from __future__ import annotations

import argparse
import csv
import json
import posixpath
import sys
from zipfile import ZipFile
from datetime import date, datetime, time, timedelta

from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.xml.functions import fromstring
from xlsx_formula_values import formula_payload, is_typed_formula


def jsonable(value):
    if isinstance(value, timedelta):
        # ISO 8601 duration, including negative values and subsecond precision;
        # elapsed hours must not wrap at midnight like a time-of-day value.
        micros = (value.days * 86400 + value.seconds) * 1000000 + value.microseconds
        sign = "-" if micros < 0 else ""
        seconds, fraction = divmod(abs(micros), 1000000)
        days, seconds = divmod(seconds, 86400)
        hours, seconds = divmod(seconds, 3600)
        minutes, seconds = divmod(seconds, 60)
        fractional = f"{seconds}.{fraction:06d}".rstrip("0") if fraction else str(seconds)
        clock = (f"{hours}H" if hours else "") + (f"{minutes}M" if minutes else "")
        if seconds or fraction or (not days and not clock): clock += fractional + "S"
        return sign + "P" + (f"{days}D" if days else "") + ("T" + clock if clock else "")
    if is_typed_formula(value):
        return formula_payload(value)
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def sheet_rows(ws):
    return [[jsonable(c) for c in row] for row in ws.iter_rows(values_only=True)]


def cmd_sheets(wb):
    info = []
    for ws in wb.worksheets:
        info.append({
            "name": ws.title,
            "dimensions": ws.dimensions,
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "merged": [str(r) for r in ws.merged_cells.ranges],
            "charts": len(getattr(ws, "_charts", [])),
            "freeze_panes": ws.freeze_panes,
            "autofilter": ws.auto_filter.ref,
            "tables": {t.displayName: t.ref for t in ws.tables.values()},
            "protected": bool(ws.protection.sheet),
        })
    names = {name: dn.attr_text for name, dn in wb.defined_names.items()}
    print(json.dumps({"sheets": info, "defined_names": names},
                     ensure_ascii=False, indent=2))


def cmd_notes(wb, sheet):
    out = []
    sheets = [sheet] if sheet else wb.sheetnames
    for name in sheets:
        for row in wb[name].iter_rows():
            for cell in row:
                if cell.comment is not None:
                    out.append({"sheet": name, "cell": cell.coordinate,
                                "text": cell.comment.text,
                                "author": cell.comment.author})
    print(json.dumps({"notes": out}, ensure_ascii=False, indent=2))


def cmd_names(wb):
    names = {name: dn.attr_text for name, dn in wb.defined_names.items()}
    print(json.dumps({"defined_names": names}, ensure_ascii=False, indent=2))


def stored_cells(path):
    """Read original XML values without recalculation or guessing missing caches.

    In particular, an empty cached string is available, whereas the empty numeric
    <v/> written by openpyxl for an uncalculated formula is not a result.
    """
    main = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    relation = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    result = {}
    with ZipFile(path) as archive:
        root_relations = fromstring(archive.read('_rels/.rels'))
        workbook_path = next(item.get('Target').lstrip('/') for item in root_relations
                             if item.get('Type', '').endswith('/officeDocument') and item.get('TargetMode') != 'External')
        workbook = fromstring(archive.read(workbook_path))
        directory, filename = posixpath.split(workbook_path)
        relationships = fromstring(archive.read(posixpath.join(directory, '_rels', filename + '.rels')))
        targets = {item.get('Id'): item.get('Target') for item in relationships
                   if item.get('TargetMode') != 'External'}
        for sheet in workbook.findall(f'{main}sheets/{main}sheet'):
            target = targets.get(sheet.get(relation))
            if not target:
                continue
            part = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(directory, target))
            cells = {}
            row_index = 0
            for row in fromstring(archive.read(part)).iter(f'{main}row'):
                row_index = int(row.get('r', row_index + 1))
                column = 0
                for cell in row.findall(f'{main}c'):
                    column = coordinate_to_tuple(cell.get('r'))[1] if cell.get('r') else column + 1
                    coordinate = cell.get('r', f'{get_column_letter(column)}{row_index}')
                    value = cell.find(f'{main}v')
                    kind = cell.get('t', 'n')
                    cells[coordinate] = {
                        'rawValue': value.text if value is not None else None,
                        'storageType': kind,
                        'valuePresent': value is not None and (kind == 'str' or value.text not in (None, '')),
                    }
            result[sheet.get('name')] = cells
    return result


def formula_details(cell):
    typed = formula_payload(cell.value)
    if typed:
        return typed
    if cell.data_type == 'f':
        return {'type': 'formula', 'formula_type': 'normal', 'text': cell.value, 'ref': cell.coordinate}
    return None


def cache_details(cell, cached, stored):
    available = bool(stored.get('valuePresent'))
    value = cached.value
    if available and value is None and stored.get('storageType') == 'str':
        value = ''
    return {'cached': jsonable(value) if available else None,
            'cachedValueAvailable': available,
            'cacheStatus': 'available' if available else 'missing',
            'warnings': [] if available else [{'code': 'missing_formula_cache', 'cell': cell.coordinate}]}


def cmd_cells(path, sheet):
    wb = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    stored = stored_cells(path)
    out = []
    try:
        for name in [sheet] if sheet else [wb.active.title]:
            ws = wb[name]
            # Keep sparse worksheets sparse: formatting a distant cell must not
            # cause inspection to enumerate a huge empty bounding rectangle.
            for cell in sorted(ws._cells.values(), key=lambda item: (item.row, item.column)):
                if cell.value is None and not cell.has_style:
                    continue
                formula = formula_details(cell)
                value = cell.value
                kind = ('formula' if formula else 'duration' if isinstance(value, timedelta)
                        else 'datetime' if isinstance(value, datetime) else 'date' if isinstance(value, date)
                        else 'time' if isinstance(value, time) else 'boolean' if isinstance(value, bool)
                        else 'error' if cell.data_type == 'e' else 'number' if isinstance(value, (int, float))
                        else 'blank' if value is None else 'text')
                original = stored.get(name, {}).get(cell.coordinate, {})
                record = {'sheet': name, 'cell': cell.coordinate, 'type': kind,
                          'value': jsonable(value), 'numberFormat': cell.number_format,
                          'rawValue': original.get('rawValue'), 'storageType': original.get('storageType')}
                if formula:
                    record.update(formula=formula, **cache_details(cell, values[name][cell.coordinate], original))
                out.append(record)
        print(json.dumps({'cells': out, 'recalculated': False, 'displayValues': 'not_rendered',
                          'note': 'value is typed data; rawValue is original OOXML storage (shared strings use indexes). numberFormat is not rendered display text. Render a copy when appearance matters.'}, ensure_ascii=False))
    finally:
        wb.close()
        values.close()


def cmd_formulas(path, sheet):
    wb_f = load_workbook(path, data_only=False)
    wb_v = load_workbook(path, data_only=True)
    stored = stored_cells(path)
    out = []
    sheets = [sheet] if sheet else wb_f.sheetnames
    for name in sheets:
        ws_f, ws_v = wb_f[name], wb_v[name]
        for row in ws_f.iter_rows():
            for cell in row:
                typed = formula_details(cell)
                if typed:
                    out.append({
                        "sheet": name,
                        "cell": cell.coordinate,
                        "formula": typed.get("text") if typed else cell.value,
                        **({key: value for key, value in typed.items() if key not in {"type", "text"}} if typed else {}),
                        "numberFormat": cell.number_format,
                        **cache_details(cell, ws_v[cell.coordinate], stored.get(name, {}).get(cell.coordinate, {})),
                    })
    print(json.dumps({"formulas": out}, ensure_ascii=False, indent=2))
    wb_f.close()
    wb_v.close()


def main(argv=None):
    ap = argparse.ArgumentParser(description="Read/inspect an .xlsx workbook.")
    ap.add_argument("file", help="path to .xlsx file")
    mode = ap.add_mutually_exclusive_group(required=True)
    mode.add_argument("--sheets", action="store_true")
    mode.add_argument("--json", action="store_true")
    mode.add_argument("--cells", action="store_true")
    mode.add_argument("--csv", action="store_true")
    mode.add_argument("--formulas", action="store_true")
    mode.add_argument("--notes", action="store_true")
    mode.add_argument("--names", action="store_true")
    ap.add_argument("--sheet", help="sheet name (default: active)")
    ap.add_argument("--data-only", action="store_true",
                    help="return cached formula results (see module docstring)")
    ap.add_argument("--encoding", default="utf-8")
    ap.add_argument("--out", help="output file for --csv")
    args = ap.parse_args(argv)

    if args.cells:
        cmd_cells(args.file, args.sheet)
        return 0
    if args.formulas:
        cmd_formulas(args.file, args.sheet)
        return 0

    wb = load_workbook(args.file, data_only=args.data_only)
    if args.sheets:
        cmd_sheets(wb)
        return 0
    if args.notes:
        cmd_notes(wb, args.sheet)
        return 0
    if args.names:
        cmd_names(wb)
        return 0

    ws = wb[args.sheet] if args.sheet else wb.active
    rows = sheet_rows(ws)
    if args.json:
        print(json.dumps({"sheet": ws.title, "rows": rows}, ensure_ascii=False))
    else:  # --csv
        rows = [[json.dumps(value, ensure_ascii=False) if isinstance(value, dict)
                 else value for value in row] for row in rows]
        if args.out:
            with open(args.out, "w", newline="", encoding=args.encoding) as fh:
                csv.writer(fh).writerows(
                    [["" if v is None else v for v in r] for r in rows])
            print(json.dumps({"ok": True, "out": args.out, "rows": len(rows)}))
        else:
            w = csv.writer(sys.stdout)
            for r in rows:
                w.writerow(["" if v is None else v for v in r])
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": str(exc)}), file=sys.stderr)
        sys.exit(1)
