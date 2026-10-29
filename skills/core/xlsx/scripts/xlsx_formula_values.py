"""Typed OOXML formula values shared by spreadsheet readers and editors."""
from copy import deepcopy

from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.utils import coordinate_to_tuple, get_column_letter, range_boundaries


def is_typed_formula(value):
    return isinstance(value, (ArrayFormula, DataTableFormula))


def formula_payload(value):
    """JSON-safe metadata; retain every supported OOXML data-table attribute."""
    if isinstance(value, ArrayFormula):
        return {"type": "formula", "formula_type": "array", "ref": value.ref,
                "text": value.text}
    if isinstance(value, DataTableFormula):
        return {"type": "formula", "formula_type": "dataTable", "ref": value.ref,
                "attributes": {key: (getattr(value, key) if key in {"r1", "r2"}
                                     else getattr(value, key) in (True, 1, "1", "true")) for key in
                               ("ca", "dt2D", "dtr", "r1", "r2", "del1", "del2")}}
    return None


def rewrite_typed_formula(value, rewrite, home):
    """Copy a formula before editing it; worksheet copies must not share it."""
    result = deepcopy(value)
    if isinstance(result, ArrayFormula) and result.text:
        result.text = rewrite(result.text, home)
    if isinstance(result, DataTableFormula):
        # openpyxl's constructor retains XML boolean strings on load.
        for key in ("ca", "dt2D", "dtr", "del1", "del2"):
            setattr(result, key, getattr(result, key) in (True, 1, "1", "true"))
        for index in (1, 2):
            key = f"r{index}"
            reference = getattr(result, key)
            if reference:
                updated = rewrite(reference, home)
                if "#REF!" in updated:
                    setattr(result, key, None)
                    setattr(result, f"del{index}", True)
                else:
                    setattr(result, key, updated)
        result.ca = True
    return result


def snapshot_typed_formulas(wb):
    return [(ws, cell.coordinate, deepcopy(cell.value))
            for ws in wb.worksheets for cell in ws._cells.values()
            if is_typed_formula(cell.value)]


def restore_typed_formulas(snapshots, target, axis, idx, count, delete,
                           rewrite, shift_point, shift_range):
    """Move formula anchors/ranges, promoting a surviving anchor on deletion."""
    changes = []
    for ws, old_anchor, original in snapshots:
        value = rewrite_typed_formula(original, rewrite, ws.title)
        anchor = old_anchor
        if ws is target:
            value.ref = shift_range(original.ref, axis, idx, count, delete)
            if value.ref is None:
                changes.append({"sheet": ws.title, "cell": old_anchor,
                                "from": formula_payload(original), "to": None})
                continue
            row, col = coordinate_to_tuple(old_anchor)
            row = shift_point(row, idx, count, delete) if axis == "rows" else row
            col = shift_point(col, idx, count, delete) if axis == "cols" else col
            if row is None or col is None:
                col, row, _, _ = range_boundaries(value.ref)
            anchor = f"{get_column_letter(col)}{row}"
        ws[anchor] = value
        if old_anchor != anchor or formula_payload(original) != formula_payload(value):
            changes.append({"sheet": ws.title, "cell": anchor, "original_cell": old_anchor,
                            "from": formula_payload(original), "to": formula_payload(value)})
    return changes
