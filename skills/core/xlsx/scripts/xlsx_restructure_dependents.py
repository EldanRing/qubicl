"""Preserve workbook dependents around a structural worksheet edit.

Call snapshot_dependents before moving cells, then apply_dependents after cell
and table formulas have been rewritten. The caller supplies its reference and
coordinate rewriters so every dependent follows the same structural semantics.
"""
from copy import deepcopy
from datetime import date, datetime, time, timedelta
from numbers import Number
import re

from openpyxl.cell.rich_text import CellRichText
from openpyxl.chart.data_source import (AxDataSource, NumRef, NumData, NumVal, StrRef, StrData,
                                      StrVal, MultiLevelStrRef, MultiLevelStrData, Level)
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor, TwoCellAnchor
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries, coordinate_to_tuple
from openpyxl.utils.datetime import to_excel, time_to_days, timedelta_to_days
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.print_settings import PrintArea


def snapshot_dependents(wb):
    """Capture pre-edit formula origins and anchors without copying cell data."""
    return [dict(
        sheet=ws,
        validations=deepcopy(ws.data_validations.dataValidation),
        formatting=[(deepcopy(cf), deepcopy(ws.conditional_formatting[cf]))
                    for cf in ws.conditional_formatting],
        drawings=[(item, deepcopy(item.anchor)) for item in [*ws._charts, *ws._images]],
        names=[(name, dn, dn.attr_text) for name, dn in ws.defined_names.items()],
        print_area=[str(rng) for rng in PrintArea.from_string(ws.print_area).sorted()] if ws.print_area else [],
        print_rows=ws.print_title_rows, print_cols=ws.print_title_cols,
    ) for ws in wb.worksheets]


def _range_origin(sqref):
    first = sqref.sorted()[0]
    return first.min_col, first.min_row


def _shift_applied(sqref, axis, idx, n, delete, shift_range):
    """Return new ranges and the first surviving cell in ORIGINAL coordinates."""
    ranges, origins = [], []
    for cell_range in sqref.sorted():
        new = shift_range(str(cell_range), axis, idx, n, delete)
        if not new:
            continue
        col, row, _, _ = range_boundaries(new)
        # Invert this edit for the first surviving cell; the formula's relative
        # origin must move before its referenced coordinates are rewritten.
        old_col, old_row = col, row
        if axis == "rows":
            old_row = row + n if delete and row >= idx else row - n if not delete and row >= idx+n else row
        else:
            old_col = col + n if delete and col >= idx else col - n if not delete and col >= idx+n else col
        ranges.append(new)
        origins.append((col, row, old_col, old_row))
    if not ranges:
        return None, None
    first = min(origins, key=lambda item: (item[0], item[1]))
    return " ".join(ranges), first[2:]


def _formula(value, home, rewrite, old_origin=None, surviving_origin=None):
    if not isinstance(value, str) or not value:
        return value
    if old_origin and surviving_origin and old_origin != surviving_origin:
        original = f"{get_column_letter(old_origin[0])}{old_origin[1]}"
        surviving = f"{get_column_letter(surviving_origin[0])}{surviving_origin[1]}"
        prefix = "" if value.startswith("=") else "="
        translator = Translator(prefix + value, origin=original)
        try:
            value = translator.translate_formula(surviving)
        except TranslatorError:
            # A surviving disjoint area can have a relative reference above or
            # left of the grid. Preserve the rest of its formula (e.g. IFERROR).
            row_delta = surviving_origin[1] - old_origin[1]
            col_delta = surviving_origin[0] - old_origin[0]
            parts = []
            for token in translator.get_tokens():
                text = token.value
                if token.subtype == "RANGE":
                    try: text = Translator.translate_range(text, row_delta, col_delta)
                    except TranslatorError: text = "#REF!"
                parts.append(text)
            value = "=" + "".join(parts)
        if prefix:
            value = value[1:]
    return rewrite(value, home)


def _walk_objects(obj, seen=None):
    """Visit serializable chart/rule children, including combined charts."""
    seen = set() if seen is None else seen
    if id(obj) in seen or isinstance(obj, (str, bytes, int, float, bool, type(None))):
        return
    seen.add(id(obj))
    if isinstance(obj, (list, tuple)):
        for child in obj:
            yield from _walk_objects(child, seen)
    elif hasattr(obj, "__elements__"):
        yield obj
        fields = set(obj.__elements__) | set(vars(obj))
        for name in fields:
            yield from _walk_objects(getattr(obj, name, None), seen)


def _shift_anchor(anchor, axis, idx, n, delete, shift_point):
    if isinstance(anchor, str):
        row, col = coordinate_to_tuple(anchor)
        if axis == "rows": row = shift_point(row, idx, n, delete) or idx
        else: col = shift_point(col, idx, n, delete) or idx
        return f"{get_column_letter(col)}{row}"
    if isinstance(anchor, AbsoluteAnchor) or getattr(anchor, "editAs", None) == "absolute":
        return anchor
    if not hasattr(anchor, "_from"):
        return anchor
    attr = "row" if axis == "rows" else "col"
    old_start = getattr(anchor._from, attr)
    new_start = (shift_point(old_start + 1, idx, n, delete) or idx) - 1
    setattr(anchor._from, attr, new_start)
    if isinstance(anchor, TwoCellAnchor):
        old_end = getattr(anchor.to, attr)
        if anchor.editAs == "oneCell":
            new_end = old_end + new_start - old_start
        else:
            shifted_end = shift_point(old_end + 1, idx, n, delete)
            new_end = (shifted_end or idx) - 1
            # Keep an object visible when its entire anchoring span is deleted.
            if delete and old_end > old_start and new_end <= new_start:
                new_end = new_start + old_end - old_start
        setattr(anchor.to, attr, new_end)
    return anchor


def _sheet(wb, title):
    return next((ws for ws in wb.worksheets if ws.title.casefold() == title.casefold()), None)


def _cells_for_reference(wb, home, formula, seen=None):
    """Resolve static range/name/table sources; dynamic formulas recalculate in Excel."""
    seen = set() if seen is None else seen
    formula = formula.lstrip("=")
    key = (home, formula)
    if key in seen or formula == "#REF!":
        return None
    seen = seen | {key}
    # Formula tokenization keeps commas inside quoted sheet names intact.
    tokens = Tokenizer("=" + formula).items
    if len(tokens) > 1:
        def evaluate(parts):
            while parts and parts[0].type == "WHITE-SPACE": parts = parts[1:]
            while parts and parts[-1].type == "WHITE-SPACE": parts = parts[:-1]
            depth, groups, start = 0, [], 0
            for index, token in enumerate(parts):
                if token.type == "PAREN": depth += 1 if token.subtype == "OPEN" else -1
                elif depth == 0 and token.value == ",":
                    groups.append(parts[start:index]); start = index+1
            if groups:
                groups.append(parts[start:])
                values = [evaluate(group) for group in groups]
                return None if any(value is None for value in values) else [row for value in values for row in value]
            if parts and parts[0].type == "PAREN" and parts[0].subtype == "OPEN":
                depth = 0
                for index, token in enumerate(parts):
                    if token.type == "PAREN": depth += 1 if token.subtype == "OPEN" else -1
                    if depth == 0: break
                if index == len(parts)-1: return evaluate(parts[1:-1])
            depth, start, groups = 0, 0, []
            for index, token in enumerate(parts):
                if token.type == "PAREN": depth += 1 if token.subtype == "OPEN" else -1
                elif depth == 0 and token.type == "WHITE-SPACE":
                    groups.append(parts[start:index]); start = index+1
            if groups:
                groups.append(parts[start:])
                values = [evaluate(group) for group in groups if group]
                if not values or any(value is None for value in values): return None
                result = values[0]
                for value in values[1:]:
                    wanted = {(cell.parent.title,cell.coordinate) for row in value for cell in row}
                    result = [[cell for cell in row if (cell.parent.title,cell.coordinate) in wanted] for row in result]
                    result = [row for row in result if row]
                return result
            if len(parts) == 1 and parts[0].subtype == "RANGE":
                return _cells_for_reference(wb, home, parts[0].value, seen)
            return None
        return evaluate(tokens)
    qualified = re.fullmatch(r"('(?:[^']|'')+'|[^!\[\]]+)!(.*)", formula)
    qualifier, coord = qualified.groups() if qualified else (home, formula)
    title = qualifier[1:-1].replace("''", "'") if qualifier.startswith("'") else qualifier
    ws = _sheet(wb, title)
    if ws is None: return None  # External workbook or unsupported 3-D source.
    name = ws.defined_names.get(coord) or wb.defined_names.get(coord)
    if name and name.attr_text:
        return _cells_for_reference(wb, ws.title, name.attr_text, seen)
    try:
        left, top, right, bottom = range_boundaries(coord)
    except ValueError:
        left = top = right = bottom = None
    if any(value is not None for value in (left, top, right, bottom)):
        left, right = left or 1, right or ws.max_column
        top, bottom = top or 1, bottom or ws.max_row
        left, right = sorted((left, right)); top, bottom = sorted((top, bottom))
        # Whole-column/row sources end at the worksheet's used cells, avoiding
        # allocation of a million empty cache entries and matching visible data.
        if range_boundaries(coord)[1] is None: bottom = min(bottom, ws.max_row)
        if range_boundaries(coord)[0] is None: right = min(right, ws.max_column)
        return [[ws.cell(row, col) for col in range(left, right+1)] for row in range(top, bottom+1)]
    match = re.fullmatch(r"([A-Za-z_\\][\w.\\]*)\[(.*)\]", coord)
    if match:
        table_name, selector = match.groups()
        for sheet in wb.worksheets:
            table = next((t for t in sheet.tables.values() if t.name.casefold() == table_name.casefold()), None)
            if table is None: continue
            left, top, right, bottom = range_boundaries(table.ref)
            header, totals = int(bool(table.headerRowCount)), int(bool(table.totalsRowCount))
            leaves, stack, pos = [], [], 0
            expression = "[" + selector + "]"
            while pos < len(expression):
                char = expression[pos]
                if char == "'" and pos+1 < len(expression) and expression[pos+1] in "[]#@'":
                    pos += 2; continue
                if char == "[":
                    if stack: stack[-1][1] = True
                    stack.append([pos+1, False])
                elif char == "]" and stack:
                    first, nested = stack.pop()
                    if not nested: leaves.append((first, pos, expression[first:pos]))
                pos += 1
            selectors = {raw for _,_,raw in leaves if raw.startswith("#")}
            if any(raw.startswith("@") for _,_,raw in leaves): return None
            rows = []
            if "#All" in selectors:
                rows = list(range(top,bottom+1))
            else:
                if "#Headers" in selectors: rows.extend(range(top,top+header))
                if not selectors or "#Data" in selectors: rows.extend(range(top+header,bottom-totals+1))
                if "#Totals" in selectors: rows.extend(range(bottom-totals+1,bottom+1))
            columns = {column.name.casefold(): left+i for i, column in enumerate(table.tableColumns)}
            requested = [(first,last,re.sub(r"'([\[\]#@'])", r"\1", raw))
                         for first,last,raw in leaves if not raw.startswith("#")]
            if any(label.casefold() not in columns for _,_,label in requested): return None
            selected = []
            for index,(first,last,label) in enumerate(requested):
                position = columns[label.casefold()]
                if index and re.fullmatch(r"\]\s*:\s*\[", expression[requested[index-1][1]:first]):
                    previous = columns[requested[index-1][2].casefold()]
                    selected.extend(range(previous+1,position+1) if previous <= position else range(previous-1,position-1,-1))
                else: selected.append(position)
            if not requested: selected = list(range(left,right+1))
            return [[sheet.cell(row,col) for col in selected] for row in rows]
    return None


def _numeric(value, epoch):
    if isinstance(value, (datetime, date)): return to_excel(value, epoch)
    if isinstance(value, time): return time_to_days(value)
    if isinstance(value, timedelta): return timedelta_to_days(value)
    if isinstance(value, Number): return float(value)
    return None


def _refresh_cache(ref, cells, wb):
    attr = "numCache" if isinstance(ref, NumRef) else "strCache" if isinstance(ref, StrRef) else "multiLvlStrCache"
    previous = getattr(ref, attr)
    flat = [cell for row in cells for cell in row] if cells is not None else []
    if cells is None or any(cell.data_type == "f" for cell in flat):
        setattr(ref, attr, None)
        return
    if isinstance(ref, NumRef):
        points = []
        for index, cell in enumerate(flat):
            if cell.value is None: continue
            value = _numeric(cell.value, wb.epoch)
            if value is None:
                # Text categories sometimes arrive in a numRef; the reference
                # remains authoritative and Excel supplies the proper cache.
                ref.numCache = None
                return
            points.append(NumVal(idx=index, v=value))
        ref.numCache = NumData(formatCode=previous.formatCode if previous else None,
                               ptCount=len(flat), pt=points)
    elif isinstance(ref, StrRef):
        ref.strCache = StrData(ptCount=len(flat), pt=[StrVal(idx=i, v=str(c.value))
                               for i,c in enumerate(flat) if c.value is not None])
    else:
        columns = list(zip(*cells)) if cells else []
        levels = [Level(pt=[StrVal(idx=i, v=str(c.value)) for i,c in enumerate(column)
                            if c.value is not None]) for column in reversed(columns)]
        ref.multiLvlStrCache = MultiLevelStrData(ptCount=len(cells), lvl=levels)


def apply_dependents(wb, snapshots, target_ws, axis, idx, n, delete,
                     rewrite, shift_point, shift_range):
    """Rewrite dependent data after the structural edit; return changed surfaces."""
    report = {name: [] for name in ("validations", "conditional_formats", "charts", "anchors", "local_defined_names", "print_settings")}
    for snapshot in snapshots:
        ws = snapshot["sheet"]
        for name, dn, text in snapshot["names"]:
            if not text: continue
            updated = rewrite(text, ws.title)
            dn.attr_text = updated
            if updated != text: report["local_defined_names"].append(dict(sheet=ws.title, name=name, old=text, new=updated))
    for snapshot in snapshots:
        ws = snapshot["sheet"]
        edited = ws is target_ws
        validations = []
        formatting = ConditionalFormattingList()
        for kind, entries in (("validations", snapshot["validations"]),
                              ("conditional_formats", snapshot["formatting"])):
            for entry in entries:
                obj, rules = (entry, []) if kind == "validations" else entry
                old = str(obj.sqref)
                if not old: continue
                origin = _range_origin(obj.sqref)
                new, survivor = _shift_applied(obj.sqref, axis, idx, n, delete, shift_range) if edited else (old, origin)
                if new is None:
                    report[kind].append(dict(sheet=ws.title, old=old, new=None))
                    continue
                obj.sqref = new
                changed = old != new
                if kind == "validations":
                    for attr in ("formula1", "formula2"):
                        before = getattr(obj, attr)
                        after = _formula(before, ws.title, rewrite, origin, survivor)
                        setattr(obj, attr, after)
                        changed |= before != after
                    validations.append(obj)
                else:
                    for rule in rules:
                        before = list(rule.formula)
                        rule.formula = [_formula(f, ws.title, rewrite, origin, survivor) for f in rule.formula]
                        changed |= before != rule.formula
                        for node in _walk_objects(rule):
                            if getattr(node, "type", None) == "formula" and hasattr(node, "val"):
                                before = node.val
                                node.val = _formula(before, ws.title, rewrite, origin, survivor)
                                changed |= before != node.val
                        formatting.add(obj, rule)
                if changed: report[kind].append(dict(sheet=ws.title, old=old, new=str(obj.sqref)))
        ws.data_validations.dataValidation = validations
        ws.conditional_formatting = formatting
        for item, anchor in snapshot["drawings"]:
            if edited:
                item.anchor = _shift_anchor(deepcopy(anchor), axis, idx, n, delete, shift_point)
                if item.anchor != anchor:
                    report["anchors"].append(dict(sheet=ws.title, object=type(item).__name__))
            for ref in _walk_objects(item):
                if isinstance(ref, AxDataSource) and ref.numRef is not None and ref.numRef.f:
                    cells = _cells_for_reference(wb, ws.title, rewrite(ref.numRef.f, ws.title))
                    flat = [cell for row in cells for cell in row] if cells is not None else []
                    if (any(isinstance(cell.value, (str, CellRichText)) for cell in flat)
                            and not any(cell.data_type in {"f", "e"} for cell in flat)):
                        # openpyxl creates numRef even for literal text labels;
                        # use the corresponding string source to cache them.
                        ref.strRef = StrRef(f=ref.numRef.f)
                        ref.numRef = None
                if not isinstance(ref, (NumRef, StrRef, MultiLevelStrRef)) or not ref.f: continue
                old = ref.f
                ref.f = rewrite(ref.f, ws.title)
                _refresh_cache(ref, _cells_for_reference(wb, ws.title, ref.f), wb)
                if ref.f != old: report["charts"].append(dict(sheet=ws.title, old=old, new=ref.f))
        if edited:
            areas = [shift_range(area, axis, idx, n, delete) for area in snapshot["print_area"]]
            areas = [area for area in areas if area]
            ws.print_area = areas
            for key, prop in (("print_rows", "print_title_rows"), ("print_cols", "print_title_cols")):
                old = snapshot[key]
                if old:
                    updated = rewrite(old, ws.title)
                    setattr(ws, prop, None if "#REF!" in updated else updated)
                    if updated != old: report["print_settings"].append(dict(sheet=ws.title, property=prop, old=old, new=getattr(ws, prop)))
            if areas != snapshot["print_area"]:
                report["print_settings"].append(dict(sheet=ws.title, property="print_area", old=snapshot["print_area"], new=areas))
    if wb.calculation is None: wb.calculation = CalcProperties()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    return report
