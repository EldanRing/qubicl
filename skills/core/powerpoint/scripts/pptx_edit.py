#!/usr/bin/env python3
"""Edit a .pptx in place (or save to --output).

Operations (repeatable / combinable):
  --replace-text OLD NEW   Replace text everywhere (slides, tables, notes).
                           Matches may span formatted runs and grouped
                           shapes. Unchanged text keeps its formatting;
                           replacement text inherits its first matched run.
  --chart-data SPEC.json   Update a chart. Full replace spec:
                           {"slide": 0, "chart": 0,
                            "categories": ["Q1", "Q2"],
                            "series": {"North": [1, 2], "South": [3, 4]}}
                           Or surgical ops (existing data is read, modified,
                           and written back via replace_data):
                           {"slide": 0, "chart": 0, "ops": [
                             {"op": "update_series", "name": "North",
                              "values": [5, 6]},
                             {"op": "add_series", "name": "East",
                              "values": [1, 2]},
                             {"op": "remove_series", "name": "South"},
                             {"op": "rename_category", "from": "Q1",
                              "to": "Q1 FY26"},
                             {"op": "set_title", "title": "New title"}]}
                           Optional "plot": N selects a combined-chart plot
                           (default 0). XY/bubble full replacement uses
                           {"series": [{"name": "S", "xValues": [1, 2],
                             "yValues": [3, 4], "bubbleSizes": [5, 6]}]}.
                           Omit bubbleSizes for scatter plots. update_series
                           can change xValues, yValues/values or bubbleSizes
                           independently; unspecified dimensions are retained.
  --swap-image SLIDE SHAPE_NAME NEW_IMAGE
                           Replace a picture's bits, keeping position/size.
  --remove-slide N         Delete slide at index N (0-based).
  --move-slide FROM TO     Reorder: move slide FROM to position TO.
  --duplicate-slide N      Append an independent deep copy of slide N
                           (text, images, tables, grouped shapes, charts,
                           embedded workbooks, notes and slide properties).
                           Layout/master definitions and external links retain
                           their presentation-wide meaning.
  --set-background N HEX   Solid background color for slide N.
  --hyperlink N TEXT URL   Make runs containing TEXT on slide N links.
  --enable-slide-number N  Copy the layout's slide-number placeholder in.
  --set-footer N TEXT      Enable the layout's footer placeholder with TEXT.
  --set-notes N TEXT       Replace slide N's speaker notes.
  --append-notes N TEXT    Append a paragraph to slide N's speaker notes.
"""
import argparse
import copy
import json
import io
import re
from pathlib import PurePosixPath, Path
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData, BubbleChartData
from pptx.chart.xmlwriter import SeriesXmlRewriterFactory
from pptx.oxml import parse_xml
from openpyxl import load_workbook, Workbook
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT, RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.package import _Relationship
from pptx.opc.packuri import PackURI

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx_chart_data import plot_kind, plot_type, plot_record

R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _run_format_key(r):
    """Canonical string for a run's <a:rPr>; None when absent."""
    rPr = r.find(qn("a:rPr"))
    if rPr is None:
        return None
    return etree.tostring(rPr)


def normalize_runs(para):
    """Merge adjacent runs whose formatting is byte-identical.

    PowerPoint splits paragraph text into runs at spell-check and edit
    boundaries even when formatting never changes; merging them back makes
    cross-run text replacement lossless for the common case.
    """
    runs = list(para.runs)
    i = 0
    while i + 1 < len(runs):
        a, b = runs[i], runs[i + 1]
        if (_run_format_key(a._r) == _run_format_key(b._r)
                and a._r.getnext() is b._r):
            a.text = a.text + b.text
            b._r.getparent().remove(b._r)
            runs.pop(i + 1)
        else:
            i += 1


def replace_in_text_frame(text_frame, old, new):
    import re
    if not old: return 0
    count = 0
    for para in text_frame.paragraphs:
        runs = list(para.runs)
        joined = "".join(run.text for run in runs)
        spans, pos = [], 0
        for run in runs:
            spans.append((run, pos, pos + len(run.text)))
            pos += len(run.text)
        matches = list(re.finditer(re.escape(old), joined))
        for match in reversed(matches):
            first = True
            for run, start, end in spans:
                if end <= match.start() or start >= match.end(): continue
                a, b = max(start, match.start()) - start, min(end, match.end()) - start
                run.text = run.text[:a] + (new if first else "") + run.text[b:]
                first = False
        count += len(matches)
    return count


def iter_shapes(shapes):
    for shape in shapes:
        yield shape
        if hasattr(shape, "shapes"):
            yield from iter_shapes(shape.shapes)


def iter_text_frames(slide):
    for shape in iter_shapes(slide.shapes):
        if shape.has_text_frame:
            yield shape.text_frame
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame
    if slide.has_notes_slide:
        yield slide.notes_slide.notes_text_frame


def replace_text(prs, old, new):
    total = 0
    for slide in prs.slides:
        for tf in iter_text_frames(slide):
            total += replace_in_text_frame(tf, old, new)
    return total


def _series_spec(value):
    if isinstance(value, dict):
        return [dict(data, name=name) if isinstance(data, dict) else {"name": name, "values": data}
                for name, data in value.items()]
    return [dict(series) for series in value]


def _chart_data(kind, categories, series):
    data = BubbleChartData() if kind == "bubble" else XyChartData() if kind == "xy" else CategoryChartData()
    if kind == "category":
        data.categories = categories
        for entry in series:
            data.add_series(entry["name"], entry["values"])
        return data
    for entry in series:
        xs, ys = entry["xValues"], entry.get("yValues", entry.get("values"))
        sizes = entry.get("bubbleSizes")
        if ys is None or len(xs) != len(ys) or (kind == "bubble" and (sizes is None or len(xs) != len(sizes))):
            raise ValueError("Each XY/bubble series needs equal-length xValues, yValues and (for bubbles) bubbleSizes")
        target = data.add_series(entry["name"])
        for index, (x, y) in enumerate(zip(xs, ys)):
            if kind == "bubble": target.add_data_point(x, y, sizes[index])
            else: target.add_data_point(x, y)
    return data


def _detach_external_workbook(chart):
    external = chart._chartSpace.find(qn("c:externalData"))
    if external is None: return
    rid = external.get(R_EMBED + "id")
    if rid and chart.part.rels[rid].is_external:
        chart._chartSpace.remove(external)
        chart.part.drop_rel(rid)


def _replace_plot_data(chart, plot_index, data):
    """Replace one plot's data without applying its type to the other plots."""
    plot = chart.plots[plot_index]
    isolated = copy.deepcopy(chart._chartSpace)
    for index, element in reversed(list(enumerate(isolated.plotArea.iter_xCharts()))):
        if index != plot_index: isolated.plotArea.remove(element)
    target = next(isolated.plotArea.iter_xCharts())
    if len(data):
        if not target.sers:
            seed = parse_xml(data.xml_bytes(plot_type(plot))).plotArea.sers[0]
            target._insert_ser(copy.deepcopy(seed))
        SeriesXmlRewriterFactory(plot_type(plot), data).replace_series_data(isolated)
        generated = load_workbook(io.BytesIO(data.xlsx_blob))
    else:
        for series in list(target.sers): target.remove(series)
        generated = Workbook()
    if len(chart.plots) == 1:
        workbook = generated
    else:
        existing = chart.part.chart_workbook.xlsx_part
        workbook = load_workbook(io.BytesIO(existing.blob)) if existing is not None else Workbook()
        source = generated.active
        # Keep existing worksheets (including formulas feeding untouched plots)
        # intact. A private new sheet owns this plot's replacement data.
        destination = workbook.create_sheet(f"QubiclPlot{plot_index + 1}")
        for row in source:
            for cell in row:
                copied = destination.cell(cell.row, cell.column, cell.value)
                copied._style = copy.copy(cell._style)
        for formula in target.xpath('.//c:f'):
            if formula.text and "!" in formula.text:
                formula.text = "'" + destination.title.replace("'", "''") + "'!" + formula.text.rsplit("!", 1)[1]
    stream = io.BytesIO(); workbook.save(stream)
    original = plot._element
    original.getparent().replace(original, target)
    # A new series in an earlier plot can otherwise reuse a later plot's idx.
    other_indices = {series.idx.val for other in chart._chartSpace.plotArea.iter_xCharts()
                     if other is not target for series in other.sers}
    used = set(other_indices)
    next_index = max((series.idx.val for series in chart._chartSpace.plotArea.sers), default=-1) + 1
    for series in target.sers:
        if series.idx.val in used:
            series.idx.val = next_index
            next_index += 1
        used.add(series.idx.val)
    for order, series in enumerate(chart._chartSpace.plotArea.sers): series.order.val = order
    chart.part.chart_workbook.update_from_xlsx_blob(stream.getvalue())


def update_chart(prs, spec_path):
    """Edit category, scatter or bubble data in the selected plot (default 0)."""
    with open(spec_path, encoding="utf-8") as fh:
        spec = json.load(fh)
    slide = prs.slides[spec.get("slide", 0)]
    charts = [shape.chart for shape in iter_shapes(slide.shapes) if shape.has_chart]
    if not charts: raise ValueError(f"no chart on slide {spec.get('slide', 0)}")
    chart = charts[spec.get("chart", 0)]
    plot_index = spec.get("plot", 0)
    plot = chart.plots[plot_index]
    record = plot_record(plot, plot_index)
    kind, categories, series = record["kind"], record["categories"], record["series"]
    dirty = "ops" not in spec
    if "ops" in spec:
        for op in spec["ops"]:
            operation = op["op"]
            if operation in {"update_series", "remove_series", "rename_series"}:
                matches = [entry for entry in series if entry["name"] == op["name"]]
                if not matches: raise ValueError(f"no series named {op['name']!r}")
                entry = matches[0]
                if operation == "update_series":
                    for key in ("values", "xValues", "yValues", "bubbleSizes"):
                        if key in op: entry[key] = op[key]
                    if "values" in op and "yValues" not in op and kind != "category": entry["yValues"] = op["values"]
                elif operation == "remove_series": series = [entry for entry in series if entry["name"] != op["name"]]
                else: entry["name"] = op["to"]
                dirty = True
            elif operation == "add_series":
                series.append({key: value for key, value in op.items() if key != "op"})
                dirty = True
            elif operation == "rename_category":
                index = int(op["index"]) if "index" in op else categories.index(op["from"])
                categories[index] = op["to"]
                dirty = True
            elif operation == "set_title":
                chart.has_title = True
                chart.chart_title.text_frame.text = op["title"]
            else: raise ValueError(f"unknown chart op {operation!r}")
    else:
        categories = spec.get("categories", categories)
        series = _series_spec(spec["series"])
    if not dirty: return
    data = _chart_data(kind, categories, series)
    # An external workbook cannot be edited in place. Materialize every cached
    # plot before detaching it so mixed charts keep their other visible data.
    external = chart._chartSpace.find(qn("c:externalData"))
    rid = external.get(R_EMBED + "id") if external is not None else None
    linked = bool(rid and chart.part.rels[rid].is_external)
    if linked and len(chart.plots) > 1:
        preserved = [plot_record(other, index) for index, other in enumerate(chart.plots)]
        _detach_external_workbook(chart)
        for original in preserved:
            if original["index"] == plot_index: continue
            unchanged = _chart_data(original["kind"], original["categories"], original["series"])
            _replace_plot_data(chart, original["index"], unchanged)
    else:
        _detach_external_workbook(chart)
    _replace_plot_data(chart, plot_index, data)


def replace_chart_data(chart, data):
    """Supplied data becomes a private embedded workbook, even for linked charts."""
    _detach_external_workbook(chart)
    chart.replace_data(data)


def swap_image(prs, slide_idx, shape_name, new_path):
    slide = prs.slides[int(slide_idx)]
    for shape in iter_shapes(slide.shapes):
        if (shape.shape_type == MSO_SHAPE_TYPE.PICTURE
                and shape.name == shape_name):
            image_part, rid = slide.part.get_or_add_image_part(new_path)
            blip = shape._element.blipFill.blip
            blip.set(R_EMBED + "embed", rid)
            return True
    raise SystemExit(f"no picture named {shape_name!r} on slide {slide_idx}")


def remove_slide(prs, index):
    sldIdLst = prs.slides._sldIdLst
    slide_id = list(sldIdLst)[int(index)]
    rid = slide_id.get(R_EMBED + "id")
    prs.part.drop_rel(rid)
    sldIdLst.remove(slide_id)


def move_slide(prs, src, dst):
    """Reorder by moving the <p:sldId> element inside <p:sldIdLst>."""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    element = ids[int(src)]
    sldIdLst.remove(element)
    sldIdLst.insert(int(dst), element)


def duplicate_slide(prs, index):
    """Append a slide with an independent copy of its owned OPC part graph.

    Chart XML, workbooks, chart styles/colors, notes, media and other owned
    parts are copied together. Relationship IDs remain unchanged inside each
    copied part, including generic XML parts not modeled by python-pptx.
    Layout/master definitions and links to other slides remain presentation
    references; self-links and notes' slide backlinks target the duplicate.
    """
    source = prs.slides[int(index)].part
    package = source.package
    used_names = {str(part.partname) for part in package.iter_parts()}
    copied = {}

    def next_name(original):
        path = PurePosixPath(str(original))
        match = re.fullmatch(r"(.*?)([0-9]+)", path.stem)
        stem = match.group(1) if match else path.stem + "_copy"
        number = 1
        while True:
            candidate = str(path.with_name(f"{stem}{number}{path.suffix}"))
            if candidate not in used_names:
                used_names.add(candidate)
                return PackURI(candidate)
            number += 1

    shared_relationships = {RT.SLIDE_LAYOUT, RT.SLIDE_MASTER, RT.NOTES_MASTER, RT.SLIDE}

    def clone(part):
        if part in copied:
            return copied[part]
        result = type(part).load(next_name(part.partname), part.content_type,
                                 package, part.blob)
        copied[part] = result  # Seed before recursion: notes and custom parts may cycle.
        for rel in part.rels.values():
            if rel.is_external:
                target, mode = rel.target_ref, RTM.EXTERNAL
            else:
                original = rel.target_part
                target = (copied[original] if original in copied else original
                          if rel.reltype in shared_relationships else clone(original))
                mode = RTM.INTERNAL
            # Preserve IDs instead of rewriting XML. Binary/generic parts can
            # own relationships too, and must retain their original references.
            result.rels._rels[rel.rId] = _Relationship(
                result.partname.baseURI, rel.rId, rel.reltype, mode, target)
        return result

    duplicate = clone(source)
    rid = prs.part.relate_to(duplicate, RT.SLIDE)
    prs.slides._sldIdLst.add_sldId(rid)
    return len(prs.slides) - 1


def set_background(slide, hex_color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor.from_string(hex_color)


def add_hyperlink(prs, slide_idx, text, url):
    """Turn every run containing `text` on the slide into a hyperlink.

    The link applies to the whole run (python-pptx links whole runs).
    """
    slide = prs.slides[int(slide_idx)]
    hits = 0
    for shape in iter_shapes(slide.shapes):
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                if text in run.text:
                    run.hyperlink.address = url
                    hits += 1
    if not hits:
        raise SystemExit(f"no run containing {text!r} on slide {slide_idx}")
    return hits


def _copy_layout_placeholder(slide, ph_idx):
    """Copy the layout placeholder with idx `ph_idx` onto the slide.

    Slide-number (idx 12) and footer (idx 11) placeholders exist on the
    layout but are not inherited by a slide until the slide carries its
    own copy — this enables them. Returns the new shape, or None when the
    layout does not provide that placeholder.
    """
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.idx == ph_idx:
            el = copy.deepcopy(ph._element)
            el.xpath(".//p:cNvPr")[0].set("id", str(slide.shapes._next_shape_id))
            slide.shapes._spTree.append(el)
            for shape in slide.placeholders:
                if shape.placeholder_format.idx == ph_idx:
                    return shape
            return None
    return None


def enable_slide_number(slide):
    if any(ph.placeholder_format.idx == 12 for ph in slide.placeholders):
        return True
    return _copy_layout_placeholder(slide, 12) is not None


def set_footer(slide, text):
    shape = next((ph for ph in slide.placeholders
                  if ph.placeholder_format.idx == 11), None)
    if shape is None:
        shape = _copy_layout_placeholder(slide, 11)
    if shape is None:
        raise SystemExit("layout provides no footer placeholder; add a "
                         "textbox instead")
    shape.text_frame.text = text
    return True


def set_notes(slide, text, append=False):
    tf = slide.notes_slide.notes_text_frame
    if append and tf.text:
        para = tf.add_paragraph()
        para.text = text
    else:
        tf.text = text


def main(argv=None):
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(
        description="Edit a .pptx: replace text, update chart data, swap "
                    "images, duplicate/remove/reorder slides, backgrounds, "
                    "hyperlinks, footers, slide numbers, speaker notes.",
        epilog=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("pptx", help="path to the .pptx file")
    parser.add_argument("--output", help="save to this path instead of "
                                         "overwriting the input")
    parser.add_argument("--replace-text", nargs=2, action="append",
                        metavar=("OLD", "NEW"), default=[])
    parser.add_argument("--chart-data", metavar="SPEC_JSON")
    parser.add_argument("--swap-image", nargs=3,
                        metavar=("SLIDE", "SHAPE_NAME", "IMAGE"))
    parser.add_argument("--remove-slide", type=int, metavar="N")
    parser.add_argument("--move-slide", nargs=2, type=int,
                        metavar=("FROM", "TO"))
    parser.add_argument("--duplicate-slide", type=int, metavar="N")
    parser.add_argument("--set-background", nargs=2,
                        metavar=("SLIDE", "HEX"))
    parser.add_argument("--hyperlink", nargs=3,
                        metavar=("SLIDE", "TEXT", "URL"))
    parser.add_argument("--enable-slide-number", type=int, metavar="N")
    parser.add_argument("--set-footer", nargs=2, metavar=("SLIDE", "TEXT"))
    parser.add_argument("--set-notes", nargs=2, metavar=("SLIDE", "TEXT"))
    parser.add_argument("--append-notes", nargs=2, metavar=("SLIDE", "TEXT"))
    args = parser.parse_args(argv)

    prs = Presentation(args.pptx)
    report = {"ok": True, "replacements": 0}

    for old, new in args.replace_text:
        report["replacements"] += replace_text(prs, old, new)
    if args.chart_data:
        update_chart(prs, args.chart_data)
        report["chart_updated"] = True
    if args.swap_image:
        swap_image(prs, *args.swap_image)
        report["image_swapped"] = True
    if args.duplicate_slide is not None:
        report["duplicated_to"] = duplicate_slide(prs, args.duplicate_slide)
    if args.set_background:
        set_background(prs.slides[int(args.set_background[0])],
                       args.set_background[1])
        report["background_set"] = True
    if args.hyperlink:
        report["hyperlinked_runs"] = add_hyperlink(prs, *args.hyperlink)
    if args.enable_slide_number is not None:
        report["slide_number_enabled"] = enable_slide_number(
            prs.slides[args.enable_slide_number])
    if args.set_footer:
        set_footer(prs.slides[int(args.set_footer[0])], args.set_footer[1])
        report["footer_set"] = True
    if args.set_notes:
        set_notes(prs.slides[int(args.set_notes[0])], args.set_notes[1])
        report["notes_set"] = True
    if args.append_notes:
        set_notes(prs.slides[int(args.append_notes[0])],
                  args.append_notes[1], append=True)
        report["notes_appended"] = True
    if args.remove_slide is not None:
        remove_slide(prs, args.remove_slide)
        report["slide_removed"] = args.remove_slide
    if args.move_slide:
        move_slide(prs, *args.move_slide)
        report["slide_moved"] = args.move_slide

    out = args.output or args.pptx
    prs.save(out)
    report["output"] = out
    print(json.dumps(report))
    return 0


if __name__ == "__main__":
    sys.exit(main())
