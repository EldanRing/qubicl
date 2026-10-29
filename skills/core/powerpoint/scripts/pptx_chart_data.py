"""Complete data inspection for category, XY and bubble PowerPoint plots."""
import io
from datetime import date, datetime, time, timedelta

from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple
from openpyxl.utils.datetime import to_excel, time_to_days, timedelta_to_days
from pptx.chart.plot import PlotTypeInspector
from pptx.oxml.ns import qn


def plot_kind(plot):
    tag = plot._element.tag.rsplit('}', 1)[-1]
    return 'bubble' if tag == 'bubbleChart' else 'xy' if tag == 'scatterChart' else 'category'


def plot_type(plot):
    try:
        return PlotTypeInspector.chart_type(plot)
    except NotImplementedError:
        return plot._element.tag.rsplit('}', 1)[-1]


def _embedded_values(source, chart, numeric=False):
    """Resolve a cacheless source from its embedded XLSX without external I/O."""
    formulas = source.xpath('.//c:f/text()')
    if not formulas or chart is None: return None
    try:
        part = chart.part.chart_workbook.xlsx_part
    except ValueError:  # Externally linked workbooks have no internal part.
        return None
    if part is None: return None
    workbook = load_workbook(io.BytesIO(part.blob), data_only=True)
    formula = formulas[0]
    if formula in workbook.defined_names:
        formula = workbook.defined_names[formula].attr_text
    try:
        sheet_name, (left, top, right, bottom) = range_to_tuple(formula)
    except ValueError:
        return None
    if sheet_name not in workbook: return None
    sheet = workbook[sheet_name]
    left, right = sorted((left or 1, right or sheet.max_column))
    top, bottom = sorted((top or 1, bottom or sheet.max_row))
    values = [sheet.cell(row, col).value for row in range(top, bottom + 1)
              for col in range(left, right + 1)]
    if numeric:
        values = [to_excel(value, workbook.epoch) if isinstance(value, (datetime, date))
                  else time_to_days(value) if isinstance(value, time)
                  else timedelta_to_days(value) if isinstance(value, timedelta)
                  else value for value in values]
    return values


def numeric_values(series, tag, chart=None):
    """Read cached or literal numeric data, preserving sparse point positions."""
    source = series._element.find(qn('c:' + tag))
    if source is None:
        return []
    if not source.xpath('.//c:numCache | .//c:numLit'):
        values = _embedded_values(source, chart, numeric=True)
        if values is not None: return values
    points = source.xpath('.//c:pt')
    counts = source.xpath('.//c:ptCount/@val')
    count = int(counts[0]) if counts else max((int(pt.get('idx')) + 1 for pt in points), default=0)
    values = [None] * count
    for point in points:
        index = int(point.get('idx'))
        value = point.find(qn('c:v'))
        if index >= len(values): values.extend([None] * (index + 1 - len(values)))
        if value is not None and value.text:
            try: values[index] = float(value.text)
            except ValueError: values[index] = value.text
    return values


def series_record(series, kind, chart=None):
    try:
        name = series.name
    except (AttributeError, KeyError):
        name = None
    if not name:
        title = series._element.find(qn('c:tx'))
        if title is not None:
            labels = _embedded_values(title, chart)
            if labels and labels[0] is not None: name = str(labels[0])
    values = numeric_values(series, 'yVal' if kind in {'xy', 'bubble'} else 'val', chart)
    result = {'name': name, 'values': values}
    if kind in {'xy', 'bubble'}:
        result.update(xValues=numeric_values(series, 'xVal', chart), yValues=values)
    if kind == 'bubble':
        result['bubbleSizes'] = numeric_values(series, 'bubbleSize', chart)
    return result


def plot_record(plot, index):
    kind = plot_kind(plot)
    categories = [str(category) for category in plot.categories]
    if not categories and kind == 'category' and len(plot.series):
        source = plot.series[0]._element.find(qn('c:cat'))
        if source is not None:
            values = _embedded_values(source, plot.chart)
            if values is not None: categories = [str(value) if value is not None else '' for value in values]
    return {'index': index, 'type': str(plot_type(plot)), 'kind': kind,
            'categories': categories,
            'series': [series_record(series, kind, plot.chart) for series in plot.series]}
