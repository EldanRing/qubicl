#!/usr/bin/env python3
"""Render captured PDF/DOCX/XLSX/PPTX bytes and report verifiable output identities.

This is a renderer and structural checker, not a claim of visual correctness.
Run with /opt/qubicl/skills-venv/bin/python. LibreOffice is needed for Office files.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import xml.etree.ElementTree as ET
import os
import signal
from pathlib import Path
import shutil
import subprocess
import tempfile
import zipfile


PROFILE = '''<?xml version="1.0"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop><prop oor:name="DisableActiveContent" oor:op="fuse"><value>true</value></prop><prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Writer/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Content/Update"><prop oor:name="Link" oor:op="fuse"><value>2</value></prop></item>
</oor:items>'''


def sha256(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def run_bounded(command, timeout, env=None):
    """A timeout must also stop converter descendants, not just its shell launcher."""
    with subprocess.Popen(command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                          text=True, encoding='utf-8', errors='replace', env=env,
                          start_new_session=True) as process:
        try:
            stdout, stderr = process.communicate(timeout=timeout)
            return subprocess.CompletedProcess(command, process.returncode, stdout, stderr)
        finally:
            try:
                os.killpg(process.pid, signal.SIGKILL)
            except ProcessLookupError:
                pass
            process.communicate()


def convert_office(source, directory, format_name, timeout=180):
    """Convert one captured file with an isolated profile and no user macro trust."""
    executable = shutil.which('soffice') or shutil.which('libreoffice')
    if not executable:
        raise RuntimeError('LibreOffice is unavailable; Office rendering/recalculation requires workstation')
    directory = Path(directory).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='qubicl-office-profile-') as temporary:
        profile = Path(temporary) / 'profile'
        (profile / 'user').mkdir(parents=True)
        (profile / 'user' / 'registrymodifications.xcu').write_text(PROFILE)
        result = run_bounded([
            executable, '-env:UserInstallation=' + profile.as_uri(), '--headless',
            '--nologo', '--nodefault', '--norestore', '--convert-to', format_name,
            '--outdir', str(directory), str(Path(source).resolve()),
        ], timeout=timeout,
            env={'PATH': '/usr/bin:/bin', 'LANG': 'C.UTF-8', 'LC_ALL': 'C.UTF-8', 'SAL_USE_VCLPLUGIN': 'svp'})
    extension = format_name.split(':')[0]
    output = directory / (Path(source).stem + '.' + extension)
    if result.returncode or not output.is_file() or output.stat().st_size == 0:
        raise RuntimeError('LibreOffice conversion failed: ' + result.stderr[-1000:] + result.stdout[-1000:])
    return output


def cached_empty_formulas(path):
    """openpyxl maps a valid cached empty string to None; inspect its explicit type."""
    main = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    rel_id = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    empty = set()
    with zipfile.ZipFile(path) as archive:
        relationships = ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))
        targets = {item.get('Id'): item.get('Target') for item in relationships if item.get('TargetMode') != 'External'}
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        for sheet in workbook.findall(main + 'sheets/' + main + 'sheet'):
            target = targets.get(sheet.get(rel_id))
            if not target:
                continue
            name = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join('xl', target))
            xml = ET.fromstring(archive.read(name))
            for cell in xml.iter(main + 'c'):
                value = cell.find(main + 'v')
                if cell.get('t') == 'str' and cell.find(main + 'f') is not None and value is not None and value.text in (None, ''):
                    empty.add((sheet.get('name'), cell.get('r')))
    return empty


def workbook_validation(path):
    """Inspect cached results, including errors; never equate cached with correct."""
    from openpyxl import load_workbook
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    empty = cached_empty_formulas(path)
    count = cached = 0
    missing, errors = [], []
    for sheet in formulas:
        for row in sheet:
            for cell in row:
                value = values[sheet.title][cell.coordinate]
                if cell.data_type == 'f':
                    count += 1
                    if value.value is None and (sheet.title, cell.coordinate) not in empty:
                        missing.append({'sheet': sheet.title, 'cell': cell.coordinate})
                    else:
                        cached += 1
                if value.data_type == 'e':
                    errors.append({'sheet': sheet.title, 'cell': cell.coordinate, 'error': value.value})
    formulas.close()
    values.close()
    return {'formulaCells': count, 'withCachedValues': cached,
            'missingFormulaResults': missing, 'formulaErrors': errors,
            'noFormulaErrors': not errors, 'allFormulaResultsAvailable': not missing,
            'formulaSemanticsVerified': False}


def workbook_changes(before, after):
    """Report format/package differences, without asserting equivalent Excel semantics."""
    with zipfile.ZipFile(before) as old, zipfile.ZipFile(after) as new:
        names = sorted(set(old.namelist()) | set(new.namelist()))
        changed = [name for name in names if name.startswith('xl/') and
                   (name not in old.namelist() or name not in new.namelist() or old.read(name) != new.read(name))]
    return {'changedPackageParts': changed, 'formatChangesPossible': bool(changed),
            'fidelityVerified': False}


def render(source, out_dir, prefix='page', dpi=100, timeout=180):
    source = Path(source).resolve()
    extension = source.suffix.lower()
    if extension not in {'.pdf', '.docx', '.xlsx', '.pptx'}:
        raise ValueError('Supported artifact types: PDF, DOCX, XLSX and PPTX')
    if not prefix or Path(prefix).name != prefix or prefix in ('.', '..'):
        raise ValueError('prefix must be a filename component')
    if dpi < 1 or timeout < 1:
        raise ValueError('dpi and timeout must be positive')
    splitter = shutil.which('pdftoppm') or shutil.which('pdftocairo')
    missing = [] if splitter else ['pdftoppm or pdftocairo']
    if extension != '.pdf' and not (shutil.which('soffice') or shutil.which('libreoffice')):
        missing.append('LibreOffice')
    if missing:
        return {'rendered': False, 'missing': missing, 'files': [],
                'warnings': ['No rendered output was produced; use workstation for Office rendering.'],
                'validation': {'visuallyInspected': False}}
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    run_dir = Path(tempfile.mkdtemp(prefix='render-', dir=Path(out_dir).resolve()))
    warnings = ['Inspect generated pages for layout and content; conversion success is not visual verification.']
    validation = {'visuallyInspected': False, 'fidelityVerified': False}
    try:
        with tempfile.TemporaryDirectory(prefix='qubicl-artifact-') as temporary:
            captured = Path(temporary) / ('source' + extension)
            # Hash the same immutable bytes that the converters consume.
            shutil.copyfile(source, captured)
            source_hash = sha256(captured)
            if extension == '.xlsx':
                recalculated = convert_office(captured, Path(temporary) / 'recalculated', 'xlsx:Calc MS Excel 2007 XML', timeout)
                validation.update({'recalculated': True, **workbook_validation(recalculated),
                                   **workbook_changes(captured, recalculated)})
                warnings.append('LibreOffice may change workbook formatting, formulas and Excel-specific features; changedPackageParts identifies changed OOXML. Volatile formulas can change on a subsequent open.')
                pdf = convert_office(recalculated, Path(temporary) / 'pdf', 'pdf', timeout)
            elif extension == '.pdf':
                pdf = captured
            else:
                pdf = convert_office(captured, Path(temporary) / 'pdf', 'pdf', timeout)
                warnings.append('LibreOffice rendering may differ from Microsoft Office; review fonts, pagination and object layout.')
            from pypdf import PdfReader
            page_count = len(PdfReader(pdf).pages)
            if page_count < 1:
                raise RuntimeError('Conversion produced an empty PDF')
            retained_pdf = run_dir / 'document.pdf'
            shutil.copyfile(pdf, retained_pdf)
            output_prefix = run_dir / prefix
            result = run_bounded([splitter, '-png', '-r', str(dpi), str(retained_pdf), str(output_prefix)], timeout)
            if result.returncode:
                raise RuntimeError('PDF rasterization failed: ' + result.stderr[-1000:])
        files = sorted(run_dir.glob(prefix + '-*.png'), key=lambda path: int(path.stem.rsplit('-', 1)[1]))
        if len(files) != page_count or any(path.stat().st_size == 0 for path in files):
            raise RuntimeError('Rasterizer output does not match the PDF page count')
        validation['generatedPages'] = page_count
        result = {'schemaVersion': 1, 'rendered': True, 'source': str(source), 'sourceSha256': source_hash,
                  'files': [str(path) for path in files],
                  'pages': [{'page': i + 1, 'path': str(path), 'sha256': sha256(path)} for i, path in enumerate(files)],
                  'pdf': {'path': str(retained_pdf), 'sha256': sha256(retained_pdf)},
                  'warnings': warnings, 'validation': validation, 'manifest': str(run_dir / 'manifest.json')}
        (run_dir / 'manifest.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
        return result
    except BaseException:
        # This unique directory belongs only to this failed invocation.
        shutil.rmtree(run_dir)
        raise


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source')
    parser.add_argument('--outdir', default='render')
    parser.add_argument('--prefix', default='page')
    parser.add_argument('--dpi', type=int, default=100)
    parser.add_argument('--timeout', type=int, default=180)
    args = parser.parse_args(argv)
    print(json.dumps(render(args.source, args.outdir, args.prefix, args.dpi, args.timeout), ensure_ascii=False))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
