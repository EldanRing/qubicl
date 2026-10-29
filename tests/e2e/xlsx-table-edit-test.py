"""Real openpyxl round trips for supported table schema transformations."""
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from copy import copy

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table, TableFormula
from openpyxl.worksheet.filters import AutoFilter, SortState, SortCondition
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / 'skills/core/xlsx/scripts/xlsx_restructure.py'
spec = importlib.util.spec_from_file_location('restructure', SCRIPT)
rewrite = importlib.util.module_from_spec(spec); sys.modules[spec.name] = rewrite; spec.loader.exec_module(rewrite)

class TableEditing(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup); self.root = Path(self.temp.name)
    def table(self, names, totals=False):
        wb = Workbook(); ws = wb.active; ws.title = 'Data'
        ws.append(names); ws.append(list(range(1,len(names)+1))); ws.append(list(range(10,10+len(names))))
        if totals: ws.append(['Totals'] + [0]*(len(names)-1))
        table = Table(displayName='Sales', ref=ws.dimensions, totalsRowCount=int(totals))
        table._initialise_columns()
        for column, name in zip(table.tableColumns,names): column.name=name
        table.autoFilter = AutoFilter(ref=f'A1:{rewrite.get_column_letter(len(names))}3')
        ws.add_table(table)
        return wb, ws, table
    def run_edit(self, wb, op, arg):
        source, dest = self.root/'source.xlsx', self.root/'output.xlsx'; wb.save(source)
        original = source.read_bytes()
        result = subprocess.run([sys.executable,str(SCRIPT),str(source),'--sheet','Data',op,arg,'--out',str(dest)],capture_output=True,text=True,timeout=20)
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertEqual(source.read_bytes(),original)
        return load_workbook(dest),json.loads(result.stdout)
    def test_insert_columns_preserves_metadata_and_creates_unique_headers(self):
        wb,ws,t=self.table(['Column1','Q1','Q2','Total'],totals=True)
        t.tableColumns[3].calculatedColumnFormula=TableFormula(attr_text='SUM([@Q1],[@Q2])')
        t.tableColumns[3].totalsRowFunction='custom'
        t.tableColumns[3].totalsRowFormula=TableFormula(attr_text='SUM(D2:D3)')
        ws['D2']='=SUM([@Q1],[@Q2])';ws['D3']='=SUM([@Q1],[@Q2])'
        t.autoFilter.add_filter_column(2,['3']);t.autoFilter.sortState=SortState(ref='A2:D3',sortCondition=[SortCondition(ref='C2:C3')]); t.sortState=copy(t.autoFilter.sortState)
        ws['A2'].fill=PatternFill('solid',fgColor='FF0000');wb.create_sheet('Report')['A1']='=SUM(Sales[Q1])'
        out,report=self.run_edit(wb,'--insert-cols','B:2'); ws=out['Data'];t=ws.tables['Sales']
        self.assertEqual(t.ref,'A1:F4'); self.assertEqual([c.name for c in t.tableColumns],['Column1','Column2','Column3','Q1','Q2','Total'])
        self.assertEqual(len({c.id for c in t.tableColumns}),6); self.assertEqual(t.tableColumns[-1].id,4)
        self.assertEqual([ws.cell(1,col).value for col in range(1,7)],[c.name for c in t.tableColumns])
        self.assertEqual(ws['D2'].value,2);self.assertEqual(ws['F2'].value,'=SUM([@Q1],[@Q2])')
        self.assertEqual(ws['B2'].fill.fgColor.rgb,'00FF0000'); self.assertEqual(t.tableColumns[-1].totalsRowFunction,'custom')
        self.assertEqual(t.tableColumns[-1].totalsRowFormula.attr_text,'SUM(F2:F3)')
        self.assertEqual(t.autoFilter.ref,'A1:F3');self.assertEqual(t.autoFilter.filterColumn[0].colId,4)
        self.assertEqual(t.autoFilter.sortState.sortCondition[0].ref,'E2:E3');self.assertEqual(t.sortState.sortCondition[0].ref,'E2:E3')
        self.assertEqual(out['Report']['A1'].value,'=SUM(Sales[Q1])');self.assertEqual(report['tables']['Sales']['insertedColumns'],['Column2','Column3'])
    def test_delete_columns_updates_local_qualified_and_range_references(self):
        wb,ws,t=self.table(['A','B','C','D']); report=wb.create_sheet('Report')
        formulas=['=SUM(Sales[B])','=SUM(Sales[[A]:[D]])','=SUM(Sales[[B]:[D]])','=SUM(Sales[[B]:[C]])','=SUM(Sales[[#Headers],[B]])','=SUM(Sales[[#Data],[B]:[D]])']
        for i,f in enumerate(formulas,1):report.cell(i,1,f)
        t.tableColumns[-1].calculatedColumnFormula=TableFormula(attr_text='SUM([@B],[@C])');ws['D2']='=SUM([@B],[@C])'
        t.autoFilter.add_filter_column(1,['2']);t.autoFilter.add_filter_column(3,['4'])
        out,_=self.run_edit(wb,'--delete-cols','B:2'); t=out['Data'].tables['Sales']
        self.assertEqual(t.ref,'A1:B3'); self.assertEqual([c.name for c in t.tableColumns],['A','D'])
        self.assertEqual([out['Report'].cell(i,1).value for i in range(1,7)],['=SUM(#REF!)','=SUM(Sales[[A]:[D]])','=SUM(Sales[[D]:[D]])','=SUM(#REF!)','=SUM(#REF!)','=SUM(Sales[[#Data],[D]:[D]])'])
        self.assertEqual(out['Data']['B2'].value,'=SUM(#REF!,#REF!)');self.assertEqual(t.tableColumns[-1].calculatedColumnFormula.attr_text,'SUM(#REF!,#REF!)')
        self.assertEqual([c.colId for c in t.autoFilter.filterColumn],[1])
    def test_special_column_names_and_unchanged_external_workbooks(self):
        wb,ws,t=self.table(['A]B',"O'Brien",'Bang!']); report=wb.create_sheet('Report')
        report['A1']="=SUM(Sales[A']B])";report['A2']="=SUM(Sales[O''Brien])";report['A3']="='[external.xlsx]Data'!A1"
        out,_=self.run_edit(wb,'--delete-cols','A')
        self.assertEqual(out['Report']['A1'].value,'=SUM(#REF!)');self.assertEqual(out['Report']['A2'].value,"=SUM(Sales[O''Brien])");self.assertEqual(out['Report']['A3'].value,"='[external.xlsx]Data'!A1")
        report['A4']='=SUM(Sales[Bang!])'
        out,_=self.run_edit(wb,'--delete-cols','C')
        self.assertEqual(out['Report']['A4'].value,'=SUM(#REF!)')
        wb,ws,t=self.table(['external.xlsx','Keep'])
        ws['B2']='=[external.xlsx]Data!A1'
        out,_=self.run_edit(wb,'--delete-cols','A')
        self.assertEqual(out['Data']['A2'].value,'=[external.xlsx]Data!A1')
    def test_whole_table_deletion_rewrites_defined_names_and_bare_references(self):
        wb,ws,t=self.table(['A','B']); report=wb.create_sheet('Report'); report['A1']='=SUM(Sales)'; report['A2']='=SUM(Sales[A])'
        wb.defined_names.add(DefinedName('SalesValue',attr_text='Sales[A]'))
        out,_=self.run_edit(wb,'--delete-cols','A:2')
        self.assertNotIn('Sales',out['Data'].tables);self.assertEqual(out['Report']['A1'].value,'=SUM(#REF!)');self.assertEqual(out['Report']['A2'].value,'=SUM(#REF!)');self.assertEqual(out.defined_names['SalesValue'].attr_text,'#REF!')
    def test_header_and_totals_row_deletion_preserves_surviving_data(self):
        wb,ws,t=self.table(['A','B'],totals=True)
        out,_=self.run_edit(wb,'--delete-rows','1');t=out['Data'].tables['Sales']
        self.assertEqual(t.headerRowCount,0);self.assertEqual(out['Data']['A1'].value,1);self.assertIsNone(t.autoFilter)
        out,_=self.run_edit(wb,'--delete-rows','4');t=out['Data'].tables['Sales']
        self.assertEqual(t.totalsRowCount,0);self.assertFalse(t.totalsRowShown);self.assertEqual(out['Data']['A3'].value,10)
    def test_inserted_table_rows_receive_calculated_formulas(self):
        wb,ws,t=self.table(['A','B'],totals=True)
        t.tableColumns[1].calculatedColumnFormula=TableFormula(attr_text='A2*2')
        ws['B2']='=A2*2';ws['B3']='=A3*2';ws['B2'].number_format='$0.00'
        out,_=self.run_edit(wb,'--insert-rows','2:2')
        self.assertEqual(out['Data'].tables['Sales'].ref,'A1:B6')
        self.assertEqual([out['Data'].cell(row,2).value for row in range(2,6)],['=A2*2','=A3*2','=A4*2','=A5*2'])
        self.assertEqual(out['Data']['A6'].value,'Totals')
        self.assertEqual([out['Data'].cell(row,2).number_format for row in range(2,5)],['$0.00']*3)
        self.assertEqual(out['Data'].tables['Sales'].tableColumns[1].calculatedColumnFormula.attr_text,'A2*2')
        out,_=self.run_edit(wb,'--delete-rows','2')
        self.assertEqual(out['Data']['B2'].value,'=A2*2')
        self.assertEqual(out['Data'].tables['Sales'].tableColumns[1].calculatedColumnFormula.attr_text,'A2*2')
    def test_whole_row_column_ranges_and_unions_shift(self):
        r=rewrite.RefRewriter('Data','cols',2,1,False)
        self.assertEqual(r.rewrite('=SUM(B:D,$B:$D,Other!B:D,Data!B:D)','Data'),'=SUM(C:E,$C:$E,Other!B:D,Data!C:E)')
        self.assertEqual(r.rewrite('=XFE1+LOG10(B2)','Data'),'=XFE1+LOG10(C2)')
        r=rewrite.RefRewriter('Data','rows',2,2,True)
        self.assertEqual(r.rewrite('=SUM(1:4,2:3,Other!2:3,$4:$9)','Data'),'=SUM(1:2,#REF!,Other!2:3,$2:$7)')
        r=rewrite.RefRewriter('Data','rows',1,1,True)
        self.assertEqual(r.rewrite('=SUM(A$5:$A1)','Data'),'=SUM(A$4:$A1)')
        r=rewrite.RefRewriter('Data','cols',1,1,True)
        self.assertEqual(r.rewrite('=SUM($E5:A$1)','Data'),'=SUM($D5:A$1)')

if __name__=='__main__':unittest.main()
