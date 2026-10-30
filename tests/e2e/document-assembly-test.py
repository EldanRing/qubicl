"""Focused reopen checks for PDF forms, slide deletion and literal DOCX templates."""
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.action import PP_ACTION
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject, ContentStream
from reportlab.pdfgen.canvas import Canvas

ROOT = Path(__file__).resolve().parents[2]


def helper(skill, script, *args, env=None, success=True):
    result = subprocess.run([sys.executable, str(ROOT / 'skills/core' / skill / 'scripts' / script),
                             *map(str, args)], capture_output=True, text=True, env=env)
    if success and result.returncode:
        raise AssertionError(result.stderr or result.stdout)
    return result


def form(path, font='Helvetica', hierarchy=False, shared_default_fonts=False):
    stream = io.BytesIO()
    canvas = Canvas(stream)
    canvas.acroForm.textfield(name='name', value='SEED_NAME', x=100, y=600, fontName=font)
    canvas.acroForm.checkbox(name='agree', checked=True, x=100, y=550)
    canvas.showPage()
    canvas.acroForm.textfield(name='second', value='SECOND_PAGE', x=100, y=600, fontName=font)
    canvas.showPage()
    canvas.save()
    reader = PdfReader(io.BytesIO(stream.getvalue()))
    if hierarchy:
        reader.add_form_topname('person')
    writer = PdfWriter(clone_from=reader)
    if shared_default_fonts:
        acroform = writer.root_object['/AcroForm']
        font_map = writer._add_object(acroform['/DR']['/Font'])
        acroform['/DR'][NameObject('/Font')] = font_map
        writer.pages[0]['/Annots'][0].get_object()['/AP']['/N']['/Resources'][NameObject('/Font')] = font_map
    writer.write(path)


class PdfAssembly(unittest.TestCase):
    def test_merge_collisions_preserve_independent_fields_fonts_and_filling(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            first, second, merged = root / 'first.pdf', root / 'second.pdf', root / 'merged.pdf'
            form(first)
            form(second, 'Courier', shared_default_fonts=True)
            result = helper('pdf', 'pdf_merge.py', first, second, '-o', merged, '--bookmarks')
            receipt = json.loads(result.stdout)
            self.assertEqual(receipt['renamed_fields'][0]['prefix'], 'source_2')
            reader = PdfReader(merged)
            fields = reader.get_fields()
            self.assertEqual(len(reader.pages), 4)
            self.assertEqual(fields['name']['/V'], 'SEED_NAME')
            self.assertEqual(fields['source_2.name']['/V'], 'SEED_NAME')
            self.assertEqual(fields['source_2.agree']['/V'], '/Yes')
            self.assertEqual(len(reader.outline), 2)
            fonts = reader.root_object['/AcroForm']['/DR']['/Font']
            self.assertEqual(fonts['/Helv']['/BaseFont'], '/Helvetica')
            self.assertEqual(fonts['/Qubicl2_Cour']['/BaseFont'], '/Courier')
            appearance = reader.pages[2]['/Annots'][0].get_object()['/AP']['/N']
            appearance_fonts = appearance['/Resources']['/Font']
            self.assertIn('/Cour', appearance_fonts, 'unchanged appearance retains its original shared font map')
            for operands, operator in ContentStream(appearance, reader).operations:
                if operator == b'Tf':
                    self.assertIn(operands[0], appearance_fonts)
            values = root / 'values.json'
            values.write_text(json.dumps({'name': 'FIRST_FILLED', 'source_2.name': 'SECOND_FILLED'}))
            filled = root / 'filled.pdf'
            helper('pdf', 'pdf_fill_form.py', merged, '--fields-json', values, '-o', filled)
            fields = PdfReader(filled).get_fields()
            self.assertEqual(fields['name']['/V'], 'FIRST_FILLED')
            self.assertEqual(fields['source_2.name']['/V'], 'SECOND_FILLED')
            original = merged.read_bytes()
            failed = helper('pdf', 'pdf_merge.py', first, second, '-o', merged, '--duplicate-fields', 'error', success=False)
            self.assertNotEqual(failed.returncode, 0)
            self.assertIn('Duplicate form field', failed.stderr)
            self.assertEqual(merged.read_bytes(), original)

    def test_split_filters_hierarchical_widgets_and_supports_independent_repeated_pages(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, split = root / 'source.pdf', root / 'split.pdf'
            form(source, hierarchy=True)
            original = source.read_bytes()
            helper('pdf', 'pdf_split.py', source, '--pages', '1', '-o', split)
            reader = PdfReader(split)
            self.assertEqual(set(reader.get_fields()), {'person', 'person.name', 'person.agree'})
            values = root / 'values.json'
            values.write_text(json.dumps({'person.name': 'AFTER_SPLIT'}))
            helper('pdf', 'pdf_fill_form.py', split, '--fields-json', values, '-o', root / 'filled.pdf')
            self.assertEqual(PdfReader(root / 'filled.pdf').get_fields()['person.name']['/V'], 'AFTER_SPLIT')
            helper('pdf', 'pdf_split.py', source, '--pages', '1,1', '--rotate', '90', '--compress', '-o', split)
            reader = PdfReader(split)
            self.assertEqual(len(reader.pages), 2)
            self.assertEqual([len(page['/Annots']) for page in reader.pages], [2, 2])
            self.assertEqual([page.rotation for page in reader.pages], [90, 90])
            self.assertEqual(reader.get_fields()['source_2.person.name']['/V'], 'SEED_NAME')
            self.assertEqual(source.read_bytes(), original)

    def test_split_retains_nonvisual_calculations_without_resurrecting_omitted_widgets(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, split = root / 'source.pdf', root / 'split.pdf'
            form(source, hierarchy=True)
            writer = PdfWriter(clone_from=source)
            acroform = writer.root_object['/AcroForm']
            parent_ref = acroform['/Fields'][0]
            hidden = writer._add_object(DictionaryObject({NameObject('/FT'): NameObject('/Tx'),
                NameObject('/T'): TextStringObject('calculated'), NameObject('/V'): TextStringObject('HIDDEN_RESULT'),
                NameObject('/Parent'): parent_ref}))
            parent_ref.get_object()['/Kids'].append(hidden)
            excluded = writer.pages[1]['/Annots'][0]
            acroform[NameObject('/CO')] = ArrayObject([excluded, hidden])
            writer.write(source)
            helper('pdf', 'pdf_split.py', source, '--pages', '1', '-o', split)
            reader = PdfReader(split)
            fields = reader.get_fields()
            self.assertEqual(fields['person.calculated']['/V'], 'HIDDEN_RESULT')
            self.assertNotIn('person.second', fields)
            self.assertEqual([ref.get_object()['/T'] for ref in reader.root_object['/AcroForm']['/CO']], ['calculated'])
            self.assertNotIn(b'SECOND_PAGE', split.read_bytes())


def png(color):
    stream = io.BytesIO()
    Image.new('RGB', (16, 16), color).save(stream, format='PNG')
    stream.seek(0)
    return stream


class SlideDeletion(unittest.TestCase):
    def test_deletion_removes_incoming_actions_and_private_parts_but_keeps_shared_content(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, output = root / 'source.pptx', root / 'output.pptx'
            presentation = Presentation()
            first, target, last = [presentation.slides.add_slide(presentation.slide_layouts[6]) for _ in range(3)]
            link = first.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            link.text = 'Go to target'
            link.click_action.target_slide = target
            hover = link._element.xpath('.//a:hlinkClick')[0]
            from copy import deepcopy
            hover = deepcopy(hover)
            hover.tag = qn('a:hlinkMouseOver')
            link._element.xpath('.//p:cNvPr')[0].append(hover)
            target.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = 'TARGET_PRIVATE_CONTENT'
            target.notes_slide.notes_text_frame.text = 'TARGET_PRIVATE_NOTES'
            first.shapes.add_picture(png('red'), 0, 0)
            target.shapes.add_picture(png('red'), 0, 0)
            target.shapes.add_picture(png('blue'), 0, 0)
            data = CategoryChartData()
            data.categories = ['TARGET_PRIVATE_CHART']
            data.add_series('Private', [5])
            target.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(2), Inches(2), data)
            retained = last.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
            retained.text = 'Go to first'
            retained.click_action.target_slide = first
            presentation.save(source)
            helper('powerpoint', 'pptx_edit.py', source, '--remove-slide', '1', '--output', output)
            reopened = Presentation(output)
            self.assertEqual(len(reopened.slides), 2)
            self.assertEqual(reopened.slides[0].shapes[0].click_action.action, PP_ACTION.NONE)
            self.assertEqual(reopened.slides[1].shapes[0].click_action.target_slide, reopened.slides[0])
            self.assertEqual(reopened.slides[0].shapes[0].text, 'Go to target')
            live = {slide.part for slide in reopened.slides}
            self.assertTrue(all(rel.target_part in live for part in reopened.part.package.iter_parts()
                                for rel in part.rels.values() if rel.reltype == RT.SLIDE and not rel.is_external))
            with ZipFile(output) as archive:
                self.assertFalse(any(b'TARGET_PRIVATE' in archive.read(name) for name in archive.namelist()))
                self.assertEqual(len([name for name in archive.namelist() if name.startswith('ppt/media/')]), 1)
                self.assertFalse(any(name.startswith('ppt/embeddings/') for name in archive.namelist()))
                self.assertFalse(any(name.startswith('ppt/charts/') for name in archive.namelist()))
                self.assertFalse(any(b'hlinksldjump' in archive.read(name) for name in archive.namelist()
                                     if name == str(reopened.slides[0].part.partname).lstrip('/')))


class DocxTemplate(unittest.TestCase):
    def test_single_pass_literals_are_seed_independent_and_preserve_split_run_styles(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source, values = root / 'source.docx', root / 'values.json'
            document = Document()
            paragraph = document.add_paragraph()
            paragraph.add_run('Before ').italic = True
            paragraph.add_run('{{ a').bold = True
            paragraph.add_run(' }}').italic = True
            paragraph.add_run(' | {{b}} | {{a}} END')
            document.sections[0].header.paragraphs[0].text = '{{a}}'
            document.add_table(rows=1, cols=1).cell(0, 0).text = '{{b}}'
            document.save(source)
            values.write_text(json.dumps({'a': '{{b}}', 'b': 'FINAL_B\\1\nNEXT\tTAB'}))
            for seed in (0, 2):
                output = root / f'seed-{seed}.docx'
                result = helper('docx', 'docx_template.py', source, values, output, '--strict',
                                env={**os.environ, 'PYTHONHASHSEED': str(seed)})
                receipt = json.loads(result.stdout)
                self.assertEqual(receipt['filled'], {'a': 3, 'b': 2})
                self.assertEqual(receipt['unfilled_tokens'], [])
                reopened = Document(output)
                self.assertEqual(reopened.paragraphs[0].text, 'Before {{b}} | FINAL_B\\1\nNEXT\tTAB | {{b}} END')
                self.assertTrue(reopened.paragraphs[0].runs[0].italic)
                self.assertTrue(reopened.paragraphs[0].runs[1].bold)
                self.assertEqual(reopened.sections[0].header.paragraphs[0].text, '{{b}}')
                self.assertEqual(reopened.tables[0].cell(0, 0).text, 'FINAL_B\\1\nNEXT\tTAB')
            values.write_text(json.dumps({'a': 'provided'}))
            output = root / 'strict-existing.docx'
            output.write_bytes(b'preserve existing output')
            failed = helper('docx', 'docx_template.py', source, values, output, '--strict', success=False)
            self.assertEqual(failed.returncode, 1)
            self.assertFalse(json.loads(failed.stdout)['output_written'])
            self.assertEqual(json.loads(failed.stdout)['unfilled_tokens'], ['b'])
            self.assertEqual(output.read_bytes(), b'preserve existing output')


if __name__ == '__main__':
    unittest.main()
