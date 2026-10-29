"""Real save/reopen regressions for spreadsheet structural-edit dependents."""
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, NumData, NumVal, StrRef, StrData, StrVal
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import (OneCellAnchor, TwoCellAnchor, AbsoluteAnchor, AnchorMarker)
from openpyxl.formatting.rule import FormulaRule, ColorScaleRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableColumn
from openpyxl.workbook.defined_name import DefinedName
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / 'skills/core/xlsx/scripts'
sys.path.insert(0, str(SCRIPTS))
import xlsx_restructure as edit
import xlsx_restructure_dependents as dep


class Dependents(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def run_edit(self, wb, operation, value):
        before, after = self.root/'before.xlsx', self.root/'after.xlsx'
        wb.save(before)
        result = subprocess.run([sys.executable, str(SCRIPTS/'xlsx_restructure.py'), str(before), '--sheet', 'Data', operation, value, '--out', str(after)], capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        return load_workbook(after)

    def fixture(self):
        wb = Workbook(); ws = wb.active; ws.title = 'Data'; dashboard = wb.create_sheet('Dashboard')
        ws.append(['Category','Amount','Computed'])
        for name, value in [('A',2),('B',3),('C',4)]:
            row=ws.max_row+1; ws.append([name,value,f'=B{row}*2'])
        chart=BarChart(); chart.add_data(Reference(ws,min_col=2,min_row=1,max_row=4),titles_from_data=True)
        chart.series[0].val.numRef.numCache=NumData(ptCount=3,pt=[NumVal(idx=0,v=999)])
        chart.series[0].cat=AxDataSource(strRef=StrRef(f="'Data'!$A$2:$A$4",strCache=StrData(ptCount=1,pt=[StrVal(idx=0,v='stale')])))
        combined=LineChart(); combined.add_data(Reference(ws,min_col=2,min_row=2,max_row=4)); chart += combined
        dashboard.add_chart(chart,'B2')
        formula_chart=BarChart(); formula_chart.add_data(Reference(ws,min_col=3,min_row=2,max_row=4))
        formula_chart.series[0].val.numRef.numCache=NumData(ptCount=3,pt=[NumVal(idx=0,v=999)])
        dashboard.add_chart(formula_chart,'J2')
        local=BarChart(); local.add_data(Reference(ws,min_col=2,min_row=2,max_row=4)); local.set_categories(Reference(ws,min_col=1,min_row=2,max_row=4)); ws.add_chart(local,'B7')
        image_bytes=io.BytesIO(); PILImage.new('RGB',(8,8),'blue').save(image_bytes,format='PNG'); image_bytes.seek(0)
        picture=Image(image_bytes); picture.anchor=TwoCellAnchor(_from=AnchorMarker(col=1,row=1),to=AnchorMarker(col=3,row=4)); ws.add_image(picture)
        ws.defined_names.add(DefinedName('LocalAmounts',attr_text="$B$2:$B$4"))
        dashboard.defined_names.add(DefinedName('SourceAmounts',attr_text="'Data'!$B$2:$B$4"))
        ws.print_area='A1:C4'; ws.print_title_cols='B:C'; ws.print_title_rows='1:1'
        dv=DataValidation(type='custom',formula1='B2>$B$1'); dv.add('B2:B4'); ws.add_data_validation(dv)
        ws.conditional_formatting.add('B2:B4',FormulaRule(formula=['B2>0']))
        dashboard.conditional_formatting.add('A1:A3',FormulaRule(formula=['Data!B2>1']))
        wb.calculation=None
        return wb

    def test_insert_updates_cross_sheet_charts_rules_names_and_print_settings(self):
        wb=self.run_edit(self.fixture(),'--insert-cols','B')
        ws=wb['Data']; dashboard=wb['Dashboard']; series=dashboard._charts[0].series[0]
        self.assertEqual(series.val.numRef.f,"'Data'!$C$2:$C$4")
        self.assertEqual([p.v for p in series.val.numRef.numCache.pt],[2,3,4])
        self.assertEqual(dashboard._charts[0]._charts[1].series[0].val.numRef.f,"'Data'!$C$2:$C$4")
        self.assertEqual([p.v for p in series.cat.strRef.strCache.pt],['A','B','C'])
        self.assertEqual(series.tx.strRef.f,"'Data'!C1")
        self.assertEqual(series.tx.strRef.strCache.pt[0].v,'Amount')
        self.assertIsNone(dashboard._charts[1].series[0].val.numRef.numCache)
        self.assertEqual(ws['D2'].value,'=C2*2')
        self.assertEqual(ws._charts[0].anchor._from.col,2)
        self.assertEqual([p.v for p in ws._charts[0].series[0].cat.strRef.strCache.pt],['A','B','C'])
        self.assertEqual((ws._images[0].anchor._from.col,ws._images[0].anchor.to.col),(2,4))
        self.assertEqual(dashboard._charts[0].anchor._from.col,1)
        self.assertEqual(ws.defined_names['LocalAmounts'].attr_text,'$C$2:$C$4')
        self.assertEqual(dashboard.defined_names['SourceAmounts'].attr_text,"'Data'!$C$2:$C$4")
        self.assertEqual(str(ws.print_area),"'Data'!$A$1:$D$4")
        self.assertEqual(ws.print_title_cols,'$C:$D'); self.assertEqual(ws.print_title_rows,'$1:$1')
        dv=ws.data_validations.dataValidation[0]
        self.assertEqual(str(dv.sqref),'C2:C4'); self.assertEqual(dv.formula1,'C2>$C$1')
        cf=next(iter(ws.conditional_formatting)); self.assertEqual(str(cf.sqref),'C2:C4')
        self.assertEqual(ws.conditional_formatting[cf][0].formula,['C2>0'])
        cf=next(iter(dashboard.conditional_formatting)); self.assertEqual(dashboard.conditional_formatting[cf][0].formula,['Data!C2>1'])
        self.assertTrue(wb.calculation.fullCalcOnLoad); self.assertTrue(wb.calculation.forceFullCalc)

    def test_deleted_chart_source_row_refreshes_cached_points(self):
        wb=self.run_edit(self.fixture(),'--delete-rows','2')
        chart=wb['Dashboard']._charts[0]
        self.assertEqual(chart.series[0].val.numRef.f,"'Data'!$B$2:$B$3")
        self.assertEqual([p.v for p in chart.series[0].val.numRef.numCache.pt],[3,4])
        self.assertEqual([p.v for p in chart.series[0].cat.strRef.strCache.pt],['B','C'])
        self.assertEqual(chart._charts[1].series[0].val.numRef.f,"'Data'!$B$2:$B$3")
        self.assertEqual([p.v for p in chart._charts[1].series[0].val.numRef.numCache.pt],[3,4])

    def test_deleted_rule_origin_translates_before_structural_rewrite(self):
        wb=Workbook(); ws=wb.active; ws.title='Data'; wb.create_sheet('Dashboard')
        for row in range(1,7): ws.append([row,row*10,row*100])
        dv=DataValidation(type='custom',formula1='A2>0',formula2='$C$6'); dv.add('A2:A5'); ws.add_data_validation(dv)
        ws.conditional_formatting.add('A2:A5',FormulaRule(formula=['A2<$C$6']))
        ws.conditional_formatting.add('B2:B5',ColorScaleRule(start_type='formula',start_value='B2',end_type='formula',end_value='$C$6',start_color='FFFFFF',end_color='000000'))
        ws.print_title_rows='2:3'; ws.print_area='A2:C5'
        reopened=self.run_edit(wb,'--delete-rows','2')
        ws=reopened['Data']; dv=ws.data_validations.dataValidation[0]
        self.assertEqual(str(dv.sqref),'A2:A4'); self.assertEqual(dv.formula1,'A2>0'); self.assertEqual(dv.formula2,'$C$5')
        rules={str(cf.sqref):ws.conditional_formatting[cf][0] for cf in ws.conditional_formatting}
        self.assertEqual(rules['A2:A4'].formula,['A2<$C$5'])
        self.assertEqual([item.val for item in rules['B2:B4'].colorScale.cfvo],['B2','$C$5'])
        self.assertEqual(ws.print_title_rows,'$2:$2'); self.assertEqual(str(ws.print_area),"'Data'!$A$2:$C$4")

    def test_anchor_modes_and_deleted_disjoint_rule_origin(self):
        wb=Workbook(); ws=wb.active; ws.title='Data'
        for row in range(1,7): ws.append([row]*6)
        anchors=['C3', OneCellAnchor(_from=AnchorMarker(col=2,row=2)),
                 TwoCellAnchor(_from=AnchorMarker(col=1,row=1),to=AnchorMarker(col=4,row=4)),
                 TwoCellAnchor(editAs='oneCell',_from=AnchorMarker(col=1,row=1),to=AnchorMarker(col=4,row=4)),
                 TwoCellAnchor(editAs='absolute',_from=AnchorMarker(col=1,row=1),to=AnchorMarker(col=4,row=4)), AbsoluteAnchor()]
        for anchor in anchors:
            chart=BarChart(); chart.add_data(Reference(ws,min_col=6,min_row=1,max_row=6)); chart.anchor=anchor; ws.add_chart(chart)
        dv=DataValidation(type='custom',formula1='A1>0'); dv.add('A1:A2'); dv.add('C1:C5'); ws.add_data_validation(dv)
        ws.conditional_formatting.add('A5:A6 C1:C2',FormulaRule(formula=['IFERROR(A1,0)>0']))
        snapshot=dep.snapshot_dependents(wb); ws.delete_cols(1,2)
        rewriter=edit.RefRewriter('Data','cols',1,2,True)
        report=dep.apply_dependents(wb,snapshot,ws,'cols',1,2,True,rewriter.rewrite,edit.shift_point,edit.shift_range)
        self.assertEqual(ws._charts[0].anchor,'A3'); self.assertEqual(ws._charts[1].anchor._from.col,0)
        self.assertEqual((ws._charts[2].anchor._from.col,ws._charts[2].anchor.to.col),(0,2))
        self.assertEqual((ws._charts[3].anchor._from.col,ws._charts[3].anchor.to.col),(0,3))
        self.assertEqual((ws._charts[4].anchor._from.col,ws._charts[4].anchor.to.col),(1,4))
        self.assertEqual(ws._charts[5].anchor,anchors[5]); self.assertEqual(len(report['anchors']),4)
        dv=ws.data_validations.dataValidation[0]; self.assertEqual(str(dv.sqref),'A1:A5'); self.assertEqual(dv.formula1,'A1>0')
        cf=next(iter(ws.conditional_formatting)); self.assertEqual(str(cf.sqref),'A1:A2')
        self.assertEqual(ws.conditional_formatting[cf][0].formula,['IFERROR(#REF!,0)>0'])
        out=self.root/'anchors.xlsx'; wb.save(out); self.assertEqual(len(load_workbook(out)['Data']._charts),6)

    def test_static_cache_sources_keep_union_intersection_and_table_selection(self):
        wb=Workbook(); ws=wb.active; ws.title='Data'
        ws.append(['A','B','C']); ws.append([1,2,3]); ws.append([4,5,6])
        table=Table(displayName='Sales',ref='A1:C3'); table.tableColumns=[TableColumn(id=i+1,name=name) for i,name in enumerate(['A','B','C'])]; ws.add_table(table)
        def values(formula):
            return [[cell.value for cell in row] for row in dep._cells_for_reference(wb,'Data',formula)]
        self.assertEqual(values('A2:A3,A2:A3'),[[1],[4],[1],[4]])
        self.assertEqual(values('A3:A2'),[[1],[4]])
        self.assertEqual(values('A2:B3 B3:C3'),[[5]])
        self.assertEqual(values('(A2:A3,A2:A3)'),[[1],[4],[1],[4]])
        self.assertEqual(values('Sales[[A],[C]]'),[[1,3],[4,6]])
        self.assertEqual(values('Sales[[A]:[C]]'),[[1,2,3],[4,5,6]])
        self.assertEqual(values('Sales[[#Headers],[A],[C]]'),[['A','C']])
        table.tableColumns[0].name='Bang!'; table.tableColumns[2].name='#REF!'
        self.assertEqual(values('Sales[Bang!]'),[[1],[4]])
        self.assertEqual(values("Sales['#REF!]"),[[3],[6]])


if __name__ == '__main__': unittest.main()
