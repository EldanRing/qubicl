"""Small real-file regressions for the 0.8 office data-preservation audit."""
import base64
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / 'skills/core'
sys.path.insert(0, str(SCRIPTS / 'docx/scripts'))
def module(kind, name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / kind / 'scripts' / (name + '.py'))
    value = importlib.util.module_from_spec(spec); sys.modules[spec.name] = value; spec.loader.exec_module(value); return value

def cli(kind, name, *args, ok=True):
    result = subprocess.run([sys.executable, str(SCRIPTS / kind / 'scripts' / (name + '.py')), *map(str,args)], capture_output=True, text=True, timeout=20)
    if (result.returncode == 0) != ok: raise AssertionError(result.stdout + result.stderr)
    return result

class OfficeAudit(unittest.TestCase):
    def setUp(self): self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup); self.dir=Path(self.tmp.name)
    def test_csv_identifiers_literals_and_ambiguous_columns(self):
        from openpyxl import load_workbook
        source=self.dir/'source.csv'; dest=self.dir/'book.xlsx'
        source.write_text('id,mixed,=header,formula\n00123,12,x,=1+1\n12345678901234567890,text,y,=2+2\n')
        for flags in [[], ['--no-infer']]:
            cli('xlsx','csv_to_xlsx',source,dest,*flags)
            ws=load_workbook(dest).active
            self.assertEqual(ws['A2'].value,'00123'); self.assertEqual(ws['A3'].value,'12345678901234567890')
            self.assertEqual(ws['B2'].value,'12'); self.assertEqual(ws['C1'].data_type,'s'); self.assertEqual(ws['D2'].data_type,'s')
        cli('xlsx','csv_to_xlsx',source,dest,'--column-types','{"D":"formula"}')
        self.assertEqual(load_workbook(dest).active['D2'].data_type,'f')
    def test_missing_formula_results_do_not_replace_output(self):
        from openpyxl import Workbook
        wb=Workbook(); wb.active['A1']='=1+1'; source=self.dir/'book.xlsx'; wb.save(source)
        output=self.dir/'data.csv'; output.write_text('preserve')
        result=cli('xlsx','xlsx_to_csv',source,output,'--data-only',ok=False)
        self.assertIn('A1',result.stderr); self.assertEqual(output.read_text(),'preserve')
        result=cli('xlsx','xlsx_to_csv',source,output,'--data-only','--allow-missing-results')
        self.assertFalse(json.loads(result.stdout)['valueComplete'])
    def test_structured_and_external_references(self):
        rewrite=module('xlsx','xlsx_restructure').RefRewriter('Sheet','rows',1,1,False)
        self.assertEqual(rewrite.rewrite('=SUM(Sales[Q1])+A1+SUM(Sales[[#Headers],[Q1]])','Sheet'), '=SUM(Sales[Q1])+A2+SUM(Sales[[#Headers],[Q1]])')
        self.assertEqual(rewrite.rewrite("='[external.xlsx]Sheet'!A1",'Sheet'),"='[external.xlsx]Sheet'!A1")
    def restructure(self, wb, *flags):
        from openpyxl import load_workbook
        source=self.dir/'source.xlsx'; out=self.dir/'out.xlsx'; wb.save(source)
        cli('xlsx','xlsx_restructure',source,'--out',out,'--sheet','Sheet',flags[0],flags[1]+':'+flags[2])
        return load_workbook(out)
    def test_dimensions_deleted_rules_tables_and_filter(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.table import Table
        from openpyxl.worksheet.datavalidation import DataValidation
        wb=Workbook(); ws=wb.active; ws.append(['Q1','Q2']); ws.append([1,2]); ws.append([3,4]); ws.add_table(Table(displayName='Sales',ref='A1:B3'))
        ws.column_dimensions['B'].width=25
        out=self.restructure(wb,'--insert-rows','1','1'); table=out.active.tables['Sales']
        self.assertEqual(table.ref,'A2:B4'); self.assertEqual(table.autoFilter.ref,'A2:B4')
        out=self.restructure(wb,'--insert-cols','1','1'); self.assertEqual(out.active.column_dimensions['C'].width,25)
        wb=Workbook(); ws=wb.active; ws['A1']='H'; ws['A2']='yes'; ws['A3']='free'; ws.row_dimensions[2].height=45
        dv=DataValidation(type='list',formula1='"yes,no"'); dv.add('A2'); ws.add_data_validation(dv)
        out=self.restructure(wb,'--delete-rows','2','1'); self.assertEqual(out.active['A2'].value,'free'); self.assertIsNone(out.active.row_dimensions[2].height); self.assertEqual(len(out.active.data_validations.dataValidation),0)
        ws.add_table(Table(displayName='Gone',ref='A1:A2'))
        out=self.restructure(wb,'--delete-rows','1','2'); self.assertNotIn('Gone',out.active.tables)
    def test_word_replacement_and_comment_preserve_image(self):
        from docx import Document
        from docx.oxml.ns import qn
        common=module('docx','docx_common'); comments=module('docx','docx_comments')
        image=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M/wHwAF/gL+X2NDWQAAAABJRU5ErkJggg==')
        doc=Document(); para=doc.add_paragraph(); run=para.add_run('before cat after'); run.add_picture(io.BytesIO(image))
        self.assertEqual(common.replace_in_paragraph(para,'cat','cats'),1)
        self.assertEqual(common.replace_in_paragraph(para,'cats','cats'),1)
        self.assertEqual(len(doc.inline_shapes),1)
        _, anchors=comments.find_anchor_runs(doc,'cats'); comments.add_comment_native(doc,anchors,'Review','Qubicl Maintainers','QM'); self.assertEqual(len(doc.inline_shapes),1)
        doc.save(self.dir/'word.docx'); reopened=Document(self.dir/'word.docx'); self.assertEqual(len(reopened.inline_shapes),1)
        self.assertEqual(para.text,'before cats after')
        self.assertEqual(comments.list_comments(reopened)[0]['anchored_text'],'cats')
        para=doc.add_paragraph(); para.add_run('ca'); para.add_run('t cat'); self.assertEqual(common.replace_in_paragraph(para,'cat','dog'),2); self.assertEqual(para.text,'dog dog')
    def test_word_structural_boundaries_and_comment_offsets(self):
        from docx import Document
        from docx.text.run import Run
        common=module('docx','docx_common'); comments=module('docx','docx_comments')
        doc=Document(); para=doc.add_paragraph(); para.add_run('ca\nt')
        self.assertEqual(common.replace_in_paragraph(para,'cat','dog'),0)
        self.assertEqual(para.text,'ca\nt')
        para=doc.add_paragraph('prefix\ttarget after')
        _, runs=comments.find_anchor_runs(doc,'target')
        self.assertEqual(''.join(Run(run,para).text for run in runs),'target')
        self.assertEqual(para.text,'prefix\ttarget after')
        para=doc.add_paragraph('cat cat')
        self.assertEqual(common.replace_in_paragraph(para,'cat','dog\nnext\tend'),2)
        self.assertEqual(para.text,'dog\nnext\tend dog\nnext\tend')
        self.assertEqual(len(para._p.xpath('.//w:br')),2)
        self.assertEqual(len(para._p.xpath('.//w:tab')),2)
    def test_word_shared_headers_and_merged_cells_edited_once(self):
        from docx import Document
        doc=Document(); doc.sections[0].header.paragraphs[0].text='cat'
        doc.add_section(); table=doc.add_table(rows=1,cols=2); table.cell(0,0).merge(table.cell(0,1)).text='cat'
        common=module('docx','docx_common')
        total=sum(common.replace_in_paragraph(para,'cat','cats') for para in common.iter_all_paragraphs(doc))
        self.assertEqual(total,2)
        self.assertEqual(doc.sections[0].header.paragraphs[0].text,'cats')
        self.assertEqual(table.cell(0,0).text,'cats')
    def test_filter_criteria_and_table_column_schema_preserved(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.table import Table
        wb=Workbook(); ws=wb.active
        ws.append(['A','B','C']); ws.append([1,2,3]); ws.append([4,5,6])
        ws.auto_filter.ref='A1:C3'; ws.auto_filter.add_filter_column(1,['2']); ws.auto_filter.add_sort_condition('B2:B3')
        out=self.restructure(wb,'--insert-cols','2','1')
        self.assertEqual(out.active.auto_filter.ref,'A1:D3')
        self.assertEqual(out.active.auto_filter.filterColumn[0].colId,2)
        self.assertEqual(out.active.auto_filter.sortState.sortCondition[0].ref,'C2:C3')
        out=self.restructure(wb,'--delete-cols','2','1')
        self.assertEqual(out.active.auto_filter.filterColumn,[])
        ws.add_table(Table(displayName='Schema',ref='A1:C3'))
        from openpyxl import load_workbook
        source=self.dir/'table.xlsx'; target=self.dir/'table-output.xlsx'; wb.save(source)
        original=source.read_bytes()
        for flag, names, ref in [('--insert-cols',['A','Column1','B','C'],'A1:D3'),
                                 ('--delete-cols',['A','C'],'A1:B3')]:
            cli('xlsx','xlsx_restructure',source,flag,'B','--out',target)
            edited=load_workbook(target); table=edited.active.tables['Schema']
            self.assertEqual(table.ref,ref)
            self.assertEqual([column.name for column in table.tableColumns],names)
            self.assertEqual(source.read_bytes(),original)

    def test_word_revisions_are_read_in_body_tables_headers(self):
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        doc=Document()
        for para in [doc.add_paragraph('Total: '),doc.add_table(1,1).cell(0,0).paragraphs[0],doc.sections[0].header.paragraphs[0]]:
            ins=OxmlElement('w:ins'); run=OxmlElement('w:r'); txt=OxmlElement('w:t'); txt.text='$200'; run.append(txt); ins.append(run); para._p.append(ins)
        path=self.dir/'revision.docx'; doc.save(path)
        result=json.loads(cli('docx','docx_read',path,'--text').stdout)
        self.assertTrue(result['completeFinalView']); self.assertTrue(result['revisions']['insertions'])
        self.assertEqual(result['body'],['Total: $200']); self.assertEqual(result['tables'],[[['$200']]])
        self.assertEqual(result['headers'],['$200']); self.assertNotIn('warnings',result)
    def test_pptx_mixed_runs_groups_background_footer(self):
        from pptx import Presentation
        from pptx.util import Inches
        from pptx.dml.color import RGBColor
        edit=module('powerpoint','pptx_edit'); prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); group=slide.shapes.add_group_shape()
        box=group.shapes.add_textbox(0,0, Inches(2), Inches(1)); para=box.text_frame.paragraphs[0]
        for text in ['cat ','ca','t']: para.add_run().text=text
        self.assertEqual(edit.replace_text(prs,'cat','dog'),2); self.assertEqual(para.text,'dog dog')
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string('FF0000')
        for i in range(10): slide.shapes.add_textbox(i,i,100,100)
        edit.set_footer(slide,'footer'); ids=[s.shape_id for s in edit.iter_shapes(slide.shapes)]; self.assertEqual(len(ids),len(set(ids)))
        index=edit.duplicate_slide(prs,0); self.assertEqual(str(prs.slides[index].background.fill.fore_color.rgb),'FF0000')
        prs.save(self.dir/'deck.pptx'); reopened=Presentation(self.dir/'deck.pptx'); self.assertEqual(str(reopened.slides[1].background.fill.fore_color.rgb),'FF0000')
    def test_grouped_chart_duplicate_has_independent_data(self):
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); data=CategoryChartData(); data.categories=['a']; data.add_series('s',[1]); group=slide.shapes.add_group_shape(); group.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,0,1000,1000,data)
        edit = module('powerpoint','pptx_edit')
        duplicate = edit.duplicate_slide(prs,0)
        copied = next(shape.chart for shape in edit.iter_shapes(prs.slides[duplicate].shapes) if shape.has_chart)
        data = CategoryChartData(); data.categories=['a']; data.add_series('s',[99]); copied.replace_data(data)
        self.assertEqual(next(shape.chart for shape in edit.iter_shapes(slide.shapes) if shape.has_chart).series[0].values, (1.0,))
        self.assertEqual(copied.series[0].values, (99.0,))
        self.assertEqual(len(prs.slides),2)

if __name__=='__main__': unittest.main()
