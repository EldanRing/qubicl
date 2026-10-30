"""Save/reopen regressions for typed formulas and elapsed-time exports."""
import csv
import io
import json
from datetime import date, timedelta
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / 'skills/core/xlsx/scripts'


class ValueTypes(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def run_cli(self, script, source, *args):
        result = subprocess.run([sys.executable, str(SCRIPTS / script), str(source), *args],
                                capture_output=True, text=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stderr)
        return result.stdout

    def workbook(self):
        wb = Workbook(); ws = wb.active; ws.title = 'Data'
        ws['A1'] = ArrayFormula(ref='A1:A2', text='=B1:B2*2')
        ws['A1'].number_format = '0.00'
        ws['B1'] = 2; ws['B2'] = 3
        ws['D1'] = DataTableFormula(ref='D1:E3', dt2D=True, r1='B1', r2='B2')
        other = wb.create_sheet('Report')
        other['A1'] = ArrayFormula(ref='A1:A2', text='=Data!B1:B2')
        other['B1'] = '=SUM(Data!B1:B2)'
        source = self.root / 'source.xlsx'; wb.save(source)
        return source

    def test_reader_inventories_array_and_data_table_formulas_and_serializes_durations(self):
        source = self.workbook(); wb = load_workbook(source); ws = wb['Data']
        for row, duration in enumerate((timedelta(hours=27, minutes=15), timedelta(hours=-1, minutes=-30), timedelta(0), timedelta(seconds=1, milliseconds=125)), 1):
            ws.cell(row, 7, duration).number_format = '[h]:mm:ss.000'
        ws['H1'] = date(2026, 10, 29); wb.save(source)
        result = json.loads(self.run_cli('xlsx_read.py', source, '--json', '--sheet', 'Data'))
        self.assertEqual([row[6] for row in result['rows']], ['P1DT3H15M', '-PT1H30M', 'PT0S', 'PT1.125S'])
        self.assertEqual(result['rows'][0][7], '2026-10-29T00:00:00')
        self.assertEqual(result['rows'][0][0], {'type': 'formula', 'formula_type': 'array', 'ref': 'A1:A2', 'text': '=B1:B2*2'})
        self.assertEqual(result['rows'][0][3]['attributes']['r2'], 'B2')
        formulas = json.loads(self.run_cli('xlsx_read.py', source, '--formulas'))['formulas']
        by_cell = {(item['sheet'], item['cell']): item for item in formulas}
        self.assertEqual(len(formulas), 4)
        self.assertEqual(by_cell['Data', 'A1']['formula'], '=B1:B2*2')
        self.assertEqual(by_cell['Data', 'A1']['ref'], 'A1:A2')
        self.assertEqual(by_cell['Data', 'D1']['formula_type'], 'dataTable')
        self.assertEqual(by_cell['Data', 'D1']['attributes'], {'ca': False, 'dt2D': True, 'dtr': False, 'r1': 'B1', 'r2': 'B2', 'del1': False, 'del2': False})
        rows = list(csv.reader(io.StringIO(self.run_cli('xlsx_read.py', source, '--csv', '--sheet', 'Data'))))
        self.assertEqual(json.loads(rows[0][0])['formula_type'], 'array')
        self.assertEqual(rows[0][6], 'P1DT3H15M')
        exported = self.root / 'standalone.csv'
        self.run_cli('xlsx_to_csv.py', source, str(exported), '--sheet', 'Data')
        with exported.open(newline='') as stream:
            standalone = list(csv.reader(stream))
        self.assertEqual(json.loads(standalone[0][0]), result['rows'][0][0])
        self.assertEqual(json.loads(standalone[0][3]), result['rows'][0][3])
        self.assertEqual([row[6] for row in standalone], [row[6] for row in rows])
        self.assertEqual(standalone[0][7], '2026-10-29', 'standalone export retains its date-only convention')

    def test_typed_cell_reads_distinguish_missing_empty_zero_and_error_formula_caches(self):
        wb = Workbook(); ws = wb.active
        ws['A1'] = '=1+2'
        ws['A2'] = ArrayFormula(ref='A2:A3', text='=B2:B3*2')
        ws['B1'] = timedelta(hours=27, minutes=15); ws['B1'].number_format = '[h]:mm'
        ws['B2'] = '00123'
        ws['C1'] = '=""'; ws['C2'] = '=0'; ws['C3'] = '=1/0'
        ws['D1'] = '=literal'; ws['D1'].data_type = 's'
        source = self.root / 'typed.xlsx'; wb.save(source)
        with ZipFile(source) as archive:
            parts = {name: archive.read(name) for name in archive.namelist()}
        ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        xml = ET.fromstring(parts['xl/worksheets/sheet1.xml'])
        cache_values = {'C1': ('str', ''), 'C2': ('n', '0'), 'C3': ('e', '#DIV/0!')}
        for cell in xml.iter(ns + 'c'):
            if cell.get('r') in cache_values:
                kind, value = cache_values[cell.get('r')]
                cell.set('t', kind); cell.find(ns + 'v').text = value
        parts['xl/worksheets/sheet1.xml'] = ET.tostring(xml)
        with ZipFile(source, 'w', ZIP_DEFLATED) as archive:
            for name, content in parts.items(): archive.writestr(name, content)
        original = source.read_bytes()
        result = json.loads(self.run_cli('xlsx_read.py', source, '--cells'))
        cells = {cell['cell']: cell for cell in result['cells']}
        self.assertEqual(cells['B1']['type'], 'duration')
        self.assertEqual(cells['B1']['value'], 'P1DT3H15M')
        self.assertAlmostEqual(float(cells['B1']['rawValue']), 1.135416666666667)
        self.assertEqual(cells['B1']['numberFormat'], '[h]:mm')
        self.assertEqual(cells['B2']['value'], '00123')
        self.assertEqual(cells['A2']['formula']['formula_type'], 'array')
        self.assertEqual(cells['A2']['formula']['ref'], 'A2:A3')
        for address in ('A1', 'A2'):
            self.assertFalse(cells[address]['cachedValueAvailable'])
            self.assertEqual(cells[address]['cacheStatus'], 'missing')
            self.assertEqual(cells[address]['warnings'][0]['cell'], address)
        for address, value in [('C1', ''), ('C2', 0), ('C3', '#DIV/0!')]:
            self.assertTrue(cells[address]['cachedValueAvailable'])
            self.assertEqual(cells[address]['cached'], value)
        formulas = json.loads(self.run_cli('xlsx_read.py', source, '--formulas'))['formulas']
        self.assertNotIn('D1', {item['cell'] for item in formulas}, 'literal equals-prefixed text is not a formula')
        self.assertFalse(next(item for item in formulas if item['cell'] == 'A2')['cachedValueAvailable'])
        self.assertFalse(result['recalculated'])
        self.assertEqual(result['displayValues'], 'not_rendered')
        self.assertEqual(source.read_bytes(), original)

    def test_row_and_column_insertions_move_typed_anchors_ranges_and_cross_sheet_formulas(self):
        for axis, script, flag, anchor, ref, formula, table_anchor, table_ref, r1, r2 in (
                ('rows', 'xlsx_edit.py', '--insert-rows', 'A2', 'A2:A3', '=B2:B3*2', 'D2', 'D2:E4', 'B2', 'B3'),
                ('cols', 'xlsx_restructure.py', '--insert-cols', 'B1', 'B1:B2', '=C1:C2*2', 'E1', 'E1:F3', 'C1', 'C2')):
            with self.subTest(axis=axis):
                source = self.workbook(); original = source.read_bytes(); output = self.root / f'{axis}.xlsx'
                self.run_cli(script, source, '--sheet', 'Data', flag, '1', '--out', str(output))
                self.assertEqual(source.read_bytes(), original)
                wb = load_workbook(output); ws = wb['Data']; value = ws[anchor].value
                self.assertIsInstance(value, ArrayFormula)
                self.assertEqual((value.ref, value.text), (ref, formula))
                self.assertEqual(ws[anchor].number_format, '0.00')
                self.assertEqual(wb['Report']['A1'].value.ref, 'A1:A2')
                self.assertEqual(wb['Report']['A1'].value.text, '=Data!' + formula[1:-2])
                table = ws[table_anchor].value
                self.assertIsInstance(table, DataTableFormula)
                self.assertEqual((table.ref, table.r1, table.r2), (table_ref, r1, r2))
                self.assertTrue(table.dt2D)
                self.assertTrue(table.ca)

    def test_deleting_an_array_anchor_promotes_surviving_range_and_full_deletion_removes_it(self):
        wb = Workbook(); ws = wb.active; ws.title = 'Data'
        ws['A1'] = ArrayFormula(ref='A1:B3', text='=C1:D3')
        ws['A2'] = 2; ws['A3'] = 3; ws['B2'] = 5; ws['B3'] = 6
        ws['A2'].number_format = '0.000'; ws['E5'] = 'keep'
        source = self.root / 'source.xlsx'; wb.save(source)
        output = self.root / 'partial.xlsx'
        self.run_cli('xlsx_restructure.py', source, '--sheet', 'Data', '--delete-rows', '1', '--out', str(output))
        ws = load_workbook(output)['Data']
        self.assertIsInstance(ws['A1'].value, ArrayFormula)
        self.assertEqual((ws['A1'].value.ref, ws['A1'].value.text), ('A1:B2', '=C1:D2'))
        self.assertEqual(ws['A1'].number_format, '0.000')
        self.assertEqual(ws['B1'].value, 5)
        self.assertEqual(ws['E4'].value, 'keep')
        final = self.root / 'deleted.xlsx'
        report = json.loads(self.run_cli('xlsx_restructure.py', output, '--sheet', 'Data', '--delete-cols', '1:2', '--out', str(final)))
        ws = load_workbook(final)['Data']
        self.assertEqual(ws.array_formulae, {})
        self.assertEqual(ws['C4'].value, 'keep')
        self.assertTrue(any(item['to'] is None for item in report['formulas']))

    def test_deleted_data_table_input_is_marked_without_a_dangling_reference(self):
        source = self.workbook(); output = self.root / 'deleted-input.xlsx'
        self.run_cli('xlsx_restructure.py', source, '--sheet', 'Data', '--delete-cols', '2', '--out', str(output))
        table = load_workbook(output)['Data']['C1'].value
        self.assertIsInstance(table, DataTableFormula)
        self.assertEqual(table.ref, 'C1:D3')
        self.assertEqual((table.r1, table.r2), (None, None))
        self.assertTrue(table.del1); self.assertTrue(table.del2)
        inventory = json.loads(self.run_cli('xlsx_read.py', output, '--formulas', '--sheet', 'Data'))
        self.assertEqual(next(item for item in inventory['formulas'] if item['cell'] == 'C1')['attributes']['del1'], True)


if __name__ == '__main__':
    unittest.main()
