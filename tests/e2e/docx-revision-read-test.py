"""Real package regressions for final/original/markup tracked-change reading."""
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.part import Part
from docx.opc.packuri import PackURI
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).resolve().parents[2]
SCRIPTS=ROOT/'skills/core/docx/scripts'
sys.path.insert(0,str(SCRIPTS))
from docx_read import read_view, extract_text
from docx_revision_view import project_package


def revision(kind,text,identifier='1'):
    node=OxmlElement('w:'+kind); node.set(qn('w:id'),identifier); node.set(qn('w:author'),'Fixture'); node.set(qn('w:date'),'2026-10-29T00:00:00Z')
    if text is not None:
        run=OxmlElement('w:r'); leaf=OxmlElement('w:delText' if kind in ('del','moveFrom') else 'w:t'); leaf.text=text; run.append(leaf); node.append(run)
    return node

def para_revision(para,kind,text,identifier='1'):
    para._p.append(revision(kind,text,identifier))

def change_property(properties,tag,previous,identifier='1'):
    node=revision(tag,None,identifier); node.append(previous); properties.append(node)

def paragraph_mark(para,kind,identifier='1'):
    marks=OxmlElement('w:rPr'); marks.append(revision(kind,None,identifier)); para._p.get_or_add_pPr().append(marks)

def save_zip(path,changes):
    with zipfile.ZipFile(path) as package: entries={info.filename:(info,package.read(info.filename)) for info in package.infolist()}
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as package:
        for name,(info,data) in entries.items(): package.writestr(info,changes.pop(name,data))
        for name,data in changes.items(): package.writestr(name,data)

class RevisionRead(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup); self.path=Path(self.tmp.name)/'fixture.docx'
    def cli(self,*args):
        result=subprocess.run([sys.executable,str(SCRIPTS/'docx_read.py'),str(self.path),*args],text=True,capture_output=True,timeout=20)
        self.assertEqual(result.returncode,0,result.stderr); return json.loads(result.stdout)
    def test_final_original_markup_cover_related_stories_and_do_not_mutate(self):
        doc=Document(); para=doc.add_paragraph('Total: ')
        para_revision(para,'del','$100','1'); para_revision(para,'ins','$200','2')
        para_revision(doc.add_paragraph(),'moveFrom','moved','3'); para_revision(doc.add_paragraph(),'moveTo','moved','4')
        cell=doc.add_table(1,1).cell(0,0); para_revision(cell.paragraphs[0],'ins','cell-new','5')
        nested=cell.add_table(1,1); para_revision(nested.cell(0,0).paragraphs[0],'del','nested-old','6')
        section=doc.sections[0]; section.different_first_page_header_footer=True
        for i,part in enumerate([section.header,section.first_page_header,section.even_page_header,section.footer,section.first_page_footer,section.even_page_footer]):
            para_revision(part.paragraphs[0],'ins',f'variant-{i}',str(10+i))
        box=OxmlElement('w:txbxContent'); p=OxmlElement('w:p'); p.append(revision('ins','textbox-new','20')); box.append(p)
        doc.add_paragraph().add_run()._r.append(box)
        note_ref=OxmlElement('w:footnoteReference'); note_ref.set(qn('w:id'),'5'); doc.add_paragraph().add_run()._r.append(note_ref)
        notes=OxmlElement('w:footnotes'); note=OxmlElement('w:footnote'); note.set(qn('w:id'),'5'); p=OxmlElement('w:p'); p.append(revision('del','note-old','21')); p.append(revision('ins','note-new','22')); note.append(p); notes.append(note)
        unused=OxmlElement('w:footnote'); unused.set(qn('w:id'),'6'); para=OxmlElement('w:p'); props=OxmlElement('w:pPr'); style=OxmlElement('w:pStyle'); style.set(qn('w:val'),'Title'); props.append(style); para.append(props); unused.append(para); notes.append(unused)
        part=Part(PackURI('/word/notes-custom.xml'),'application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml',etree.tostring(notes),doc.part.package)
        doc.part.relate_to(part,RT.FOOTNOTES); doc.save(self.path)
        before=hashlib.sha256(self.path.read_bytes()).hexdigest()
        final=self.cli('--text'); original=self.cli('--text','--mode','original'); markup=self.cli('--text','--mode','markup')
        self.assertEqual(final['body'][:3],['Total: $200','','moved']); self.assertEqual(original['body'][:3],['Total: $100','moved',''])
        self.assertEqual(final['tables'][0][0][0],'cell-new\n\n'); self.assertIn('nested-old',original['tables'][0][0][0])
        self.assertEqual(final['headers'],['variant-0','variant-1','variant-2']); self.assertEqual(final['footers'],['variant-3','variant-4','variant-5'])
        self.assertEqual(final['footnotes'][0]['blocks'][0]['text'],'note-new'); self.assertEqual(original['footnotes'][0]['blocks'][0]['text'],'note-old')
        self.assertEqual(final['textboxes'][0]['blocks'][0]['text'],'textbox-new')
        self.assertNotIn('Title',self.cli('--styles')['styles'])
        self.assertEqual(markup['final']['body'],final['body']); self.assertEqual(markup['original']['body'],original['body'])
        self.assertTrue(all(record['part'] and record['path'] for record in markup['revisions']['records']))
        self.assertTrue(final['completeFinalView']); self.assertTrue(original['completeView']); self.assertNotIn('warnings',final)
        self.assertEqual(hashlib.sha256(self.path.read_bytes()).hexdigest(),before)
        loaded=Document(self.path); xml=loaded.element.xml
        self.assertEqual(extract_text(loaded)['body'],final['body']); self.assertEqual(loaded.element.xml,xml)
    def test_revised_rows_cells_and_merge_state(self):
        doc=Document(); table=doc.add_table(3,2)
        for i,row in enumerate(table.rows):
            for j,cell in enumerate(row.cells): cell.text=f'{i}-{j}'
        table.rows[0]._tr.get_or_add_trPr().append(revision('ins',None,'1'))
        table.rows[1]._tr.get_or_add_trPr().append(revision('del',None,'2'))
        table.cell(2,0)._tc.get_or_add_tcPr().append(revision('cellIns',None,'3'))
        table.cell(2,1)._tc.get_or_add_tcPr().append(revision('cellDel',None,'4'))
        doc.save(self.path)
        self.assertEqual(read_view(self.path,'final')['tables'],[[['0-0','0-1'],['2-0']]])
        self.assertEqual(read_view(self.path,'original')['tables'],[[['1-0','1-1'],['2-1']]])
        # A revised vertical merge affects structure even when the text is unchanged.
        merge=revision('cellMerge',None,'5'); merge.set(qn('w:vMerge'),'cont'); merge.set(qn('w:vMergeOrig'),'rest')
        table.cell(0,0)._tc.get_or_add_tcPr().append(merge); doc.save(self.path)
        final=read_view(self.path,'final')['blocks'][0]['rows'][0][0]
        self.assertEqual(final['verticalMerge'],'continue')
    def test_paragraph_marks_and_property_snapshots(self):
        doc=Document(); first=doc.add_paragraph('one'); second=doc.add_paragraph('two'); doc.add_paragraph('three')
        for para in [first,second]: paragraph_mark(para,'del','7')
        heading=doc.add_paragraph('heading',style='Heading 2'); old=OxmlElement('w:pPr'); style=OxmlElement('w:pStyle'); style.set(qn('w:val'),'Heading1'); old.append(style); change_property(heading._p.get_or_add_pPr(),'pPrChange',old,'8')
        run=heading.add_run(' format'); run.bold=True; prior=OxmlElement('w:rPr'); prior.append(OxmlElement('w:i')); change_property(run._r.get_or_add_rPr(),'rPrChange',prior,'9')
        doc.save(self.path)
        self.assertEqual(read_view(self.path,'final')['body'][:1],['onetwothree'])
        self.assertEqual(read_view(self.path,'original')['body'][:3],['one','two','three'])
        self.assertEqual(self.cli('--structure')['outline'][0]['level'],2); self.assertEqual(self.cli('--structure','--mode','original')['outline'][0]['level'],1)
        self.assertIn('Heading 2',self.cli('--styles')['styles']); self.assertNotIn('Heading 2',self.cli('--styles','--mode','original')['styles'])
        import io
        original=Document(io.BytesIO(project_package(self.path,False)))
        self.assertTrue(original.paragraphs[-1].runs[-1].italic); self.assertIsNone(original.paragraphs[-1].runs[-1].bold)
    def test_inserted_paragraph_mark_and_listnum_previous_cache(self):
        doc=Document(); first=doc.add_paragraph('left'); paragraph_mark(first,'ins','1'); doc.add_paragraph('right')
        para=doc.add_paragraph(); run=para.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); change=revision('numberingChange',None,'2'); change.set(qn('w:original'),'1.'); begin.append(change); run._r.append(begin)
        for tag,text in [('instrText','LISTNUM'),('fldChar','separate'),('t','2.'),('fldChar','end')]:
            el=OxmlElement('w:'+tag)
            if tag=='fldChar': el.set(qn('w:fldCharType'),text)
            else: el.text=text
            para.add_run()._r.append(el)
        doc.save(self.path)
        self.assertEqual(read_view(self.path,'final')['body'],['left','right','2.'])
        self.assertEqual(read_view(self.path,'original')['body'],['leftright','1.'])
    def test_resolve_cli_handles_same_structures_without_destroying_objects(self):
        doc=Document(); p=doc.add_paragraph(); para_revision(p,'del','old','1'); para_revision(p,'ins','new','2')
        anchor=doc.add_paragraph().add_run('anchored'); doc.add_comment([anchor],text='preserved comment',author='Fixture')
        doc.save(self.path)
        save_zip(self.path,{'custom/opaque.bin':b'preserve unrelated package bytes'})
        output=self.path.with_name('resolved.docx')
        result=subprocess.run([sys.executable,str(SCRIPTS/'docx_revisions.py'),'reject-all',str(self.path),'-o',str(output)],text=True,capture_output=True,timeout=20)
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertEqual(Document(output).paragraphs[0].text,'old'); self.assertFalse(read_view(output)['revisions']['has_tracked_changes'])
        self.assertEqual(read_view(self.path)['body'],['new','anchored'])
        self.assertEqual(next(iter(Document(output).comments)).text,'preserved comment')
        with zipfile.ZipFile(output) as package: self.assertEqual(package.read('custom/opaque.bin'),b'preserve unrelated package bytes')

    def test_section_properties_and_control_wrapped_table(self):
        from copy import deepcopy
        doc=Document(); section=doc.sections[0]; section.header.paragraphs[0].text='current header'; section.first_page_header.paragraphs[0].text='previous header'
        properties=section._sectPr; previous=deepcopy(properties)
        references=previous.findall(qn('w:headerReference'))
        default=next(ref for ref in references if ref.get(qn('w:type'))=='default')
        first=next(ref for ref in references if ref.get(qn('w:type'))=='first')
        default.set(qn('r:id'),first.get(qn('r:id'))); previous.remove(first)
        change_property(properties,'sectPrChange',previous,'10')
        table=doc.add_table(1,1); cell=table.cell(0,0); para_revision(cell.paragraphs[0],'ins','inside control','11')
        row=table.rows[0]._tr; row.getparent().remove(row); control=OxmlElement('w:sdt'); content=OxmlElement('w:sdtContent'); content.append(row); control.append(content); table._tbl.append(control)
        doc.save(self.path)
        final=read_view(self.path); original=read_view(self.path,'original')
        self.assertEqual(final['tables'],[[['inside control']]]); self.assertEqual(original['tables'],[[['']]])
        self.assertEqual(final['headers'],['current header','previous header']); self.assertEqual(original['headers'],['previous header'])
        section.header.paragraphs[0].style='Subtitle'; section.first_page_header.paragraphs[0].style='Title'; doc.save(self.path)
        self.assertIn('Subtitle',self.cli('--styles')['styles']); self.assertNotIn('Subtitle',self.cli('--styles','--mode','original')['styles'])
        self.assertIn('Title',self.cli('--styles','--mode','original')['styles'])

    def test_custom_xml_ranges_and_alternate_textbox_are_not_duplicate_prose(self):
        doc=Document(); p=doc.add_paragraph('kept ')
        p._p.append(revision('customXmlDelRangeStart',None,'15')); p.add_run('custom XML text'); p._p.append(revision('customXmlDelRangeEnd',None,'15'))
        mc='http://schemas.openxmlformats.org/markup-compatibility/2006'
        alternate=etree.Element('{'+mc+'}AlternateContent')
        for tag in ('Choice','Fallback'):
            branch=etree.SubElement(alternate,'{'+mc+'}'+tag)
            box=OxmlElement('w:txbxContent'); para=OxmlElement('w:p'); para.append(revision('ins','one text box','16')); box.append(para); branch.append(box)
            props=OxmlElement('w:pPr'); style=OxmlElement('w:pStyle'); style.set(qn('w:val'),'Subtitle' if tag == 'Choice' else 'Title'); props.append(style); para.insert(0,props)
        p.add_run()._r.append(alternate); doc.save(self.path)
        final=read_view(self.path); original=read_view(self.path,'original')
        self.assertEqual(final['body'],['kept custom XML text']); self.assertEqual(original['body'],final['body'])
        self.assertEqual(len(final['textboxes']),1); self.assertEqual(final['textboxes'][0]['blocks'][0]['text'],'one text box')
        styles=self.cli('--styles')['styles']
        self.assertIn('Subtitle',styles); self.assertNotIn('Title',styles)

    def test_paragraph_mark_join_preserves_boundary_bookmarks_and_comment_ranges(self):
        from io import BytesIO
        doc=Document(); first=doc.add_paragraph('before'); paragraph_mark(first,'del','1'); second=doc.add_paragraph('anchored')
        for tag in ('bookmarkStart','commentRangeStart'):
            marker=OxmlElement('w:'+tag); marker.set(qn('w:id'),'2'); first._p.addnext(marker)
        for tag in ('bookmarkEnd','commentRangeEnd'):
            marker=OxmlElement('w:'+tag); marker.set(qn('w:id'),'2'); second._p.append(marker)
        doc.save(self.path); projected=Document(BytesIO(project_package(self.path,True)))
        self.assertEqual(projected.paragraphs[0].text,'beforeanchored')
        for start,end in [('bookmarkStart','bookmarkEnd'),('commentRangeStart','commentRangeEnd')]:
            active=False; text=[]
            for element in projected.element.iter():
                if element.tag==qn('w:'+start): active=True
                elif element.tag==qn('w:'+end): active=False
                elif active and element.tag==qn('w:t'): text.append(element.text)
            self.assertEqual(''.join(text),'anchored')

if __name__=='__main__': unittest.main()
