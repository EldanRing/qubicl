"""CLI save/reopen checks for CSV typing, capacity and lossless publication."""
import csv
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[2] / 'skills/core/xlsx/scripts/csv_to_xlsx.py'


class CsvConversion(unittest.TestCase):
    def test_capacity_errors_preserve_existing_output(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'input.csv', Path(directory) / 'output.xlsx'
            output.write_bytes(b'original')
            command = [sys.executable, str(SCRIPT), str(source), str(output), '--plain']
            cases = [
                ('Value\n' + 'a' * 40000 + 'END\n', [], 'A2'),
                (','.join(['column'] * 16385) + '\n', [], 'XFE1'),
                ('\n' * 1048577, [], 'CSV row 1048577'),
                ('Value\n=' + '1+' * 4096 + '1\n', ['--allow-formulas'], 'A2'),
            ]
            for value, options, location in cases:
                source.write_text(value)
                result = subprocess.run(command + options, capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0, result.stdout)
                self.assertIn(location, json.loads(result.stderr)['error'])
                self.assertEqual(output.read_bytes(), b'original')
                self.assertEqual(list(Path(directory).glob('.qubicl-csv-*')), [])

    def test_cell_and_column_boundaries_save_without_truncation(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'input.csv', Path(directory) / 'output.xlsx'
            text = 'a' * 32761 + '\n日本END'
            row = [text, '=1+2'] + [''] * (16384 - 3) + ['LAST']
            with source.open('w', newline='') as stream:
                csv.writer(stream).writerow(row)
            result = subprocess.run([sys.executable, str(SCRIPT), str(source), str(output), '--plain'], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            workbook = load_workbook(output)
            self.assertEqual(workbook.active.max_column, 16384)
            self.assertEqual(workbook.active['A1'].value, text)
            self.assertEqual(workbook.active['B1'].value, '=1+2')
            self.assertEqual(workbook.active['B1'].data_type, 's')
            self.assertEqual(workbook.active['XFD1'].value, 'LAST')
            workbook.close()

    def test_serialization_failure_cannot_replace_destination(self):
        spec = importlib.util.spec_from_file_location('csv_to_xlsx', SCRIPT)
        converter = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(converter)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / 'output.xlsx'
            output.write_bytes(b'original')

            def fail_after_write(path):
                Path(path).write_bytes(b'partial archive')
                raise OSError('fixture disk write failure')

            workbook = converter.Workbook()
            with patch.object(workbook, 'save', fail_after_write):
                with self.assertRaisesRegex(OSError, 'fixture disk write failure'):
                    converter.save_verified(workbook, output, [])
            self.assertEqual(output.read_bytes(), b'original')
            self.assertEqual(list(Path(directory).glob('.qubicl-csv-*')), [])

    def test_overflow_never_succeeds_or_replaces_existing_output(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'input.csv', Path(directory) / 'output.xlsx'
            output.write_bytes(b'original')
            for value in ('1e309', '-1e309', 'NaN', 'Infinity', '9' * 310):
                source.write_text(f'Value\n{value}\n2\n')
                result = subprocess.run([sys.executable, str(SCRIPT), str(source), str(output), '--column-types', '{"A":"number"}'], capture_output=True, text=True)
                self.assertNotEqual(result.returncode, 0)
                self.assertIn('A2', json.loads(result.stderr)['error'])
                self.assertEqual(output.read_bytes(), b'original')

    def test_bom_nullable_values_and_contextual_strict_errors(self):
        with tempfile.TemporaryDirectory() as directory:
            source, output = Path(directory) / 'input.csv', Path(directory) / 'output.xlsx'
            source.write_text('\ufeffName,Value\none,1\ninner\ufeffname,\nthree,3\n', encoding='utf-8')
            command = [sys.executable, str(SCRIPT), str(source), str(output), '--column-types', '{"B":"number"}']
            result = subprocess.run(command, capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            workbook = load_workbook(output)
            self.assertEqual(list(workbook.active.values), [('Name', 'Value'), ('one', 1), ('inner\ufeffname', None), ('three', 3)])
            workbook.close()
            original = output.read_bytes()
            result = subprocess.run(command + ['--empty-cells', 'reject'], capture_output=True, text=True)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('CSV row 3, column 2 (B3)', json.loads(result.stderr)['error'])
            self.assertEqual(output.read_bytes(), original)


if __name__ == '__main__': unittest.main()
