"""Hidden-slide provenance through real PPTX inspection and LibreOffice export."""
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
spec = importlib.util.spec_from_file_location('renderer', ROOT / 'packages/control/src/artifact-render.py')
renderer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(renderer)


class HiddenSlides(unittest.TestCase):
    def test_source_order_visibility_and_output_mapping(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary); source = root / 'deck.pptx'
            deck = Presentation()
            for number in range(1, 6):
                slide = deck.slides.add_slide(deck.slide_layouts[6])
                slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = f'SOURCE SLIDE {number}'
                if number in (1, 3, 5):slide._element.set('show', '0' if number != 3 else 'false')
            deck.save(source); before = source.read_bytes()
            read = subprocess.run([sys.executable, str(ROOT / 'skills/core/powerpoint/scripts/pptx_read.py'), str(source)],
                                  capture_output=True, text=True, check=True, timeout=20)
            outline = json.loads(read.stdout)
            self.assertEqual((outline['slide_count'], outline['visible_slide_count'], outline['hidden_slide_count']), (5, 2, 3))
            self.assertEqual([slide['hidden'] for slide in outline['slides']], [True, False, True, False, True])
            self.assertIn('SOURCE SLIDE 3', outline['slides'][2]['texts'])
            receipt = renderer.render(source, root / 'render', dpi=36)
            self.assertTrue(receipt['rendered'])
            self.assertEqual(receipt['sourceSha256'], hashlib.sha256(before).hexdigest())
            self.assertEqual(receipt, json.loads(Path(receipt['manifest']).read_text()))
            self.assertEqual(receipt['presentation']['policy'], 'visible_slides_only')
            self.assertEqual([s['renderedPage'] for s in receipt['presentation']['slides']], [None, 1, None, 2, None])
            self.assertEqual([p['sourceSlide']['number'] for p in receipt['pages']], [2, 4])
            pdf = PdfReader(receipt['pdf']['path'])
            for page, number in zip(pdf.pages, (2, 4)):
                self.assertIn(f'SOURCE SLIDE {number}', page.extract_text())
            self.assertEqual(source.read_bytes(), before)
            for slide in deck.slides:slide._element.set('show', '0')
            deck.save(source); hidden = source.read_bytes()
            receipt = renderer.render(source, root / 'render', dpi=36)
            self.assertFalse(receipt['rendered']);self.assertEqual(receipt['reason'], 'no_visible_slides')
            self.assertEqual(receipt['pages'], []);self.assertEqual(receipt['presentation']['hiddenSlideCount'], 5)
            self.assertEqual(source.read_bytes(), hidden)


if __name__ == '__main__':unittest.main()
