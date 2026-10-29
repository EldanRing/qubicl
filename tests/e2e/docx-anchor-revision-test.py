"""Saved-package regressions for tracked row resolution and inline comment ranges."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / 'skills/core/docx/scripts'
sys.path.insert(0, str(SCRIPTS))
from docx_comments import list_comments


class DocxAnchorsAndRows(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.directory = Path(self.temp.name)

    def cli(self, helper, *arguments):
        result = subprocess.run([sys.executable, str(SCRIPTS / helper), *map(str, arguments)],
                                text=True, capture_output=True, timeout=20)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        return json.loads(result.stdout)

    def test_accept_and_reject_deleted_row_after_save_and_reopen(self):
        source = self.directory / 'deleted-row.docx'
        doc = Document()
        table = doc.add_table(3, 2)
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.text = f'row-{row_index}-column-{column_index}'
        deletion = OxmlElement('w:del')
        deletion.set(qn('w:id'), '6')
        deletion.set(qn('w:author'), 'Qubicl Maintainers')
        table.rows[1]._tr.get_or_add_trPr().append(deletion)
        doc.save(source)
        original = source.read_bytes()
        for action, expected in [('accept-all', [0, 2]), ('reject-all', [0, 1, 2])]:
            with self.subTest(action=action):
                output = self.directory / f'{action}.docx'
                receipt = self.cli('docx_revisions.py', action, source, '-o', output)
                self.assertEqual(receipt['resolved'], 1)
                reopened = Document(output)
                self.assertEqual([[cell.text for cell in row.cells] for row in reopened.tables[0].rows],
                                 [[f'row-{index}-column-0', f'row-{index}-column-1'] for index in expected])
                self.assertEqual(reopened.element.xpath('.//w:del'), [])
                self.assertEqual(self.cli('docx_revisions.py', 'list', output)['revisions'], [])
        self.assertEqual(source.read_bytes(), original)

    def test_comment_targets_and_inline_controls_survive_native_and_xml_roundtrips(self):
        for xml in (False, True):
            with self.subTest(xml=xml):
                doc = Document()
                # Exact reported case: A + tab + Target, with the target in the
                # same run, where raw w:t lengths formerly anchored only arget.
                exact = doc.add_paragraph('A\tTarget')
                # The less common run controls must count exactly like Run.text.
                complex_para = doc.add_paragraph()
                run = complex_para.add_run('prefix')
                for tag in ('ptab', 'noBreakHyphen', 'cr'):
                    run._r.append(OxmlElement(f'w:{tag}'))
                run.add_text('Alpha\tBeta\nGamma')
                run.add_break()
                run.add_text('end suffix')
                # Non-text breaks remain in XML and do not shift the target.
                page = doc.add_paragraph()
                page_run = page.add_run('page-before')
                page_break = OxmlElement('w:br'); page_break.set(qn('w:type'), 'page')
                page_run._r.append(page_break)
                page_run.add_text('PageTarget after')
                source = self.directory / f'comments-{xml}.docx'
                doc.save(source)
                expected_text = [para.text for para in doc.paragraphs]
                expected_controls = {tag: len(doc.element.xpath(f'.//w:{tag}')) for tag in ('tab', 'ptab', 'noBreakHyphen', 'cr', 'br')}
                targets = ['Target', 'Alpha\tBeta\nGamma\nend', 'PageTarget']
                for index, target in enumerate(targets):
                    output = self.directory / f'comments-{xml}-{index}.docx'
                    self.cli('docx_comments.py', 'add', source, '--target', target,
                             '--text', 'Review\twith\nlines', '--author', 'Qubicl Maintainers',
                             '-o', output, *(['--xml'] if xml else []))
                    source = output
                reopened = Document(source)
                comments = list_comments(reopened)
                self.assertEqual([comment['anchored_text'] for comment in comments], targets)
                self.assertEqual([comment['text'] for comment in comments], ['Review\twith\nlines'] * 3)
                self.assertEqual([para.text for para in reopened.paragraphs], expected_text)
                self.assertEqual({tag: len(reopened.element.xpath(f'.//w:{tag}')) for tag in expected_controls}, expected_controls)
                self.assertEqual(exact.text, 'A\tTarget')


if __name__ == '__main__':
    unittest.main()
