"""Real package regressions for independent PowerPoint chart-slide duplication."""
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile
from lxml import etree
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / 'skills/core/powerpoint/scripts/pptx_edit.py'
spec = importlib.util.spec_from_file_location('pptx_chart_edit', SCRIPT)
edit = importlib.util.module_from_spec(spec); spec.loader.exec_module(edit)
STYLE_REL = 'http://schemas.microsoft.com/office/2011/relationships/chartStyle'
COLOR_REL = 'http://schemas.microsoft.com/office/2011/relationships/chartColorStyle'

def charts(slide):
    return [shape.chart for shape in edit.iter_shapes(slide.shapes) if shape.has_chart]

def related(chart, reltype):
    return next(rel.target_part for rel in chart.part.rels.values() if rel.reltype == reltype)

def validate_package(path):
    from posixpath import dirname, join, normpath
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)): raise AssertionError('Duplicate ZIP part name')
        for name in names:
            if not name.endswith('.rels'): continue
            root = etree.fromstring(archive.read(name))
            ids = [rel.get('Id') for rel in root]
            if len(ids) != len(set(ids)): raise AssertionError('Duplicate relationship ID')
            directory = dirname(dirname(name)) if '/_rels/' in name else ''
            for rel in root:
                if rel.get('TargetMode') == 'External': continue
                target = rel.get('Target')
                resolved = target.lstrip('/') if target.startswith('/') else normpath(join(directory, target))
                if resolved not in names: raise AssertionError(f'Missing relationship target {resolved}')

class ChartCopies(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_nested_charts_clone_workbooks_styles_notes_and_relationships(self):
        prs = Presentation(); source = prs.slides.add_slide(prs.slide_layouts[6])
        target = prs.slides.add_slide(prs.slide_layouts[6])
        data = CategoryChartData(); data.categories = ['Q1','Q2']; data.add_series('North',[1,2])
        category = source.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,0, Inches(4), Inches(3),data).chart
        nested = source.shapes.add_group_shape().shapes.add_group_shape()
        xy = XyChartData(); series = xy.add_series('XY'); series.add_data_point(1,2); series.add_data_point(3,4)
        nested.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER, Inches(4),0, Inches(4), Inches(3),xy)
        source.background.fill.solid(); source.background.fill.fore_color.rgb = RGBColor.from_string('AA2244')
        para = source.notes_slide.notes_text_frame.paragraphs[0]
        para.text = 'Formatted note'; para.runs[0].font.bold = True
        para.runs[0].hyperlink.address = 'https://example.test/notes'
        package = prs.part.package
        style = Part(PackURI('/ppt/charts/style1.xml'), 'application/vnd.ms-office.chartstyle+xml', package,
                     b'<cs:chartStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle" id="201"/>')
        colors = Part(PackURI('/ppt/charts/colors1.xml'), 'application/vnd.ms-office.chartcolorstyle+xml', package,
                      b'<cs:colorStyle xmlns:cs="http://schemas.microsoft.com/office/drawing/2012/chartStyle" meth="cycle" id="10"/>')
        category.part.relate_to(style, STYLE_REL); category.part.relate_to(colors,COLOR_REL)
        # An opaque XML owned part has relationships too; preserve rIds and cycles.
        style.relate_to(colors,COLOR_REL); colors.relate_to(style,STYLE_REL)
        original_style = style.blob
        image = io.BytesIO(); Image.new('RGB', (8, 8), 'blue').save(image, format='PNG'); image.seek(0)
        picture = source.shapes.add_picture(image, 0, Inches(4), width=Inches(1))
        image_rid = picture._element.blipFill.blip.get(edit.R_EMBED + 'embed')
        original_image = source.part.rels[image_rid].target_part
        self_link = source.part.relate_to(source.part,RT.SLIDE)
        other_link = source.part.relate_to(target.part,RT.SLIDE)
        duplicate = prs.slides[edit.duplicate_slide(prs,0)]
        copy_charts = charts(duplicate)
        copied_image = duplicate.part.rels[image_rid].target_part
        self.assertIsNot(copied_image, original_image)
        self.assertEqual(copied_image.blob, original_image.blob)
        self.assertEqual(etree.tostring(source._element), etree.tostring(duplicate._element))
        self.assertEqual(len(copy_charts),2)
        for left,right in zip(charts(source),copy_charts):
            self.assertIsNot(left.part,right.part)
            self.assertIsNot(left.part.chart_workbook.xlsx_part,right.part.chart_workbook.xlsx_part)
            self.assertEqual(left.part.chart_workbook.xlsx_part.blob,right.part.chart_workbook.xlsx_part.blob)
        self.assertIs(duplicate.part.rels[self_link].target_part,duplicate.part)
        self.assertIs(duplicate.part.rels[other_link].target_part,target.part)
        self.assertIsNot(source.notes_slide.part,duplicate.notes_slide.part)
        self.assertIs(duplicate.notes_slide.part.part_related_by(RT.SLIDE),duplicate.part)
        self.assertEqual(etree.tostring(source.notes_slide._element),etree.tostring(duplicate.notes_slide._element))
        copied_style = related(copy_charts[0],STYLE_REL)
        copied_colors = related(copy_charts[0],COLOR_REL)
        self.assertIsNot(copied_style,style); self.assertIsNot(copied_colors,colors)
        self.assertIs(next(iter(copied_colors.rels.values())).target_part,copied_style)
        copied_style._blob = copied_style.blob.replace(b'201',b'202')
        replacement = CategoryChartData(); replacement.categories=['Q1','Q2']; replacement.add_series('North',[99,100])
        edit.replace_chart_data(copy_charts[0],replacement)
        replacement_xy = XyChartData(); series=replacement_xy.add_series('XY'); series.add_data_point(5,6)
        copy_charts[1].replace_data(replacement_xy)
        # A second copy exercises package-wide unique part naming.
        second = edit.duplicate_slide(prs,0)
        path=self.root/'copies.pptx'; prs.save(path); validate_package(path)
        reopened=Presentation(path)
        self.assertEqual(charts(reopened.slides[0])[0].series[0].values,(1.0,2.0))
        self.assertEqual(charts(reopened.slides[2])[0].series[0].values,(99.0,100.0))
        self.assertEqual(charts(reopened.slides[second])[0].series[0].values,(1.0,2.0))
        self.assertEqual(charts(reopened.slides[0])[1].series[0].values,(2.0,4.0))
        self.assertEqual(charts(reopened.slides[2])[1].series[0].values,(6.0,))
        self.assertEqual(related(charts(reopened.slides[0])[0],STYLE_REL).blob,original_style)
        self.assertIn(b'202',related(charts(reopened.slides[2])[0],STYLE_REL).blob)
        self.assertEqual(str(reopened.slides[2].background.fill.fore_color.rgb),'AA2244')
        self.assertTrue(reopened.slides[2].notes_slide.notes_text_frame.paragraphs[0].runs[0].font.bold)
        for slide_index, expected in [(0,1),(2,99),(second,1)]:
            workbook=load_workbook(io.BytesIO(charts(reopened.slides[slide_index])[0].part.chart_workbook.xlsx_part.blob))
            self.assertEqual(workbook.active['B2'].value,expected)

    def test_chart_data_cli_updates_grouped_copy_without_touching_source(self):
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); group=slide.shapes.add_group_shape()
        data=CategoryChartData(); data.categories=['Q1']; data.add_series('North',[1])
        group.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,0, Inches(4), Inches(3),data)
        edit.duplicate_slide(prs,0)
        path=self.root/'grouped.pptx'; prs.save(path)
        spec_path=self.root/'change.json'; spec_path.write_text(json.dumps({'slide':1,'chart':0,'ops':[{'op':'update_series','name':'North','values':[99]}]}))
        subprocess.run([sys.executable,str(SCRIPT),str(path),'--chart-data',str(spec_path)],check=True,capture_output=True,timeout=20)
        reopened=Presentation(path); validate_package(path)
        self.assertEqual(charts(reopened.slides[0])[0].series[0].values,(1.0,))
        self.assertEqual(charts(reopened.slides[1])[0].series[0].values,(99.0,))

    def test_linked_workbook_is_preserved_until_explicit_data_replacement(self):
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); data=CategoryChartData(); data.categories=['A']; data.add_series('S',[1])
        chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,0, Inches(4), Inches(3),data).chart
        external=chart._chartSpace.find(qn('c:externalData'))
        original_rid=external.get(edit.R_EMBED+'id'); chart.part.drop_rel(original_rid)
        rid=chart.part.relate_to('https://example.test/book.xlsx',RT.PACKAGE,is_external=True); external.set(edit.R_EMBED+'id',rid)
        duplicate=prs.slides[edit.duplicate_slide(prs,0)]; copied=charts(duplicate)[0]
        replacement=CategoryChartData(); replacement.categories=['A']; replacement.add_series('S',[7]); edit.replace_chart_data(copied,replacement)
        self.assertTrue(chart.part.rels[rid].is_external)
        self.assertEqual(chart.part.rels[rid].target_ref,'https://example.test/book.xlsx')
        self.assertIsNotNone(copied.part.chart_workbook.xlsx_part)
        path=self.root/'linked.pptx'; prs.save(path); validate_package(path)
        reopened=Presentation(path)
        self.assertEqual(charts(reopened.slides[0])[0].series[0].values,(1.0,))
        self.assertEqual(charts(reopened.slides[1])[0].series[0].values,(7.0,))

if __name__ == '__main__': unittest.main()
