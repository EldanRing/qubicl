#!/usr/bin/env python3
"""Small real-file regressions for audited flattening and artifact rendering."""
import hashlib
import importlib.util
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[2]
HELPER = ROOT / 'packages/control/src/artifact-render.py'
os.environ['QUBICL_ARTIFACT_RENDERER'] = str(HELPER)
spec = importlib.util.spec_from_file_location('artifact_renderer', HELPER)
renderer = importlib.util.module_from_spec(spec)
spec.loader.exec_module(renderer)


def run(script, *args):
    result = subprocess.run([sys.executable, str(ROOT / script), *map(str, args)], text=True, capture_output=True, check=True, timeout=120)
    return json.loads(result.stdout)


with tempfile.TemporaryDirectory(prefix='qubicl-artifact-regression-') as directory:
    tmp = Path(directory)
    from reportlab.pdfgen import canvas
    from pypdf import PdfReader
    form = tmp / 'form.pdf'
    drawing = canvas.Canvas(str(form))
    drawing.drawString(20, 780, 'Form fixture')
    drawing.acroForm.textfield(name='change', x=20, y=700, width=150, height=25, value='old')
    drawing.acroForm.textfield(name='retain', x=20, y=650, width=150, height=25, value='KEEP')
    drawing.acroForm.checkbox(name='agree', x=20, y=600, checked=True)
    drawing.showPage()
    drawing.save()
    fields = tmp / 'fields.json'
    fields.write_text(json.dumps({'change': 'NEW'}))
    flattened = tmp / 'flat.pdf'
    receipt = run('skills/core/pdf/scripts/pdf_fill_form.py', form, '--fields-json', fields, '--flatten', '-o', flattened)
    pdf = PdfReader(flattened)
    assert receipt['flattened'] is True and not pdf.get_fields()
    assert not any(a.get_object().get('/Subtype') == '/Widget' for page in pdf.pages for a in page.get('/Annots', []))
    assert 'NEW' in pdf.pages[0].extract_text() and 'KEEP' in pdf.pages[0].extract_text()
    assert pdf.pages[0]['/Resources']['/XObject'], 'Flattened appearances must be merged into page resources'
    interactive = tmp / 'interactive.pdf'
    run('skills/core/pdf/scripts/pdf_fill_form.py', form, '--fields-json', fields, '-o', interactive)
    assert PdfReader(interactive).get_fields()['change']['/V'] == 'NEW'
    print('PASS: filled text, untouched text and checkbox flattened; interactive mode preserved', flush=True)
    # Flattened checkbox appearance must actually survive rasterization.
    from PIL import Image, ImageChops
    for source, name in [(interactive, 'interactive'), (flattened, 'flat')]:
        subprocess.run(['pdftoppm', '-singlefile', '-png', '-r', '72', str(source), str(tmp / name)], check=True, capture_output=True)
    image = Image.open(tmp / 'interactive.png').convert('RGB')
    flat_image = Image.open(tmp / 'flat.png').convert('RGB')
    height = image.height
    region = (18, height - 625, 45, height - 598)
    difference = ImageChops.difference(image.crop(region), flat_image.crop(region))
    assert difference.getbbox() is None, 'Checkbox appearance changed when flattened'
    # Nonterminal field groups should not prevent flattening their leaf fields.
    from pypdf import PdfWriter
    from pypdf.generic import ArrayObject, DictionaryObject, NameObject, TextStringObject
    grouped = PdfWriter(clone_from=form)
    leaf = grouped.get_fields()['retain'].indirect_reference
    parent = grouped._add_object(DictionaryObject({NameObject('/T'): TextStringObject('group'), NameObject('/Kids'): ArrayObject([leaf])}))
    leaf.get_object()[NameObject('/Parent')] = parent
    root_fields = grouped._root_object['/AcroForm']['/Fields']
    root_fields[root_fields.index(leaf)] = parent
    grouped_path = tmp / 'grouped.pdf'
    grouped.write(grouped_path)
    grouped_output = tmp / 'grouped-flat.pdf'
    run('skills/core/pdf/scripts/pdf_fill_form.py', grouped_path, '--fields-json', fields, '--flatten', '-o', grouped_output)
    assert 'KEEP' in PdfReader(grouped_output).pages[0].extract_text()
    from pypdf.generic import NumberObject
    hidden = PdfWriter(clone_from=form)
    hidden.get_fields()['retain'].indirect_reference.get_object()[NameObject('/F')] = NumberObject(2)
    hidden_path = tmp / 'hidden.pdf'
    hidden.write(hidden_path)
    hidden_flat = tmp / 'hidden-flat.pdf'
    run('skills/core/pdf/scripts/pdf_fill_form.py', hidden_path, '--fields-json', fields, '--flatten', '-o', hidden_flat)
    assert 'KEEP' not in PdfReader(hidden_flat).pages[0].extract_text()
    unsupported = PdfWriter(clone_from=form)
    unsupported.get_fields()['retain'].indirect_reference.get_object()[NameObject('/FT')] = NameObject('/Sig')
    unsupported_path = tmp / 'unsupported.pdf'
    unsupported.write(unsupported_path)
    previous = flattened.read_bytes()
    failure = subprocess.run([sys.executable, str(ROOT / 'skills/core/pdf/scripts/pdf_fill_form.py'), str(unsupported_path), '--fields-json', str(fields), '--flatten', '-o', str(flattened)], capture_output=True, text=True)
    assert failure.returncode and flattened.read_bytes() == previous
    assert 'Cannot safely flatten unsupported field' in failure.stderr
    if '--pdf-only' in sys.argv:
        print('PASS: PDF-only flatten appearance and grouped-field checks', flush=True)
        raise SystemExit(0)


    output = tmp / 'render'
    output.mkdir()
    unrelated = output / 'slide-9.png'
    unrelated.write_bytes(b'keep unrelated old preview')
    from pptx import Presentation
    from pptx.util import Inches
    from docx import Document
    from openpyxl import Workbook, load_workbook
    deck = Presentation()
    for title in ['FIRST', 'SECOND']:
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = title
    pptx = tmp / 'deck.pptx'
    deck.save(pptx)
    first = run('skills/core/powerpoint/scripts/pptx_render.py', pptx, '--outdir', output, '--dpi', '36')
    assert len(first['files']) == 2
    short = Presentation()
    short.slides.add_slide(short.slide_layouts[6])
    short.save(pptx)
    second = run('skills/core/powerpoint/scripts/pptx_render.py', pptx, '--outdir', output, '--dpi', '36')
    assert len(second['files']) == 1 and set(first['files']).isdisjoint(second['files'])
    assert unrelated.read_bytes() == b'keep unrelated old preview'
    assert second['sourceSha256'] == hashlib.sha256(pptx.read_bytes()).hexdigest()
    assert json.loads(Path(second['manifest']).read_text()) == second
    assert all(hashlib.sha256(Path(page['path']).read_bytes()).hexdigest() == page['sha256'] for page in second['pages'])
    print('PASS: fresh PPTX page manifests, shorter rerender and unrelated file preservation', flush=True)

    docx = tmp / 'document.docx'
    document = Document()
    document.add_paragraph('DOCX fixture')
    document.save(docx)
    word = renderer.render(docx, output, dpi=36)
    assert word['rendered'] and len(word['files']) == 1
    direct = renderer.render(flattened, output, dpi=36)
    assert direct['rendered'] and len(direct['files']) == 1
    workbook = Workbook()
    workbook.active['A1'] = '=1+1'
    workbook.active['A2'] = '=1/0'
    workbook.active['A3'] = '=""'
    xlsx = tmp / 'workbook.xlsx'
    workbook.save(xlsx)
    result = run('skills/core/xlsx/scripts/xlsx_recalc.py', xlsx, '--out', tmp / 'recalculated.xlsx')
    assert result['recalculated'] and result['validation']['formulaErrors'] == [{'sheet': 'Sheet', 'cell': 'A2', 'error': '#DIV/0!'}]
    assert result['validation']['noFormulaErrors'] is False
    assert result['validation']['missingFormulaResults'] == [] and result['with_cached_values'] == 3
    assert result['validation']['changedPackageParts'] and result['validation']['fidelityVerified'] is False
    assert load_workbook(xlsx)['Sheet']['A1'].value == '=1+1'
    exported = tmp / 'results.csv'
    csv_receipt = run('skills/core/xlsx/scripts/xlsx_to_csv.py', tmp / 'recalculated.xlsx', exported, '--data-only')
    assert csv_receipt['valueComplete'] and csv_receipt['missingFormulaResults'] == []
    assert exported.read_text().splitlines() == ['2', '#DIV/0!', '""']
    prior = exported.read_bytes()
    failure = subprocess.run([sys.executable, str(ROOT / 'skills/core/xlsx/scripts/xlsx_to_csv.py'), str(xlsx), str(exported), '--data-only'], capture_output=True, text=True)
    assert failure.returncode != 0 and exported.read_bytes() == prior
    assert 'Missing cached formula results' in failure.stderr

    sheet = renderer.render(xlsx, output, dpi=36)
    assert sheet['validation']['recalculated'] and not sheet['validation']['noFormulaErrors']
    assert sheet['files']
    print('PASS: DOCX/PDF/XLSX render manifests and recalculation error/fidelity distinctions', flush=True)
