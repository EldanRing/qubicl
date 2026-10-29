"""Save/reopen regressions for complete XY/bubble reads and plot-safe edits."""
import copy
from datetime import datetime
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData, BubbleChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT/'skills/core/powerpoint/scripts'
sys.path.insert(0,str(SCRIPTS))
import pptx_read as read
import pptx_edit as edit


class XYCharts(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def chart(self, prs, kind, points, name='Series'):
        data=BubbleChartData() if kind == 'bubble' else XyChartData()
        series=data.add_series(name)
        for point in points: series.add_data_point(*point)
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        return slide.shapes.add_chart(XL_CHART_TYPE.BUBBLE if kind == 'bubble' else XL_CHART_TYPE.XY_SCATTER,
                                      Inches(1), Inches(1), Inches(6), Inches(4), data).chart

    def save_reopen(self, prs, name='deck'):
        path=self.root/f'{name}.pptx'; prs.save(path)
        return Presentation(path), path

    def change(self, path, spec):
        spec_path=self.root/'change.json'; spec_path.write_text(json.dumps(spec))
        result=subprocess.run([sys.executable,str(SCRIPTS/'pptx_edit.py'),str(path),'--chart-data',str(spec_path)],capture_output=True,text=True,timeout=20)
        self.assertEqual(result.returncode,0,result.stderr)
        return Presentation(path)

    def test_distinct_x_values_and_sparse_literals_are_visible_in_outline(self):
        prs=Presentation(); self.chart(prs,'xy',[(1,4),(2,5)]); second=self.chart(prs,'xy',[(10,4),(20,5)])
        # Numeric literals are valid OOXML too; missing idx=1 must stay aligned.
        source=second.series[0]._element.find(qn('c:xVal'))
        source.clear()
        from lxml import etree
        literal=etree.SubElement(source,qn('c:numLit')); etree.SubElement(literal,qn('c:ptCount'),val='3')
        for index,value in ((0,10),(2,30)):
            point=etree.SubElement(literal,qn('c:pt'),idx=str(index)); etree.SubElement(point,qn('c:v')).text=str(value)
        reopened,path=self.save_reopen(prs)
        first_record=read.chart_info(reopened.slides[0].shapes[0].chart)
        second_record=read.chart_info(reopened.slides[1].shapes[0].chart)
        self.assertEqual(first_record['series'][0]['values'],[4,5])
        self.assertEqual(first_record['series'][0]['xValues'],[1,2])
        self.assertEqual(second_record['series'][0]['xValues'],[10,None,30])
        result=subprocess.run([sys.executable,str(SCRIPTS/'pptx_read.py'),str(path)],capture_output=True,text=True,check=True,timeout=20)
        self.assertEqual(json.loads(result.stdout)['slides'][1]['charts'][0]['plots'][0]['kind'],'xy')

    def test_cacheless_xy_sources_read_embedded_workbook(self):
        prs=Presentation(); chart=self.chart(prs,'bubble',[(1,2,3),(4,5,6)])
        workbook=load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob))
        workbook.active['A2']=datetime(2024,1,2)
        blob=io.BytesIO(); workbook.save(blob); chart.part.chart_workbook.update_from_xlsx_blob(blob.getvalue())
        for cache in chart.series[0]._element.xpath('.//c:numCache | .//c:tx/c:strRef/c:strCache'):
            cache.getparent().remove(cache)
        reopened,_=self.save_reopen(prs)
        series=read.chart_info(reopened.slides[0].shapes[0].chart)['series'][0]
        self.assertEqual(series['xValues'],[to_excel(datetime(2024,1,2)),4]); self.assertEqual(series['yValues'],[2,5])
        self.assertEqual(series['bubbleSizes'],[3,6]); self.assertEqual(series['name'],'Series')

    def test_scatter_full_and_incremental_edits_preserve_x_y_and_workbook(self):
        prs=Presentation(); self.chart(prs,'xy',[(1,4),(2,5)])
        _,path=self.save_reopen(prs)
        prs=self.change(path,{'series':[{'name':'New','xValues':[10,20],'yValues':[30,40]}]})
        prs=self.change(path,{'ops':[{'op':'update_series','name':'New','values':[31,41]},
                                  {'op':'add_series','name':'Added','xValues':[7],'yValues':[8]}]})
        chart=prs.slides[0].shapes[0].chart; info=read.chart_info(chart)
        self.assertEqual(info['series'][0]['xValues'],[10,20]); self.assertEqual(info['series'][0]['yValues'],[31,41])
        self.assertEqual(info['series'][1]['xValues'],[7]); self.assertEqual(info['series'][1]['yValues'],[8])
        workbook=load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob))
        self.assertEqual(workbook.active['A2'].value,10); self.assertEqual(workbook.active['B2'].value,31)

    def test_bubble_edits_preserve_sizes_and_can_empty_then_repopulate_plot(self):
        prs=Presentation(); self.chart(prs,'bubble',[(1,2,3),(4,5,6)])
        _,path=self.save_reopen(prs)
        prs=self.change(path,{'ops':[{'op':'update_series','name':'Series','xValues':[7,8],'values':[9,10]}]})
        info=read.chart_info(prs.slides[0].shapes[0].chart)
        self.assertEqual(info['series'][0]['bubbleSizes'],[3,6]); self.assertEqual(info['series'][0]['xValues'],[7,8])
        prs=self.change(path,{'series':{'Bubbles':{'xValues':[2],'yValues':[3],'bubbleSizes':[4]}}})
        self.assertEqual(read.chart_info(prs.slides[0].shapes[0].chart)['series'][0]['bubbleSizes'],[4])
        prs=self.change(path,{'ops':[{'op':'remove_series','name':'Bubbles'}]})
        self.assertEqual(len(prs.slides[0].shapes[0].chart.plots),1)
        self.assertEqual(read.chart_info(prs.slides[0].shapes[0].chart)['series'],[])
        prs=self.change(path,{'ops':[{'op':'add_series','name':'Again','xValues':[3],'yValues':[4],'bubbleSizes':[5]}]})
        self.assertEqual(read.chart_info(prs.slides[0].shapes[0].chart)['series'][0]['bubbleSizes'],[5])

    def mixed(self):
        prs=Presentation(); slide=prs.slides.add_slide(prs.slide_layouts[6]); data=CategoryChartData()
        data.categories=['A','B']; data.add_series('Bars',[12,14])
        chart=slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED,0,0, Inches(6),Inches(4),data).chart
        scatter=self.chart(prs,'xy',[(1,5),(2,6)],'Scatter')
        plot=copy.deepcopy(scatter.plots[0]._element)
        for formula in plot.xpath('.//c:f'): formula.text="'Scatter'!"+formula.text.rsplit('!',1)[1]
        chart._chartSpace.plotArea.insert(1,plot)
        workbook=load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob)); workbook.active['B2']='=10+2'
        source=load_workbook(io.BytesIO(scatter.part.chart_workbook.xlsx_part.blob)).active
        dest=workbook.create_sheet('Scatter')
        for row in source:
            for cell in row: dest.cell(cell.row,cell.column,cell.value)
        blob=io.BytesIO(); workbook.save(blob); chart.part.chart_workbook.update_from_xlsx_blob(blob.getvalue())
        return prs

    def test_mixed_plot_edit_keeps_other_plot_type_data_and_workbook_formulas(self):
        prs,path=self.save_reopen(self.mixed())
        info=read.chart_info(prs.slides[0].shapes[0].chart)
        self.assertEqual([plot['kind'] for plot in info['plots']],['category','xy'])
        self.assertEqual(info['categories'],['A','B']); self.assertEqual(info['series'][1]['plot'],1)
        prs=self.change(path,{'plot':1,'ops':[{'op':'update_series','name':'Scatter','xValues':[10,20],'yValues':[50,60]}]})
        chart=prs.slides[0].shapes[0].chart; info=read.chart_info(chart)
        self.assertEqual(info['plots'][0]['series'][0]['values'],[12,14])
        self.assertEqual(info['plots'][1]['series'][0]['xValues'],[10,20])
        self.assertEqual(info['plots'][1]['series'][0]['yValues'],[50,60])
        self.assertIn('COLUMN_CLUSTERED',info['plots'][0]['type']); self.assertIn('XY_SCATTER',info['plots'][1]['type'])
        workbook=load_workbook(io.BytesIO(chart.part.chart_workbook.xlsx_part.blob))
        self.assertEqual(workbook['Sheet1']['B2'].value,'=10+2')
        self.assertEqual(workbook['QubiclPlot2']['A2'].value,10)
        prs=self.change(path,{'plot':0,'ops':[{'op':'update_series','name':'Bars','values':[20,30]},{'op':'add_series','name':'More bars','values':[40,50]}]})
        info=read.chart_info(prs.slides[0].shapes[0].chart)
        self.assertEqual(info['plots'][0]['series'][0]['values'],[20,30]); self.assertEqual(info['plots'][1]['series'][0]['xValues'],[10,20])
        ids=[series.idx.val for series in prs.slides[0].shapes[0].chart._chartSpace.plotArea.sers]
        self.assertEqual(len(ids),len(set(ids)))


if __name__ == '__main__': unittest.main()
