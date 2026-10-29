"""Saved-package worksheet-copy regression for independent drawable/data objects."""
import base64
from io import BytesIO
import json
from pathlib import Path
import posixpath
import subprocess
import sys
import tempfile
import unittest
import zipfile
from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableFormula, TableStyleInfo

ROOT = Path(__file__).resolve().parents[2]
SCRIPTS = ROOT / 'skills/core/xlsx/scripts'
sys.path.insert(0, str(SCRIPTS))
from xlsx_edit import copy_sheet
from xlsx_restructure_dependents import _walk_objects
from openpyxl.chart.data_source import NumRef, StrRef, NumData, NumVal

PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M/wHwAF/gL+X2NDWQAAAABJRU5ErkJggg==')


def refs(chart):
    return [obj for obj in _walk_objects(chart) if isinstance(obj, (NumRef, StrRef))]


class SheetCopy(unittest.TestCase):
    def test_repeated_cli_copy_preserves_objects_and_mutation_leaves_source_independent(self):
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            source, output, changed = (directory / name for name in ('source.xlsx', 'copied.xlsx', 'changed.xlsx'))
            wb = Workbook(); ws = wb.active; ws.title = 'Source Data'
            wb.create_sheet('Other')['A1'] = 9
            for row in [['Label', 'Value', 'Bang!'], ['One', 10, 3], ['Two', 20, 4]]: ws.append(row)
            table = Table(displayName='Sales', ref='A1:C3')
            table._initialise_columns()
            for column, name in zip(table.tableColumns, ['Label', 'Value', 'Bang!']): column.name = name
            table.tableColumns[2].calculatedColumnFormula = TableFormula(attr_text='Sales[Value]*2')
            table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium9', showRowStripes=True)
            ws.add_table(table)
            wb.defined_names.add(DefinedName('Sales_Copy', attr_text="'Other'!$A$1"))
            wb.defined_names.add(DefinedName('WorkbookSource', attr_text="'Source Data'!$B$2"))
            ws.defined_names.add(DefinedName('LocalValues', attr_text="'Source Data'!$B$2:$B$3"))
            ws['E1'] = "=SUM(Sales[Value])+'Source Data'!B2+Other!A1"
            ws['E2'] = "='[other.xlsx]Source Data'!B2"
            ws['E3'] = '=IF(TRUE,"Sales[Value] Source Data!B2",SUM(Sales[Bang!]))'
            ws['E4'] = '=SUM(LocalValues)+WorkbookSource'
            ws['E5'] = "=SUM('Source Data:Other'!B2)"
            ws['F1'] = ArrayFormula(ref='F1:F2', text="='Source Data'!B2:B3*2")
            ws['G1'] = DataTableFormula(ref='G1:H2', r1='B2', r2='B3', dt2D=True)
            ws['E6'].hyperlink = "#'Source Data'!B2"
            ws['E7'].hyperlink = 'https://example.test/Source%20Data'
            validation = DataValidation(type='custom', formula1="'Source Data'!B2>0")
            validation.add('B2:B3'); ws.add_data_validation(validation)
            ws.conditional_formatting.add('B2:B3', FormulaRule(formula=["'Source Data'!B2>5"], fill=PatternFill('solid', fgColor='FFFF00')))
            ws.auto_filter.ref = 'A1:C3'; ws.auto_filter.add_filter_column(0, ['One'])
            ws.freeze_panes = 'B2'; ws.sheet_view.zoomScale = 85; ws.sheet_view.showGridLines = False
            ws.print_area = 'A1:H20'; ws.print_title_rows = '1:1'; ws.print_title_cols = 'A:A'
            ws.oddHeader.center.text = 'Independent copy'; ws.row_breaks.append(Break(id=10))
            ws.protection.sheet = True; ws.row_dimensions[2].height = 28; ws.column_dimensions['B'].width = 21
            chart = BarChart(); chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=3), titles_from_data=True)
            chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=3))
            line = LineChart(); line.add_data(Reference(ws, min_col=3, min_row=1, max_row=3), titles_from_data=True)
            line.ser[0].val.numRef.f = "'[external.xlsx]Other'!$A$2:$A$3"
            line.ser[0].val.numRef.numCache = NumData(ptCount=2, pt=[NumVal(idx=0, v=77), NumVal(idx=1, v=88)])
            chart += line; ws.add_chart(chart, 'J2')
            image = Image(BytesIO(PNG)); image.width = 24; image.height = 24; ws.add_image(image, 'J15')
            wb.save(source); source_bytes = source.read_bytes()
            result = subprocess.run([sys.executable, str(SCRIPTS / 'xlsx_edit.py'), str(source),
                                     '--copy-sheet', 'Source Data:Copy One', '--copy-sheet', 'Source Data:Copy Two', '--out', str(output)],
                                    capture_output=True, text=True, timeout=20)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(json.loads(result.stdout)['ok'])
            self.assertEqual(source.read_bytes(), source_bytes)
            loaded = load_workbook(output)
            original, first, second = (loaded[name] for name in ('Source Data', 'Copy One', 'Copy Two'))
            self.assertEqual(list(original.tables), ['Sales'])
            self.assertEqual(list(first.tables), ['Sales_Copy2']); self.assertEqual(list(second.tables), ['Sales_Copy3'])
            self.assertEqual(original['E1'].value, "=SUM(Sales[Value])+'Source Data'!B2+Other!A1")
            self.assertEqual(first['E1'].value, "=SUM(Sales_Copy2[Value])+'Copy One'!B2+Other!A1")
            self.assertEqual(first['E2'].value, original['E2'].value)
            self.assertEqual(first['E3'].value, '=IF(TRUE,"Sales[Value] Source Data!B2",SUM(Sales_Copy2[Bang!]))')
            self.assertEqual(first['E4'].value, original['E4'].value)
            self.assertEqual(first['E5'].value, original['E5'].value, '3D intervals retain their original sheet membership')
            self.assertEqual(first['F1'].value.text, "='Copy One'!B2:B3*2")
            self.assertEqual(original['F1'].value.text, "='Source Data'!B2:B3*2")
            self.assertEqual(first['G1'].value.r1, 'B2'); self.assertTrue(first['G1'].value.dt2D)
            self.assertEqual(first['E6'].hyperlink.target, "#'Copy One'!B2")
            self.assertEqual(first['E7'].hyperlink.target, original['E7'].hyperlink.target)
            self.assertEqual(first.tables['Sales_Copy2'].tableColumns[2].calculatedColumnFormula.attr_text, 'Sales_Copy2[Value]*2')
            self.assertEqual(first.defined_names['LocalValues'].attr_text, "'Copy One'!$B$2:$B$3")
            self.assertEqual(loaded.defined_names['WorkbookSource'].attr_text, "'Source Data'!$B$2")
            self.assertEqual(first.data_validations.dataValidation[0].formula1, "'Copy One'!B2>0")
            self.assertEqual(next(iter(first.conditional_formatting._cf_rules.values()))[0].formula, ["'Copy One'!B2>5"])
            for sheet in (original, first, second):
                self.assertEqual(len(sheet._charts), 1); self.assertEqual(len(sheet._charts[0]._charts), 2)
                self.assertEqual(len(sheet._images), 1)
                self.assertEqual(sheet._charts[0].anchor._from.col, 9); self.assertEqual(sheet._images[0].anchor._from.row, 14)
                self.assertEqual(sheet.auto_filter.ref, 'A1:C3'); self.assertEqual(sheet.auto_filter.filterColumn[0].filters.filter, ['One'])
                self.assertEqual(sheet.freeze_panes, 'B2'); self.assertEqual(sheet.sheet_view.zoomScale, 85); self.assertFalse(sheet.sheet_view.showGridLines)
                self.assertEqual(sheet.print_area, f"'{sheet.title}'!$A$1:$H$20")
                self.assertEqual(sheet.print_title_rows, '$1:$1'); self.assertEqual(sheet.print_title_cols, '$A:$A')
                self.assertEqual(sheet.oddHeader.center.text, 'Independent copy'); self.assertEqual(sheet.row_breaks.brk[0].id, 10)
                self.assertTrue(sheet.protection.sheet); self.assertEqual(sheet.row_dimensions[2].height, 28)
                self.assertEqual(sheet.column_dimensions['B'].width, 21)
            self.assertTrue(all(ref.f.startswith("'Copy One'!") for ref in refs(first._charts[0]) if "[external.xlsx]" not in ref.f))
            self.assertTrue(all(ref.f.startswith("'Source Data'!") for ref in refs(original._charts[0]) if "[external.xlsx]" not in ref.f))
            external = next(ref for ref in refs(first._charts[0]) if "[external.xlsx]" in ref.f)
            self.assertEqual(external.f, "'[external.xlsx]Other'!$A$2:$A$3")
            self.assertEqual([point.v for point in external.numCache.pt], [77, 88])
            # Mutate only the first copy, then reopen all parts and verify the
            # original/second copy retained separate graph/data/settings objects.
            first['B2'] = 99
            first.tables['Sales_Copy2'].ref = 'A1:C2'
            first.data_validations.dataValidation[0].formula1 = 'B2>50'
            first._charts[0].ser[0].val.numRef.f = "'Copy One'!$B$2"
            first._charts[0].anchor._from.col = 12
            first._images[0].anchor._from.row = 30
            loaded.save(changed)
            reopened = load_workbook(changed)
            self.assertEqual(reopened['Source Data']['B2'].value, 10); self.assertEqual(reopened['Copy Two']['B2'].value, 10)
            self.assertEqual(reopened['Copy One']['B2'].value, 99)
            self.assertEqual(reopened['Source Data'].tables['Sales'].ref, 'A1:C3')
            self.assertEqual(reopened['Source Data'].data_validations.dataValidation[0].formula1, "'Source Data'!B2>0")
            self.assertEqual(reopened['Source Data']._charts[0].anchor._from.col, 9)
            self.assertEqual(reopened['Copy Two']._images[0].anchor._from.row, 14)
            self.assertEqual(reopened['Copy One']._images[0].anchor._from.row, 30)
            with zipfile.ZipFile(changed) as package:
                names = set(package.namelist())
                for prefix, count in [('xl/charts/chart', 3), ('xl/tables/table', 3), ('xl/media/image', 3)]:
                    self.assertEqual(sum(name.startswith(prefix) and '/_rels/' not in name for name in names), count)
                for name in names:
                    if not name.endswith('.rels'): continue
                    root = etree.fromstring(package.read(name))
                    parent = posixpath.dirname(posixpath.dirname(name))
                    for relationship in root:
                        if relationship.get('TargetMode') == 'External': continue
                        target = relationship.get('Target')
                        resolved = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(parent, target))
                        self.assertIn(resolved, names, (name, target))

    def test_pivot_cache_records_and_rich_text_copy_independently(self):
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        from openpyxl.chart.pivot import PivotSource
        from openpyxl.pivot.cache import CacheDefinition, CacheSource, WorksheetSource, CacheField, SharedItems
        from openpyxl.pivot.fields import Text, Number, Index
        from openpyxl.pivot.record import Record, RecordList
        from openpyxl.pivot.table import TableDefinition, Location, PivotField, RowColField, DataField, FieldItem
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            source, output, changed = (directory / name for name in ('pivots.xlsx', 'pivot-copies.xlsx', 'pivot-changed.xlsx'))
            wb = Workbook(); ws = wb.active; ws.title = 'Data'
            for row in [['Label', 'Value'], ['One', 10], ['Two', 20]]: ws.append(row)
            ws['D1'] = CellRichText('Plain ', TextBlock(InlineFont(b=True, color='FF0088'), 'bold'), ' tail')
            ws['D2'] = CellRichText(TextBlock(InlineFont(i=True), 'Second label'))
            rich_chart = BarChart(); rich_chart.add_data(Reference(ws, min_col=2, min_row=2, max_row=3))
            rich_chart.set_categories(Reference(ws, min_col=4, min_row=1, max_row=2)); ws.add_chart(rich_chart, 'J1')
            cache = CacheDefinition(cacheSource=CacheSource(type='worksheet', worksheetSource=WorksheetSource(sheet='Data', ref='A1:B3')),
                recordCount=2, cacheFields=[CacheField(name='Label', sharedItems=SharedItems(_fields=[Text(v='One'), Text(v='Two')])),
                                           CacheField(name='Value', sharedItems=SharedItems(_fields=[Number(v=10), Number(v=20)]))])
            cache.records = RecordList(r=[Record(_fields=[Index(v=0), Index(v=0)]), Record(_fields=[Index(v=1), Index(v=1)])])
            pivot = TableDefinition(name='Summary', cacheId=7, dataCaption='Values', location=Location(ref='F1:G4', firstHeaderRow=1, firstDataRow=2, firstDataCol=1),
                pivotFields=[PivotField(axis='axisRow', items=[FieldItem(x=0), FieldItem(x=1)]), PivotField(dataField=True)],
                rowFields=[RowColField(x=0)], dataFields=[DataField(name='Total', fld=1)])
            pivot.cache = cache; ws.add_pivot(pivot)
            rich_chart.pivotSource = PivotSource(name='[pivots.xlsx]Data!Summary', fmtId=0)
            wb.save(source)
            result = subprocess.run([sys.executable, str(SCRIPTS / 'xlsx_edit.py'), str(source), '--copy-sheet', 'Data:Copied', '--out', str(output)],
                                    capture_output=True, text=True, timeout=20)
            self.assertEqual(result.returncode, 0, result.stderr)
            loaded = load_workbook(output, rich_text=True)
            original, copied = loaded['Data'], loaded['Copied']
            self.assertEqual(len(copied._pivots), 1)
            first, second = original._pivots[0], copied._pivots[0]
            self.assertEqual(second.name, 'Summary_Copy'); self.assertNotEqual(first.cacheId, second.cacheId)
            self.assertEqual(original._charts[0].pivotSource.name, '[pivots.xlsx]Data!Summary')
            self.assertEqual(copied._charts[0].pivotSource.name, "'[pivots.xlsx]Copied'!Summary_Copy")
            self.assertIsNot(first.cache, second.cache); self.assertIsNot(first.cache.records, second.cache.records)
            self.assertEqual(first.cache.cacheSource.worksheetSource.sheet, 'Data')
            self.assertEqual(second.cache.cacheSource.worksheetSource.sheet, 'Copied')
            self.assertEqual(len(second.cache.records.r), 2)
            self.assertEqual(second.cache.records.r[1]._fields[1].v, 1)
            self.assertIsInstance(original['D1'].value, CellRichText); self.assertIsInstance(copied['D1'].value, CellRichText)
            self.assertTrue(original['D1'].value[1].font.b); self.assertEqual(copied['D1'].value[1].font.color.rgb, '00FF0088')
            second.cache.cacheFields[1].sharedItems._fields[0].v = 99
            copied['D1'].value[1].font.b = False
            copied['D1'].value[1].text = 'changed'
            self.assertTrue(original['D1'].value[1].font.b)
            self.assertEqual(first.cache.cacheFields[1].sharedItems._fields[0].v, 10)
            loaded.save(changed)
            reopened = load_workbook(changed, rich_text=True)
            self.assertEqual(reopened['Data']._pivots[0].cache.cacheFields[1].sharedItems._fields[0].v, 10)
            self.assertEqual(reopened['Copied']._pivots[0].cache.cacheFields[1].sharedItems._fields[0].v, 99)
            self.assertEqual(str(reopened['Data']['D1'].value), 'Plain bold tail')
            self.assertEqual(str(reopened['Copied']['D1'].value), 'Plain changed tail')
            self.assertTrue(reopened['Data']['D1'].value[1].font.b)
            with zipfile.ZipFile(changed) as package:
                names = package.namelist()
                self.assertEqual(sum(name.startswith('xl/pivotTables/pivotTable') and name.endswith('.xml') for name in names), 2)
                self.assertEqual(sum(name.startswith('xl/pivotCache/pivotCacheDefinition') and name.endswith('.xml') for name in names), 2)
                self.assertEqual(sum(name.startswith('xl/pivotCache/pivotCacheRecords') and name.endswith('.xml') for name in names), 2)
            # The other native writer must retain mixed formatting too.
            restructured = directory / 'pivot-restructured.xlsx'
            result = subprocess.run([sys.executable, str(SCRIPTS / 'xlsx_restructure.py'), str(source), '--sheet', 'Data', '--insert-rows', '6:1', '--out', str(restructured)],
                                    capture_output=True, text=True, timeout=20)
            self.assertEqual(result.returncode, 0, result.stderr)
            rebuilt = load_workbook(restructured, rich_text=True)['Data']
            rich = rebuilt['D1'].value
            self.assertIsInstance(rich, CellRichText); self.assertTrue(rich[1].font.b)
            self.assertEqual([point.v for point in rebuilt._charts[0].ser[0].cat.strRef.strCache.pt], ['Plain bold tail', 'Second label'])

    def test_in_memory_typed_formulas_and_name_conflicts_are_independent(self):
        wb = Workbook(); source = wb.active; source.title = 'Data'
        source['A1'] = ArrayFormula(ref='A1:A2', text='=Data!B1:B2')
        source['C1'] = DataTableFormula(ref='C1:D2', r1='B1')
        from openpyxl.cell.rich_text import CellRichText, TextBlock
        from openpyxl.cell.text import InlineFont
        source['E1'] = CellRichText(TextBlock(InlineFont(b=True), 'bold'))
        copied = copy_sheet(wb, 'Data', 'Copied')
        copied['A1'].value.text = '=42'; copied['C1'].value.r1 = 'B9'
        copied['E1'].value[0].font.b = False
        self.assertTrue(source['E1'].value[0].font.b)
        self.assertEqual(source['A1'].value.text, '=Data!B1:B2')
        self.assertEqual(source['C1'].value.r1, 'B1')
        before = wb.sheetnames[:]
        with self.assertRaises(ValueError): copy_sheet(wb, 'Data', 'copied')
        self.assertEqual(wb.sheetnames, before)
        with self.assertRaises(ValueError): copy_sheet(wb, 'Data', 'invalid/title')
        self.assertEqual(wb.sheetnames, before)


if __name__ == '__main__': unittest.main()
