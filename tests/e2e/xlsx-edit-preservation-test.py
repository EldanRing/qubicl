"""CLI regressions for reference-aware sheet names and non-overwriting appends."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.table import Table, TableColumn, TableFormula
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / 'skills/core/xlsx/scripts/xlsx_edit.py'


class EditPreservation(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(); self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name)

    def run_edit(self, wb, *arguments):
        source,output=self.root/'source.xlsx',self.root/'output.xlsx';wb.save(source)
        original=source.read_bytes()
        result=subprocess.run([sys.executable,str(SCRIPT),str(source),'--sheet','Data',*arguments,'--out',str(output)],capture_output=True,text=True,timeout=20)
        self.assertEqual(result.returncode,0,result.stderr)
        self.assertEqual(source.read_bytes(),original)
        return load_workbook(output),json.loads(result.stdout)

    def table(self, totals=False):
        wb=Workbook();ws=wb.active;ws.title='Data';report=wb.create_sheet('Report')
        ws.append(['Amount','Computed']);ws.append([2,'=A2*2'])
        if totals: ws.append(['Totals','=SUM(B2:B2)'])
        table=Table(displayName='Sales',ref=f'A1:B{3 if totals else 2}',totalsRowCount=int(totals))
        table.tableColumns=[TableColumn(id=1,name='Amount'),TableColumn(id=2,name='Computed',calculatedColumnFormula=TableFormula(attr_text='A2*2'))]
        if totals:
            table.tableColumns[1].totalsRowFunction='custom'
            table.tableColumns[1].totalsRowFormula=TableFormula(attr_text='SUM(B2:B2)')
        table.autoFilter=AutoFilter(ref='A1:B2');ws.add_table(table);ws['B2'].number_format='$0.00'
        return wb,ws,report,table

    def test_rename_updates_formula_bearing_surfaces_and_preserves_external_links(self):
        wb,ws,report,table=self.table(totals=True)
        report['A1']='=Data!A2+LEN("Data!A2")'
        report['A2']="='[external.xlsx]Data'!A2"
        report['A3']='=[external.xlsx]Data!A2'
        report['A4']='jump';report['A4'].hyperlink=Hyperlink(ref='A4',location="'Data'!A2")
        wb.defined_names.add(DefinedName('GlobalAmount',attr_text="'Data'!$A$2"))
        report.defined_names.add(DefinedName('LocalAmount',attr_text='Data!$A$2'))
        table.tableColumns[1].calculatedColumnFormula.attr_text='Data!A2*2'
        table.tableColumns[1].totalsRowFormula.attr_text='SUM(Data!B2:B2)'
        dv=DataValidation(type='list',formula1='Data!$A$2:$A$2');dv.add('C1:C2');report.add_data_validation(dv)
        report.conditional_formatting.add('C1:C2',FormulaRule(formula=['Data!A2>0']))
        chart=BarChart();chart.add_data(Reference(ws,min_col=1,min_row=2,max_row=2));report.add_chart(chart,'E2')
        ws.print_area='A1:B3'
        out,result=self.run_edit(wb,'--rename-sheet',"Data:New O'Brien")
        target="'New O''Brien'!";report=out['Report'];ws=out["New O'Brien"]
        self.assertEqual(result['sheet'],"New O'Brien")
        self.assertEqual(report['A1'].value,'='+target+'A2+LEN("Data!A2")')
        self.assertEqual(report['A2'].value,"='[external.xlsx]Data'!A2")
        self.assertEqual(report['A3'].value,'=[external.xlsx]Data!A2')
        self.assertEqual(report['A4'].hyperlink.location,target+'A2')
        self.assertEqual(out.defined_names['GlobalAmount'].attr_text,target+'$A$2')
        self.assertEqual(report.defined_names['LocalAmount'].attr_text,target+'$A$2')
        self.assertEqual(ws.tables['Sales'].tableColumns[1].calculatedColumnFormula.attr_text,target+'A2*2')
        self.assertEqual(ws.tables['Sales'].tableColumns[1].totalsRowFormula.attr_text,'SUM('+target+'B2:B2)')
        self.assertEqual(report.data_validations.dataValidation[0].formula1,target+'$A$2:$A$2')
        cf=next(iter(report.conditional_formatting));self.assertEqual(report.conditional_formatting[cf][0].formula,[target+'A2>0'])
        self.assertEqual(report._charts[0].series[0].val.numRef.f,target+'$A$2')
        self.assertEqual(str(ws.print_area),target+'$A$1:$B$3')

    def test_case_only_rename_and_newly_created_table_append(self):
        wb=Workbook();ws=wb.active;ws.title='Data';ws.append(['A','B']);ws.append([1,2]);ws['A3']='keep';wb.create_sheet('Report')['A1']='=Data!A2'
        out,_=self.run_edit(wb,'--rename-sheet','Data:data','--add-table','Sales:A1:B2','--table-append','Sales=[3,4]')
        self.assertIn('data',out.sheetnames);self.assertNotIn('data1',out.sheetnames)
        self.assertEqual(out['data']['A4'].value,'keep');self.assertEqual(out['data']['A3'].value,3)
        self.assertEqual(out['Report']['A1'].value,"='data'!A2")

    def test_append_moves_unrelated_cells_and_extends_calculated_chart_sources(self):
        wb,ws,report,table=self.table();ws['A3']='keep';ws['C3']=99;report['A1']='=Data!C3'
        chart=BarChart();chart.add_data(Reference(ws,min_col=1,min_row=2,max_row=2));report.add_chart(chart,'E2')
        wb.defined_names.add(DefinedName('Amounts',attr_text="'Data'!$A$2:$A$2"))
        report['C1']=77;other=BarChart();other.add_data(Reference(report,min_col=3,min_row=1,max_row=1));report.add_chart(other,'K2')
        out,_=self.run_edit(wb,'--table-append','Sales=[4]');ws=out['Data'];table=ws.tables['Sales']
        self.assertEqual(ws['A3'].value,4);self.assertEqual(ws['B3'].value,'=A3*2');self.assertEqual(ws['B3'].number_format,'$0.00')
        self.assertEqual(ws['A4'].value,'keep');self.assertEqual(ws['C4'].value,99);self.assertEqual(out['Report']['A1'].value,'=Data!C4')
        self.assertEqual(table.ref,'A1:B3');self.assertEqual(table.autoFilter.ref,'A1:B3')
        self.assertEqual(out.defined_names['Amounts'].attr_text,"'Data'!$A$2:$A$2")
        ref=out['Report']._charts[0].series[0].val.numRef
        self.assertEqual(ref.f,"'Data'!$A$2:$A$3")
        self.assertEqual([point.v for point in ref.numCache.pt],[2,4])
        self.assertEqual(out['Report']._charts[1].series[0].val.numRef.f,"'Report'!$C$1")

    def test_append_keeps_totals_below_data_and_expands_total_formula(self):
        wb,ws,report,table=self.table(totals=True);ws['A4']='after table';report['A1']='=Data!A4'
        out,_=self.run_edit(wb,'--table-append','Sales=[4]');ws=out['Data'];table=ws.tables['Sales']
        self.assertEqual(table.ref,'A1:B4');self.assertEqual(table.autoFilter.ref,'A1:B3')
        self.assertEqual(ws['A3'].value,4);self.assertEqual(ws['B3'].value,'=A3*2')
        self.assertEqual(ws['A4'].value,'Totals');self.assertEqual(ws['B4'].value,'=SUM(B2:B3)')
        self.assertEqual(table.tableColumns[1].totalsRowFormula.attr_text,'SUM(B2:B3)')
        self.assertEqual(ws['A5'].value,'after table');self.assertEqual(out['Report']['A1'].value,'=Data!A5')

    def test_append_preserves_row_local_ranges_and_calculated_templates(self):
        wb=Workbook();ws=wb.active;ws.title='Data';report=wb.create_sheet('Report')
        ws.append(['A','B','Total']);ws.append([1,2,'=SUM(A2:B2)'])
        table=Table(displayName='Sales',ref='A1:C2')
        table.tableColumns=[TableColumn(id=1,name='A'),TableColumn(id=2,name='B'),TableColumn(id=3,name='Total',calculatedColumnFormula=TableFormula(attr_text='SUM(A2:B2)'))]
        ws.add_table(table);report['A1']='=SUM(Data!A2:A2)';report['B1']='=Data!A2'
        out,_=self.run_edit(wb,'--table-append','Sales=[3,4]');ws=out['Data']
        self.assertEqual(ws['C2'].value,'=SUM(A2:B2)')
        self.assertEqual(ws['C3'].value,'=SUM(A3:B3)')
        self.assertEqual(ws.tables['Sales'].tableColumns[2].calculatedColumnFormula.attr_text,'SUM(A2:B2)')
        self.assertEqual(out['Report']['A1'].value,'=SUM(Data!A2:A2)')
        self.assertEqual(out['Report']['B1'].value,'=Data!A2')

    def test_editor_insert_shifts_local_names_and_validation_sources(self):
        wb=Workbook();ws=wb.active;ws.title='Data';ws.append(['A']);ws.append([1]);ws.append([2])
        ws.defined_names.add(DefinedName('Choices',attr_text='$A$2:$A$3'))
        dv=DataValidation(type='list',formula1="'Data'!$A$2:$A$3");dv.add('C2:C3');ws.add_data_validation(dv)
        out,_=self.run_edit(wb,'--insert-rows','2:1');ws=out['Data']
        self.assertEqual(ws.defined_names['Choices'].attr_text,'$A$3:$A$4')
        self.assertEqual(ws.data_validations.dataValidation[0].formula1,"'Data'!$A$3:$A$4")
        self.assertEqual(str(ws.data_validations.dataValidation[0].sqref),'C3:C4')


if __name__=='__main__':unittest.main()
